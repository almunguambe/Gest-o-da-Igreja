import os

# =========================================================================
# 1. ATUALIZAÇÃO DO APP.PY COM TODAS AS NOVAS TABELAS E ROTAS
# =========================================================================
app_code = """from flask import Flask, render_template, request, redirect, url_for, session, send_file
import sqlite3
from datetime import datetime
import json
import io
import os
import zipfile
from werkzeug.utils import secure_filename
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# ReportLab para PDFs de alta precisão
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import cm, mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas

app = Flask(__name__)
app.secret_key = "iead_chicuque_chave_super_segura_2026"
DB_NAME = "gestao_chicuque.db"
UPLOAD_FOLDER = os.path.join("static", "uploads")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

def init_db():
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    
    # 1. Utilizadores
    c.execute('''CREATE TABLE IF NOT EXISTS usuarios (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        usuario TEXT UNIQUE NOT NULL,
        senha TEXT NOT NULL,
        cargo TEXT NOT NULL
    )''')
    if c.execute("SELECT COUNT(*) FROM usuarios").fetchone()[0] == 0:
        c.execute("INSERT INTO usuarios (usuario, senha, cargo) VALUES (?, ?, ?)",
                  ('admin', 'chicuque123', 'Pastor Presidente'))

    # 2. Membros Completo
    c.execute('''CREATE TABLE IF NOT EXISTS membros (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nome TEXT NOT NULL,
        telefone TEXT,
        genero TEXT,
        data_nascimento TEXT,
        faixa_etaria TEXT,
        naturalidade TEXT,
        bairro TEXT,
        filiacao TEXT,
        tipo_documento TEXT,
        numero_documento TEXT,
        ano_conversao INTEGER,
        data_batismo TEXT,
        posicao_atual TEXT,
        progressoes TEXT,
        departamento TEXT,
        foto_path TEXT,
        observacoes TEXT,
        data_registo TEXT
    )''')

    # Migração segura de colunas para membros
    c.execute("PRAGMA table_info(membros)")
    cols = [col[1] for col in c.fetchall()]
    novas = [
        ('genero', 'TEXT'), ('data_nascimento', 'TEXT'), ('faixa_etaria', 'TEXT'),
        ('naturalidade', 'TEXT'), ('filiacao', 'TEXT'), ('tipo_documento', 'TEXT'),
        ('numero_documento', 'TEXT'), ('ano_conversao', 'INTEGER'), ('posicao_atual', 'TEXT'),
        ('progressoes', 'TEXT'), ('foto_path', 'TEXT')
    ]
    for n, t in novas:
        if n not in cols:
            c.execute(f"ALTER TABLE membros ADD COLUMN {n} {t}")

    # 3. Financeiro
    c.execute('''CREATE TABLE IF NOT EXISTS financeiro (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        tipo TEXT CHECK(tipo IN ('Entrada', 'Saída')),
        local_movimento TEXT CHECK(local_movimento IN ('Caixa', 'Banco')) DEFAULT 'Caixa',
        departamento TEXT DEFAULT 'Geral',
        categoria TEXT NOT NULL,
        valor REAL NOT NULL,
        data_movimento TEXT NOT NULL,
        dia INTEGER,
        mes INTEGER,
        ano INTEGER,
        data_registo TEXT NOT NULL,
        membro_id INTEGER,
        descricao TEXT,
        FOREIGN KEY (membro_id) REFERENCES membros (id)
    )''')

    # 4. Transferências
    c.execute('''CREATE TABLE IF NOT EXISTS transferencias (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        data_movimento TEXT NOT NULL,
        origem_local TEXT NOT NULL,
        origem_depto TEXT NOT NULL,
        destino_local TEXT NOT NULL,
        destino_depto TEXT NOT NULL,
        valor REAL NOT NULL,
        motivo TEXT NOT NULL,
        data_registo TEXT NOT NULL
    )''')

    # 5. Casamentos e Óbitos
    c.execute('''CREATE TABLE IF NOT EXISTS casamentos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        noivo TEXT NOT NULL,
        noiva TEXT NOT NULL,
        data_casamento TEXT NOT NULL,
        pastor_oficiante TEXT,
        data_registo TEXT
    )''')

    c.execute('''CREATE TABLE IF NOT EXISTS mortes (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nome_falecido TEXT NOT NULL,
        data_falecimento TEXT NOT NULL,
        observacoes TEXT,
        data_registo TEXT
    )''')

    # 6. Listas Configuráveis
    c.execute('''CREATE TABLE IF NOT EXISTS departamentos_lista (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nome TEXT UNIQUE NOT NULL
    )''')
    c.execute('''CREATE TABLE IF NOT EXISTS categorias_financeiras (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        tipo TEXT CHECK(tipo IN ('Entrada', 'Saída')) NOT NULL,
        nome TEXT NOT NULL
    )''')
    c.execute('''CREATE TABLE IF NOT EXISTS zonas_lista (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nome TEXT UNIQUE NOT NULL
    )''')

    # 7. NOVOS MÓDULOS EXPANDIDOS
    # A. Cultos & Presenças
    c.execute('''CREATE TABLE IF NOT EXISTS cultos_frequencia (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        data_culto TEXT NOT NULL,
        tipo_culto TEXT NOT NULL,
        homens INTEGER DEFAULT 0,
        mulheres INTEGER DEFAULT 0,
        jovens INTEGER DEFAULT 0,
        criancas INTEGER DEFAULT 0,
        visitantes INTEGER DEFAULT 0,
        novos_convertidos INTEGER DEFAULT 0,
        total_presentes INTEGER DEFAULT 0,
        pregador TEXT,
        tema_mensagem TEXT,
        data_registo TEXT
    )''')

    # B. Novos Convertidos / Discipulado
    c.execute('''CREATE TABLE IF NOT EXISTS novos_convertidos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nome TEXT NOT NULL,
        telefone TEXT,
        bairro TEXT,
        data_decisao TEXT NOT NULL,
        culto_origem TEXT,
        quem_convidou TEXT,
        status_discipulado TEXT DEFAULT 'Decisão Inicial',
        observacoes TEXT
    )''')

    # C. Património & Inventário da Igreja
    c.execute('''CREATE TABLE IF NOT EXISTS patrimonio (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        item TEXT NOT NULL,
        departamento TEXT DEFAULT 'Geral',
        quantidade INTEGER DEFAULT 1,
        estado_conservacao TEXT DEFAULT 'Bom',
        localizacao TEXT,
        observacoes TEXT
    )''')

    # D. Escalas de Culto / Púlpito
    c.execute('''CREATE TABLE IF NOT EXISTS escalas (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        data_escala TEXT NOT NULL,
        tipo_culto TEXT NOT NULL,
        dirigente TEXT,
        pregador TEXT,
        leitura_palavra TEXT,
        louvor_grupo TEXT,
        diaconos_servico TEXT,
        observacoes TEXT
    )''')

    # E. Metas e Campanhas de Construção
    c.execute('''CREATE TABLE IF NOT EXISTS campanhas_metas (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nome_campanha TEXT NOT NULL,
        departamento TEXT DEFAULT 'Construção',
        valor_meta REAL NOT NULL,
        status TEXT DEFAULT 'Ativa'
    )''')

    deptos = [
        ('Activista',), ('Juventude',), ('Mulher (Senhoras)',), 
        ('Boa Esperança (Crianças)',), ('Homens / Obreiros',), 
        ('Louvor / Música',), ('Ação Social',), ('Construção',)
    ]
    for d in deptos:
        c.execute("INSERT OR IGNORE INTO departamentos_lista (nome) VALUES (?)", d)

    if c.execute("SELECT COUNT(*) FROM categorias_financeiras").fetchone()[0] == 0:
        padroes = [
            ('Entrada', 'Dízimo'), ('Entrada', 'Oferta'), ('Entrada', 'Doação / Voto'), ('Entrada', 'Campanha de Construção'),
            ('Saída', 'Manutenção do Templo'), ('Saída', 'Ação Social / Ajudas'), ('Saída', 'Combustível / Transporte'),
            ('Saída', 'Água e Eletricidade'), ('Saída', 'Material de Culto')
        ]
        c.executemany("INSERT INTO categorias_financeiras (tipo, nome) VALUES (?, ?)", padroes)

    if c.execute("SELECT COUNT(*) FROM zonas_lista").fetchone()[0] == 0:
        c.executemany("INSERT INTO zonas_lista (nome) VALUES (?)",
                      [('Chicuque Sede',), ('Maxixe Cidade',), ('Nhacoongo',), ('Conguiana',), ('Bairro 1',)])

    # Campanha Padrão Inicial
    if c.execute("SELECT COUNT(*) FROM campanhas_metas").fetchone()[0] == 0:
        c.execute("INSERT INTO campanhas_metas (nome_campanha, departamento, valor_meta, status) VALUES (?, ?, ?, ?)",
                  ('Campanha de Obras e Ampliação do Templo', 'Construção', 100000.0, 'Ativa'))

    conn.commit()
    conn.close()

init_db()

def get_db():
    conn = sqlite3.connect(DB_NAME)
    conn.row_factory = sqlite3.Row
    return conn

def is_admin():
    cargo = session.get('cargo', '')
    return session.get('usuario') == 'admin' or 'Pastor' in cargo

def can_cadastro():
    cargo = session.get('cargo', '')
    return is_admin() or 'Secretário' in cargo or 'Líder' in cargo

def can_tesouraria():
    cargo = session.get('cargo', '')
    return is_admin() or 'Tesoureiro' in cargo

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        usuario = request.form.get('usuario', '').strip()
        senha = request.form.get('senha', '').strip()
        conn = get_db()
        user = conn.execute("SELECT * FROM usuarios WHERE usuario = ? AND senha = ?", (usuario, senha)).fetchone()
        conn.close()
        if user:
            session['usuario'] = user['usuario']
            session['cargo'] = user['cargo']
            return redirect(url_for('dashboard'))
        return render_template('login.html', erro="Credenciais incorretas.")
    return render_template('login.html', erro=None)

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.route('/')
def dashboard():
    if 'usuario' not in session:
        return redirect(url_for('login'))
        
    conn = get_db()
    total_membros = conn.execute("SELECT COUNT(*) FROM membros").fetchone()[0]
    total_casamentos = conn.execute("SELECT COUNT(*) FROM casamentos").fetchone()[0]
    total_mortes = conn.execute("SELECT COUNT(*) FROM mortes").fetchone()[0]

    # Saldos
    entrada_caixa = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Entrada' AND (local_movimento = 'Caixa' OR local_movimento IS NULL)").fetchone()[0] or 0.0
    saida_caixa = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Saída' AND (local_movimento = 'Caixa' OR local_movimento IS NULL)").fetchone()[0] or 0.0
    saldo_caixa = entrada_caixa - saida_caixa

    entrada_banco = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Entrada' AND local_movimento = 'Banco'").fetchone()[0] or 0.0
    saida_banco = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Saída' AND local_movimento = 'Banco'").fetchone()[0] or 0.0
    saldo_banco = entrada_banco - saida_banco
    saldo_total = (entrada_caixa + entrada_banco) - (saida_caixa + saida_banco)

    # Aniversariantes do Dia e do Mês
    hoje_md = datetime.now().strftime("-%m-%d")
    mes_atual_str = datetime.now().strftime("-%m-")
    aniversariantes_dia = conn.execute("SELECT nome, data_nascimento, telefone, departamento FROM membros WHERE data_nascimento LIKE ?", (f"%{hoje_md}",)).fetchall()
    aniversariantes_mes = conn.execute("SELECT nome, data_nascimento, telefone, departamento FROM membros WHERE data_nascimento LIKE ? ORDER BY data_nascimento ASC", (f"%{mes_atual_str}%",)).fetchall()

    # Campanhas de Construção com Arrecadação Real
    campanhas_raw = conn.execute("SELECT * FROM campanhas_metas WHERE status = 'Ativa'").fetchall()
    campanhas = []
    for c in campanhas_raw:
        arrecadado = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Entrada' AND (categoria LIKE '%Construção%' OR departamento = ?)", (c['departamento'],)).fetchone()[0] or 0.0
        perc = min(round((arrecadado / c['valor_meta']) * 100, 1), 100) if c['valor_meta'] > 0 else 0
        campanhas.append({
            'id': c['id'], 'nome': c['nome_campanha'], 'meta': c['valor_meta'],
            'arrecadado': arrecadado, 'percentual': perc
        })

    # Gráficos Analíticos de Tesouraria
    ano_atual = datetime.now().year
    evolucao_entradas, evolucao_saidas = [], []
    for m in range(1, 13):
        e_mes = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Entrada' AND mes = ? AND ano = ?", (m, ano_atual)).fetchone()[0] or 0.0
        s_mes = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Saída' AND mes = ? AND ano = ?", (m, ano_atual)).fetchone()[0] or 0.0
        evolucao_entradas.append(e_mes)
        evolucao_saidas.append(s_mes)

    deptos_lista = [d['nome'] for d in conn.execute("SELECT nome FROM departamentos_lista ORDER BY nome ASC").fetchall()]
    deptos_financeiro = ['Geral'] + [d for d in deptos_lista if d != 'Geral']
    depto_fin_labels, depto_fin_entradas, depto_fin_saidas = [], [], []
    for d in deptos_financeiro:
        ent = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Entrada' AND departamento = ?", (d,)).fetchone()[0] or 0.0
        sai = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Saída' AND departamento = ?", (d,)).fetchone()[0] or 0.0
        if ent > 0 or sai > 0 or d in ['Geral', 'Activista', 'Juventude', 'Mulher (Senhoras)', 'Boa Esperança (Crianças)']:
            depto_fin_labels.append(d)
            depto_fin_entradas.append(ent)
            depto_fin_saidas.append(sai)

    fontes_rows = conn.execute("SELECT categoria, SUM(valor) as total FROM financeiro WHERE tipo = 'Entrada' GROUP BY categoria ORDER BY total DESC").fetchall()
    fontes_labels = [r['categoria'] for r in fontes_rows] if fontes_rows else ['Sem Entradas']
    fontes_valores = [r['total'] for r in fontes_rows] if fontes_rows else [0]

    saidas_rows = conn.execute("SELECT categoria, SUM(valor) as total FROM financeiro WHERE tipo = 'Saída' GROUP BY categoria ORDER BY total DESC").fetchall()
    saidas_labels = [r['categoria'] for r in saidas_rows] if saidas_rows else ['Sem Saídas']
    saidas_valores = [r['total'] for r in saidas_rows] if saidas_rows else [0]

    # Listagens
    todos_membros = conn.execute("SELECT * FROM membros ORDER BY id DESC").fetchall()
    todas_financas = conn.execute("SELECT * FROM financeiro ORDER BY id DESC").fetchall()
    todos_casamentos = conn.execute("SELECT * FROM casamentos ORDER BY id DESC").fetchall()
    todas_mortes = conn.execute("SELECT * FROM mortes ORDER BY id DESC").fetchall()
    todos_cultos = conn.execute("SELECT * FROM cultos_frequencia ORDER BY id DESC LIMIT 25").fetchall()
    todos_convertidos = conn.execute("SELECT * FROM novos_convertidos ORDER BY id DESC").fetchall()
    todo_patrimonio = conn.execute("SELECT * FROM patrimonio ORDER BY departamento, item ASC").fetchall()
    todas_escalas = conn.execute("SELECT * FROM escalas ORDER BY data_escala DESC LIMIT 20").fetchall()
    lista_usuarios = conn.execute("SELECT id, usuario, cargo FROM usuarios ORDER BY id ASC").fetchall()
    ultimas_transferencias = conn.execute("SELECT * FROM transferencias ORDER BY id DESC LIMIT 20").fetchall()

    lista_deptos = conn.execute("SELECT * FROM departamentos_lista ORDER BY nome ASC").fetchall()
    lista_categorias = conn.execute("SELECT * FROM categorias_financeiras ORDER BY tipo, nome ASC").fetchall()
    lista_zonas = conn.execute("SELECT * FROM zonas_lista ORDER BY nome ASC").fetchall()

    categorias_json = json.dumps([{'tipo': c['tipo'], 'nome': c['nome']} for c in lista_categorias])
    membros_json = json.dumps([dict(m) for m in todos_membros])

    conn.close()

    alerta_duplicado = session.pop('alerta_duplicado', None)
    sucesso_cadastro = session.pop('sucesso_cadastro', None)

    return render_template('dashboard.html',
                           pode_cadastro=can_cadastro(),
                           pode_tesouraria=can_tesouraria(),
                           e_admin=is_admin(),
                           alerta_duplicado=alerta_duplicado,
                           sucesso_cadastro=sucesso_cadastro,
                           total_membros=total_membros,
                           total_casamentos=total_casamentos,
                           total_mortes=total_mortes,
                           saldo_caixa=saldo_caixa,
                           saldo_banco=saldo_banco,
                           saldo_total=saldo_total,
                           aniversariantes_dia=aniversariantes_dia,
                           aniversariantes_mes=aniversariantes_mes,
                           campanhas=campanhas,
                           todos_membros=todos_membros,
                           todas_financas=todas_financas,
                           todos_casamentos=todos_casamentos,
                           todas_mortes=todas_mortes,
                           todos_cultos=todos_cultos,
                           todos_convertidos=todos_convertidos,
                           todo_patrimonio=todo_patrimonio,
                           todas_escalas=todas_escalas,
                           lista_deptos=lista_deptos,
                           lista_categorias=lista_categorias,
                           lista_zonas=lista_zonas,
                           lista_usuarios=lista_usuarios,
                           ultimas_transferencias=ultimas_transferencias,
                           categorias_json=categorias_json,
                           membros_json=membros_json,
                           evolucao_entradas=json.dumps(evolucao_entradas),
                           evolucao_saidas=json.dumps(evolucao_saidas),
                           depto_fin_labels=json.dumps(depto_fin_labels),
                           depto_fin_entradas=json.dumps(depto_fin_entradas),
                           depto_fin_saidas=json.dumps(depto_fin_saidas),
                           fontes_labels=json.dumps(fontes_labels),
                           fontes_valores=json.dumps(fontes_valores),
                           saidas_labels=json.dumps(saidas_labels),
                           saidas_valores=json.dumps(saidas_valores))

# CADASTRO COM BLOQUEIO DE DUPLICADOS
@app.route('/membros/novo', methods=['POST'])
def novo_membro():
    if not can_cadastro(): return redirect(url_for('dashboard'))
    nome = request.form.get('nome', '').strip()
    num_doc = request.form.get('numero_documento', '').strip()
    tipo_doc = request.form.get('tipo_documento', 'BI')

    conn = get_db()
    if num_doc:
        existente_doc = conn.execute("SELECT id, nome FROM membros WHERE numero_documento = ? AND numero_documento != ''", (num_doc,)).fetchone()
        if existente_doc:
            conn.close()
            session['alerta_duplicado'] = f"Atenção: Já existe um membro registado com o documento nº {num_doc} ({existente_doc['nome']})."
            return redirect(url_for('dashboard'))

    existente_nome = conn.execute("SELECT id FROM membros WHERE LOWER(TRIM(nome)) = LOWER(?)", (nome,)).fetchone()
    if existente_nome:
        conn.close()
        session['alerta_duplicado'] = f"Atenção: O membro '{nome}' já se encontra registado."
        return redirect(url_for('dashboard'))

    foto_path = ""
    if 'foto' in request.files:
        foto = request.files['foto']
        if foto and foto.filename:
            ext = foto.filename.rsplit('.', 1)[-1].lower()
            if ext in ['png', 'jpg', 'jpeg', 'webp']:
                nome_foto = f"membro_{int(datetime.now().timestamp())}.{ext}"
                salvar_em = os.path.join(app.config['UPLOAD_FOLDER'], nome_foto)
                foto.save(salvar_em)
                foto_path = f"/static/uploads/{nome_foto}"

    ano_conv = request.form.get('ano_conversao')
    ano_conv_val = int(ano_conv) if ano_conv and ano_conv.isdigit() else None

    conn.execute('''
        INSERT INTO membros (
            nome, telefone, genero, data_nascimento, faixa_etaria,
            naturalidade, bairro, filiacao, tipo_documento, numero_documento,
            ano_conversao, data_batismo, posicao_atual, progressoes,
            departamento, foto_path, observacoes, data_registo
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', (
        nome, request.form.get('telefone', '').strip(), request.form.get('genero', 'Masculino'),
        request.form.get('data_nascimento', ''), request.form.get('faixa_etaria', 'Adulto'),
        request.form.get('naturalidade', '').strip(), request.form.get('bairro', ''),
        request.form.get('filiacao', '').strip(), tipo_doc, num_doc, ano_conv_val,
        request.form.get('data_batismo', ''), request.form.get('posicao_atual', 'Membro em Comunhão'),
        request.form.get('progressoes', '').strip(), request.form.get('departamento', 'Geral'),
        foto_path, request.form.get('observacoes', '').strip(), datetime.now().strftime("%d/%m/%Y")
    ))
    conn.commit()
    conn.close()
    session['sucesso_cadastro'] = f"Membro '{nome}' registado com sucesso!"
    return redirect(url_for('dashboard'))

# 1. ROTAS: REGISTO DE CULTOS & PRESENÇAS
@app.route('/cultos/novo', methods=['POST'])
def novo_culto():
    if not can_cadastro(): return redirect(url_for('dashboard'))
    h = int(request.form.get('homens') or 0)
    m = int(request.form.get('mulheres') or 0)
    j = int(request.form.get('jovens') or 0)
    c = int(request.form.get('criancas') or 0)
    v = int(request.form.get('visitantes') or 0)
    nc = int(request.form.get('novos_convertidos') or 0)
    total = h + m + j + c + v

    conn = get_db()
    conn.execute('''INSERT INTO cultos_frequencia (data_culto, tipo_culto, homens, mulheres, jovens, criancas, visitantes, novos_convertidos, total_presentes, pregador, tema_mensagem, data_registo)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
                 (request.form['data_culto'], request.form['tipo_culto'], h, m, j, c, v, nc, total,
                  request.form.get('pregador', ''), request.form.get('tema_mensagem', ''), datetime.now().strftime("%d/%m/%Y")))
    conn.commit()
    conn.close()
    session['sucesso_cadastro'] = "Culto e frequência registados com sucesso!"
    return redirect(url_for('dashboard'))

# 2. ROTAS: NOVOS CONVERTIDOS & DISCIPULADO
@app.route('/convertidos/novo', methods=['POST'])
def novo_convertido():
    if not can_cadastro(): return redirect(url_for('dashboard'))
    conn = get_db()
    conn.execute('''INSERT INTO novos_convertidos (nome, telefone, bairro, data_decisao, culto_origem, quem_convidou, status_discipulado, observacoes)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)''',
                 (request.form['nome'].strip(), request.form.get('telefone', '').strip(), request.form.get('bairro', ''),
                  request.form['data_decisao'], request.form.get('culto_origem', ''), request.form.get('quem_convidou', ''),
                  request.form.get('status_discipulado', 'Decisão Inicial'), request.form.get('observacoes', '')))
    conn.commit()
    conn.close()
    session['sucesso_cadastro'] = f"Novo convertido '{request.form['nome']}' registado para discipulado!"
    return redirect(url_for('dashboard'))

@app.route('/convertidos/atualizar_status/<int:id>', methods=['POST'])
def atualizar_status_convertido(id):
    if not can_cadastro(): return redirect(url_for('dashboard'))
    novo_status = request.form.get('status_discipulado')
    conn = get_db()
    conn.execute("UPDATE novos_convertidos SET status_discipulado = ? WHERE id = ?", (novo_status, id))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

# 3. ROTAS: PATRIMÓNIO & INVENTÁRIO DO TEMPLO
@app.route('/patrimonio/novo', methods=['POST'])
def novo_patrimonio():
    if not is_admin(): return redirect(url_for('dashboard'))
    conn = get_db()
    conn.execute('''INSERT INTO patrimonio (item, departamento, quantidade, estado_conservacao, localizacao, observacoes)
                    VALUES (?, ?, ?, ?, ?, ?)''',
                 (request.form['item'].strip(), request.form.get('departamento', 'Geral'),
                  int(request.form.get('quantidade') or 1), request.form.get('estado_conservacao', 'Bom'),
                  request.form.get('localizacao', 'Templo Sede'), request.form.get('observacoes', '')))
    conn.commit()
    conn.close()
    session['sucesso_cadastro'] = "Item patrimonial inventariado com sucesso!"
    return redirect(url_for('dashboard'))

# 4. ROTAS: ESCALA DE CULTO & GERADOR EM PDF
@app.route('/escalas/novo', methods=['POST'])
def nova_escala():
    if not is_admin() and not can_cadastro(): return redirect(url_for('dashboard'))
    conn = get_db()
    conn.execute('''INSERT INTO escalas (data_escala, tipo_culto, dirigente, pregador, leitura_palavra, louvor_grupo, diaconos_servico, observacoes)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)''',
                 (request.form['data_escala'], request.form['tipo_culto'], request.form.get('dirigente', ''),
                  request.form.get('pregador', ''), request.form.get('leitura_palavra', ''),
                  request.form.get('louvor_grupo', ''), request.form.get('diaconos_servico', ''), request.form.get('observacoes', '')))
    conn.commit()
    conn.close()
    session['sucesso_cadastro'] = "Escala de culto registada com sucesso!"
    return redirect(url_for('dashboard'))

@app.route('/escalas/pdf/<int:id>')
def escala_pdf(id):
    conn = get_db()
    e = conn.execute("SELECT * FROM escalas WHERE id = ?", (id,)).fetchone()
    conn.close()
    if not e: return redirect(url_for('dashboard'))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2*cm, rightMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm)
    elementos = []

    elementos.append(Paragraph("<font color='#1e3a8a' size=14><b>IGREJA EVANGÉLICA ASSEMBLEIA DE DEUS</b></font>", ParagraphStyle('H1', alignment=1)))
    elementos.append(Paragraph("<font color='#475569' size=9><b>CONGREGAÇÃO DE CHICUQUE • ESCALA OFICIAL DE CULTO</b></font>", ParagraphStyle('H2', alignment=1, spaceAfter=20)))
    elementos.append(Paragraph(f"<font color='#0f172a' size=14><b>ESCALA DE {e['tipo_culto'].upper()} — {e['data_escala']}</b></font>", ParagraphStyle('T', alignment=1, spaceAfter=20)))

    dados = [
        ["Função Litúrgica", "Obreiro / Equipa Responsável"],
        ["Dirigente do Culto:", e['dirigente'] or 'A definir'],
        ["Ministração da Palavra (Pregador):", e['pregador'] or 'A definir'],
        ["Leitura Bíblica Inicial:", e['leitura_palavra'] or 'A definir'],
        ["Grupo de Louvor / Música:", e['louvor_grupo'] or 'Geral'],
        ["Diáconos / Serviço de Recepção:", e['diaconos_servico'] or 'A definir'],
        ["Notas Pastorais:", e['observacoes'] or 'Comparecer 30 minutos antes com traje solene.']
    ]
    t = Table(dados, colWidths=[7*cm, 9*cm])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#1e3a8a")),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 8),
        ('TOPPADDING', (0,0), (-1,-1), 8),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
    ]))
    elementos.append(t)
    elementos.append(Spacer(1, 2*cm))

    ass = [
        [
            Paragraph("____________________________________<br/><b>Secretaria / Direção de Culto</b>", ParagraphStyle('A1', alignment=1)),
            Paragraph("____________________________________<br/><b>Visto do Pastor Presidente</b>", ParagraphStyle('A2', alignment=1))
        ]
    ]
    elementos.append(Table(ass, colWidths=[8*cm, 8*cm]))

    doc.build(elementos)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"Escala_Culto_{e['data_escala']}.pdf", mimetype="application/pdf")

# 5. ROTAS DE DOCUMENTOS: CARTÃO, CERTIFICADOS, BALANCETE E BACKUP
@app.route('/membro/cartao_pdf/<int:id>')
def cartao_membro_pdf(id):
    if not can_cadastro(): return redirect(url_for('dashboard'))
    conn = get_db()
    m = conn.execute("SELECT * FROM membros WHERE id = ?", (id,)).fetchone()
    conn.close()
    if not m: return redirect(url_for('dashboard'))

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(8.5*cm, 5.4*cm))
    c.setFillColor(colors.HexColor("#0f172a"))
    c.rect(0, 0, 8.5*cm, 5.4*cm, fill=1, stroke=0)
    c.setFillColor(colors.HexColor("#1e3a8a"))
    c.rect(0, 4.2*cm, 8.5*cm, 1.2*cm, fill=1, stroke=0)
    c.setFillColor(colors.HexColor("#f59e0b"))
    c.rect(0, 4.15*cm, 8.5*cm, 0.06*cm, fill=1, stroke=0)

    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 7.5)
    c.drawCentredString(4.25*cm, 4.85*cm, "IGREJA EVANGÉLICA ASSEMBLEIA DE DEUS")
    c.setFont("Helvetica", 6)
    c.setFillColor(colors.HexColor("#93c5fd"))
    c.drawCentredString(4.25*cm, 4.45*cm, "CONGREGAÇÃO DE CHICUQUE • CARTÃO DE MEMBRO")

    c.setFillColor(colors.white)
    c.roundRect(0.5*cm, 1.2*cm, 2.0*cm, 2.6*cm, 3, fill=1, stroke=1)
    
    foto_desenhada = False
    if m['foto_path']:
        caminho_real = m['foto_path'].lstrip('/')
        if os.path.exists(caminho_real):
            try:
                c.drawImage(caminho_real, 0.55*cm, 1.25*cm, width=1.9*cm, height=2.5*cm, preserveAspectRatio=True)
                foto_desenhada = True
            except Exception: pass
            
    if not foto_desenhada:
        c.setFillColor(colors.HexColor("#1e3a8a"))
        c.setFont("Helvetica-Bold", 14)
        c.drawCentredString(1.5*cm, 2.3*cm, m['nome'][:2].upper())

    c.setFillColor(colors.HexColor("#f8fafc"))
    c.setFont("Helvetica-Bold", 8.5)
    nome_curto = m['nome'] if len(m['nome']) <= 25 else m['nome'][:23] + "..."
    c.drawString(2.8*cm, 3.5*cm, nome_curto)
    c.setFont("Helvetica-Bold", 6.5)
    c.setFillColor(colors.HexColor("#fbbf24"))
    c.drawString(2.8*cm, 3.1*cm, f"CARGO: {m['posicao_atual'] or 'Membro'}")
    c.setFont("Helvetica", 6)
    c.setFillColor(colors.HexColor("#cbd5e1"))
    c.drawString(2.8*cm, 2.65*cm, f"Ministério: {m['departamento'] or 'Geral'}")
    c.drawString(2.8*cm, 2.25*cm, f"Doc: {m['tipo_documento'] or 'BI'} - {m['numero_documento'] or 'S/N'}")
    c.drawString(2.8*cm, 1.85*cm, f"Batismo: {m['data_batismo'] or '-'} | Conv: {m['ano_conversao'] or '-'}")
    c.drawString(2.8*cm, 1.45*cm, f"Bairro/Zona: {m['bairro'] or 'Chicuque'}")

    c.setFont("Helvetica", 5)
    c.setFillColor(colors.HexColor("#64748b"))
    c.drawString(0.5*cm, 0.4*cm, f"ID #{m['id']} • Válido com carimbo pastoral")
    c.drawRightString(8.0*cm, 0.4*cm, "Pastor Presidente")
    c.setStrokeColor(colors.HexColor("#475569"))
    c.line(5.5*cm, 0.65*cm, 8.0*cm, 0.65*cm)

    c.showPage()
    c.save()
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"Cartao_Membro_{m['nome'].replace(' ', '_')}.pdf", mimetype="application/pdf")

@app.route('/membro/certificado_pdf/<int:id>/<tipo>')
def certificado_membro_pdf(id, tipo):
    if not can_cadastro(): return redirect(url_for('dashboard'))
    conn = get_db()
    m = conn.execute("SELECT * FROM membros WHERE id = ?", (id,)).fetchone()
    conn.close()
    if not m: return redirect(url_for('dashboard'))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=2*cm, rightMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm)
    titulo_cert = {
        'batismo': "CERTIFICADO DE BATISMO NAS ÁGUAS",
        'apresentacao': "CERTIFICADO DE APRESENTAÇÃO AO SENHOR",
        'recomendacao': "CARTA DE RECOMENDAÇÃO PASTORAL"
    }.get(tipo, "CERTIFICADO ECLESIÁSTICO")

    versiculo = {
        'batismo': '"Portanto ide, fazei discípulos de todas as nações, batizando-os em nome do Pai, e do Filho, e do Espírito Santo." — Mateus 28:19',
        'apresentacao': '"Deixai vir a mim os pequeninos e não os impeçais, porque dos tais é o Reino de Deus." — Marcos 10:14',
        'recomendacao': '"Recomendo-vos a nossa irmã, que vos seja apresentada no Senhor de um modo digno dos santos." — Romanos 16:1-2'
    }.get(tipo, '')

    elementos = [
        Paragraph("<font color='#1e3a8a' size=14><b>IGREJA EVANGÉLICA ASSEMBLEIA DE DEUS</b></font>", ParagraphStyle('H1', alignment=1, spaceAfter=4)),
        Paragraph("<font color='#475569' size=10><b>CONGREGAÇÃO DE CHICUQUE • INHAMBANE, MOÇAMBIQUE</b></font>", ParagraphStyle('H2', alignment=1, spaceAfter=15)),
        Spacer(1, 0.5*cm),
        Paragraph(f"<font color='#d97706' size=20><b>{titulo_cert}</b></font>", ParagraphStyle('T', alignment=1, spaceAfter=20))
    ]

    if tipo == 'batismo':
        corpo = f"Certificamos que <b>{m['nome'].upper()}</b>, nascido(a) em <b>{m['data_nascimento'] or '---'}</b>, natural de <b>{m['naturalidade'] or 'Moçambique'}</b>, portador(a) do documento <b>{m['tipo_documento'] or 'BI'} nº {m['numero_documento'] or '---'}</b>, desceu às águas batismais no dia <b>{m['data_batismo'] or datetime.now().strftime('%d/%m/%Y')}</b>, confessando publicamente a Jesus Cristo como seu Salvador."
    elif tipo == 'apresentacao':
        corpo = f"Certificamos que a criança <b>{m['nome'].upper()}</b>, nascida aos <b>{m['data_nascimento'] or '---'}</b>, filha de <b>{m['filiacao'] or 'seus pais'}</b>, foi solenemente apresentada a Deus nesta congregação, nos termos da Palavra Sagrada, sob oração e bênção pastoral."
    else:
        corpo = f"Pela presente recomendamos o(a) nosso(a) irmão(ã) em Cristo <b>{m['nome'].upper()}</b>, que congregou connosco em plena comunhão fraterna como <b>{m['posicao_atual']}</b>, prestando serviços no departamento <b>{m['departamento']}</b>. Rogamos que seja acolhido(a) no amor fraterno pelo vosso ministério."

    elementos.append(Paragraph(f"<font color='#1e293b' size=13 leading=22>{corpo}</font>", ParagraphStyle('B', alignment=4, spaceAfter=25)))
    elementos.append(Paragraph(f"<i><font color='#64748b' size=10>{versiculo}</font></i>", ParagraphStyle('V', alignment=1, spaceAfter=35)))
    elementos.append(Spacer(1, 1*cm))
    elementos.append(Paragraph(f"<font color='#334155' size=10><b>Chicuque, aos {datetime.now().strftime('%d de %B de %Y')}.</b></font>", ParagraphStyle('D', alignment=1, spaceAfter=30)))

    ass = [[
        Paragraph("_______________________________________<br/><b>Pastor Presidente da IEAD Chicuque</b>", ParagraphStyle('Ass1', alignment=1)),
        Paragraph("_______________________________________<br/><b>Secretaria Geral / Ministro Oficiante</b>", ParagraphStyle('Ass2', alignment=1))
    ]]
    elementos.append(Table(ass, colWidths=[12*cm, 12*cm]))

    doc.build(elementos)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"Certificado_{tipo}_{m['nome'].replace(' ', '_')}.pdf", mimetype="application/pdf")

@app.route('/financeiro/recibo_pdf/<int:id>')
def recibo_financeiro_pdf(id):
    if not can_tesouraria(): return redirect(url_for('dashboard'))
    conn = get_db()
    f = conn.execute("SELECT * FROM financeiro WHERE id = ?", (id,)).fetchone()
    conn.close()
    if not f: return redirect(url_for('dashboard'))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2*cm, rightMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm)
    elementos = [
        Paragraph("<font color='#1e3a8a' size=14><b>IGREJA EVANGÉLICA ASSEMBLEIA DE DEUS</b></font>", ParagraphStyle('H1', alignment=1)),
        Paragraph("<font color='#475569' size=9><b>CONGREGAÇÃO DE CHICUQUE • DEPARTAMENTO DE TESOURARIA</b></font>", ParagraphStyle('H2', alignment=1, spaceAfter=20)),
        Paragraph(f"<font color='#059669' size=16><b>COMPROVATIVO DE MOVIMENTO DE CAIXA Nº #{f['id']}</b></font>", ParagraphStyle('T', alignment=1, spaceAfter=20))
    ]

    dados = [
        ["Data do Movimento:", f['data_movimento']],
        ["Tipo de Operação:", f['tipo']],
        ["Conta / Caixa:", f['local_movimento']],
        ["Departamento / Fundo:", f['departamento'] or 'Geral'],
        ["Categoria / Área:", f['categoria']],
        ["Valor Registado:", f"{f['valor']:,.2f} MT"],
        ["Descrição / Finalidade:", f['descricao']],
        ["Data do Registo:", f['data_registo']]
    ]
    t = Table(dados, colWidths=[6*cm, 10*cm])
    t.setStyle(TableStyle([
        ('FONTNAME', (0,0), (-1,-1), 'Helvetica'),
        ('FONTSIZE', (0,0), (-1,-1), 10),
        ('TEXTCOLOR', (0,0), (0,-1), colors.HexColor("#1e3a8a")),
        ('FONTNAME', (0,0), (0,-1), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('LINEBELOW', (0,0), (-1,-1), 0.5, colors.HexColor("#e2e8f0")),
    ]))
    elementos.append(t)
    elementos.append(Spacer(1, 2*cm))

    ass = [[
        Paragraph("____________________________________<br/><b>Responsável da Tesouraria</b>", ParagraphStyle('A1', alignment=1)),
        Paragraph("____________________________________<br/><b>Visto Pastoral / Controlo</b>", ParagraphStyle('A2', alignment=1))
    ]]
    elementos.append(Table(ass, colWidths=[8*cm, 8*cm]))

    doc.build(elementos)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"Recibo_Caixa_{f['id']}.pdf", mimetype="application/pdf")

@app.route('/financeiro/balancete_pdf')
def balancete_financeiro_pdf():
    if not can_tesouraria(): return redirect(url_for('dashboard'))
    conn = get_db()
    entrada_caixa = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Entrada' AND (local_movimento = 'Caixa' OR local_movimento IS NULL)").fetchone()[0] or 0.0
    saida_caixa = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Saída' AND (local_movimento = 'Caixa' OR local_movimento IS NULL)").fetchone()[0] or 0.0
    saldo_caixa = entrada_caixa - saida_caixa

    entrada_banco = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Entrada' AND local_movimento = 'Banco'").fetchone()[0] or 0.0
    saida_banco = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Saída' AND local_movimento = 'Banco'").fetchone()[0] or 0.0
    saldo_banco = entrada_banco - saida_banco
    saldo_total = (entrada_caixa + entrada_banco) - (saida_caixa + saida_banco)

    movimentos = conn.execute("SELECT data_movimento, tipo, local_movimento, departamento, categoria, valor FROM financeiro ORDER BY id DESC LIMIT 50").fetchall()
    conn.close()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=1.5*cm, rightMargin=1.5*cm, topMargin=1.5*cm, bottomMargin=1.5*cm)
    elementos = [
        Paragraph("<font color='#1e3a8a' size=14><b>IGREJA EVANGÉLICA ASSEMBLEIA DE DEUS</b></font>", ParagraphStyle('H1', alignment=1)),
        Paragraph("<font color='#475569' size=9><b>CONGREGAÇÃO DE CHICUQUE • BALANCETE FINANCEIRO CONSOLIDADO</b></font>", ParagraphStyle('H2', alignment=1, spaceAfter=15))
    ]

    resumo = [
        ["Conta / Caixa", "Entradas Totais", "Saídas Totais", "Saldo Atual"],
        ["Caixa Físico", f"{entrada_caixa:,.2f} MT", f"{saida_caixa:,.2f} MT", f"{saldo_caixa:,.2f} MT"],
        ["Conta Bancária", f"{entrada_banco:,.2f} MT", f"{saida_banco:,.2f} MT", f"{saldo_banco:,.2f} MT"],
        ["TOTAL GERAL", f"{entrada_caixa+entrada_banco:,.2f} MT", f"{saida_caixa+saida_banco:,.2f} MT", f"{saldo_total:,.2f} MT"]
    ]
    t_res = Table(resumo, colWidths=[4.5*cm, 4.5*cm, 4.5*cm, 4.5*cm])
    t_res.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#1e3a8a")),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('FONTNAME', (0,-1), (-1,-1), 'Helvetica-Bold'),
        ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor("#f1f5f9")),
    ]))
    elementos.append(t_res)
    elementos.append(Spacer(1, 1*cm))

    elementos.append(Paragraph("<b>EXTRATO DE LANÇAMENTOS RECENTES</b>", ParagraphStyle('H3', spaceAfter=8)))
    linhas = [["Data", "Tipo", "Conta", "Departamento", "Categoria", "Valor (MT)"]]
    for m in movimentos:
        linhas.append([m['data_movimento'], m['tipo'], m['local_movimento'], m['departamento'], m['categoria'], f"{m['valor']:,.2f}"])
    
    t_mov = Table(linhas, colWidths=[2.5*cm, 2*cm, 2.5*cm, 3.5*cm, 4.5*cm, 3*cm])
    t_mov.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#334155")),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTSIZE', (0,0), (-1,-1), 8),
        ('ALIGN', (-1,0), (-1,-1), 'RIGHT'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#e2e8f0")),
    ]))
    elementos.append(t_mov)
    elementos.append(Spacer(1, 1.5*cm))

    ass = [[
        Paragraph("____________________________<br/><b>1º Tesoureiro</b>", ParagraphStyle('A1', alignment=1)),
        Paragraph("____________________________<br/><b>Conselho Fiscal</b>", ParagraphStyle('A2', alignment=1)),
        Paragraph("____________________________<br/><b>Pastor Presidente</b>", ParagraphStyle('A3', alignment=1))
    ]]
    elementos.append(Table(ass, colWidths=[6*cm, 6*cm, 6*cm]))

    doc.build(elementos)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"Balancete_IEAD_Chicuque_{datetime.now().strftime('%Y_%m')}.pdf", mimetype="application/pdf")

@app.route('/sistema/backup')
def backup_sistema():
    if not is_admin(): return redirect(url_for('dashboard'))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zipf:
        if os.path.exists(DB_NAME):
            zipf.write(DB_NAME, arcname=DB_NAME)
        for root, _, files in os.walk(UPLOAD_FOLDER):
            for file in files:
                caminho = os.path.join(root, file)
                rel_path = os.path.relpath(caminho, start=".")
                zipf.write(caminho, arcname=rel_path)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"BACKUP_COMPLETO_IEAD_CHICUQUE_{datetime.now().strftime('%Y%m%d_%H%M')}.zip", mimetype="application/zip")

# ROTAS RESTANTES (TESOURARIA, TRANSFERÊNCIAS, CONFIGURAÇÕES, UTILIZADORES)
@app.route('/financeiro/novo', methods=['POST'])
def novo_financeiro():
    if not can_tesouraria(): return redirect(url_for('dashboard'))
    data_raw = request.form.get('data_movimento') or datetime.now().strftime("%Y-%m-%d")
    dt_obj = datetime.strptime(data_raw, "%Y-%m-%d")
    conn = get_db()
    conn.execute('''INSERT INTO financeiro (tipo, local_movimento, departamento, categoria, valor, data_movimento, dia, mes, ano, data_registo, membro_id, descricao)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
                 (request.form['tipo'], request.form.get('local_movimento', 'Caixa'), request.form.get('departamento', 'Geral'),
                  request.form['categoria'], float(request.form['valor']), dt_obj.strftime("%d/%m/%Y"), dt_obj.day, dt_obj.month, dt_obj.year,
                  datetime.now().strftime("%d/%m/%Y %H:%M"), request.form.get('membro_id') or None, request.form['descricao']))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/financeiro/transferir', methods=['POST'])
def transferir_fundos():
    if not can_tesouraria(): return redirect(url_for('dashboard'))
    data_raw = request.form.get('data_movimento') or datetime.now().strftime("%Y-%m-%d")
    dt_obj = datetime.strptime(data_raw, "%Y-%m-%d")
    origem_local, origem_depto = request.form['origem_local'], request.form['origem_depto']
    destino_local, destino_depto = request.form['destino_local'], request.form['destino_depto']
    valor = float(request.form['valor'])
    motivo = request.form['motivo']
    agora = datetime.now().strftime("%d/%m/%Y %H:%M")

    conn = get_db()
    conn.execute('''INSERT INTO financeiro (tipo, local_movimento, departamento, categoria, valor, data_movimento, dia, mes, ano, data_registo, descricao)
                    VALUES ('Saída', ?, ?, 'Transferência de Fundos', ?, ?, ?, ?, ?, ?, ?)''',
                 (origem_local, origem_depto, valor, dt_obj.strftime("%d/%m/%Y"), dt_obj.day, dt_obj.month, dt_obj.year, agora, f"Transf. p/ [{destino_local} - {destino_depto}]: {motivo}"))
    conn.execute('''INSERT INTO financeiro (tipo, local_movimento, departamento, categoria, valor, data_movimento, dia, mes, ano, data_registo, descricao)
                    VALUES ('Entrada', ?, ?, 'Transferência de Fundos', ?, ?, ?, ?, ?, ?, ?)''',
                 (destino_local, destino_depto, valor, dt_obj.strftime("%d/%m/%Y"), dt_obj.day, dt_obj.month, dt_obj.year, agora, f"Transf. de [{origem_local} - {origem_depto}]: {motivo}"))
    conn.execute('''INSERT INTO transferencias (data_movimento, origem_local, origem_depto, destino_local, destino_depto, valor, motivo, data_registo)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)''',
                 (dt_obj.strftime("%d/%m/%Y"), origem_local, origem_depto, destino_local, destino_depto, valor, motivo, agora))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/casamentos/novo', methods=['POST'])
def novo_casamento():
    if not can_cadastro(): return redirect(url_for('dashboard'))
    conn = get_db()
    conn.execute('''INSERT INTO casamentos (noivo, noiva, data_casamento, pastor_oficiante, data_registo)
                    VALUES (?, ?, ?, ?, ?)''',
                 (request.form['noivo'], request.form['noiva'], request.form['data_casamento'],
                  request.form.get('pastor_oficiante', ''), datetime.now().strftime("%d/%m/%Y")))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/mortes/novo', methods=['POST'])
def novo_morte():
    if not can_cadastro(): return redirect(url_for('dashboard'))
    conn = get_db()
    conn.execute('''INSERT INTO mortes (nome_falecido, data_falecimento, observacoes, data_registo)
                    VALUES (?, ?, ?, ?)''',
                 (request.form['nome_falecido'], request.form['data_falecimento'],
                  request.form.get('observacoes', ''), datetime.now().strftime("%d/%m/%Y")))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/usuarios/novo', methods=['POST'])
def novo_usuario():
    if not is_admin(): return redirect(url_for('dashboard'))
    user, senha, cargo = request.form['usuario'].strip(), request.form['senha'].strip(), request.form['cargo'].strip()
    if user and senha:
        conn = get_db()
        try:
            conn.execute("INSERT INTO usuarios (usuario, senha, cargo) VALUES (?, ?, ?)", (user, senha, cargo))
            conn.commit()
            session['sucesso_cadastro'] = f"Utilizador '{user}' criado com sucesso!"
        except sqlite3.IntegrityError:
            session['alerta_duplicado'] = f"Utilizador '{user}' já existe."
        conn.close()
    return redirect(url_for('dashboard'))

@app.route('/usuarios/apagar/<int:id>')
def apagar_usuario(id):
    if not is_admin(): return redirect(url_for('dashboard'))
    conn = get_db()
    u = conn.execute("SELECT usuario FROM usuarios WHERE id = ?", (id,)).fetchone()
    if u and u['usuario'] != 'admin':
        conn.execute("DELETE FROM usuarios WHERE id = ?", (id,))
        conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/apagar/<tabela>/<int:id>')
def apagar_registo(tabela, id):
    if not is_admin(): return redirect(url_for('dashboard'))
    mapa = {'membro': ('membros', 'id'), 'financeiro': ('financeiro', 'id'), 'casamento': ('casamentos', 'id'), 'morte': ('mortes', 'id'), 'patrimonio': ('patrimonio', 'id'), 'escala': ('escalas', 'id'), 'culto': ('cultos_frequencia', 'id'), 'convertido': ('novos_convertidos', 'id')}
    if tabela in mapa:
        tab, col = mapa[tabela]
        conn = get_db()
        conn.execute(f"DELETE FROM {tab} WHERE {col} = ?", (id,))
        conn.commit()
        conn.close()
    return redirect(url_for('dashboard'))

@app.route('/config/departamento/novo', methods=['POST'])
def novo_depto():
    if not is_admin(): return redirect(url_for('dashboard'))
    nome = request.form.get('nome', '').strip()
    if nome:
        conn = get_db()
        try:
            conn.execute("INSERT INTO departamentos_lista (nome) VALUES (?)", (nome,))
            conn.commit()
        except sqlite3.IntegrityError: pass
        conn.close()
    return redirect(url_for('dashboard'))

@app.route('/config/departamento/apagar/<int:id>')
def apagar_depto(id):
    if not is_admin(): return redirect(url_for('dashboard'))
    conn = get_db()
    conn.execute("DELETE FROM departamentos_lista WHERE id = ?", (id,))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/config/categoria/novo', methods=['POST'])
def nova_categoria():
    if not is_admin(): return redirect(url_for('dashboard'))
    tipo, nome = request.form.get('tipo'), request.form.get('nome', '').strip()
    if tipo and nome:
        conn = get_db()
        conn.execute("INSERT INTO categorias_financeiras (tipo, nome) VALUES (?, ?)", (tipo, nome))
        conn.commit()
        conn.close()
    return redirect(url_for('dashboard'))

@app.route('/config/categoria/apagar/<int:id>')
def apagar_categoria(id):
    if not is_admin(): return redirect(url_for('dashboard'))
    conn = get_db()
    conn.execute("DELETE FROM categorias_financeiras WHERE id = ?", (id,))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/config/zona/novo', methods=['POST'])
def nova_zona():
    if not is_admin(): return redirect(url_for('dashboard'))
    nome = request.form.get('nome', '').strip()
    if nome:
        conn = get_db()
        try:
            conn.execute("INSERT INTO zonas_lista (nome) VALUES (?)", (nome,))
            conn.commit()
        except sqlite3.IntegrityError: pass
        conn.close()
    return redirect(url_for('dashboard'))

@app.route('/config/zona/apagar/<int:id>')
def apagar_zona(id):
    if not is_admin(): return redirect(url_for('dashboard'))
    conn = get_db()
    conn.execute("DELETE FROM zonas_lista WHERE id = ?", (id,))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/exportar/financeiro')
def exportar_financeiro():
    if not can_tesouraria(): return redirect(url_for('dashboard'))
    conn = get_db()
    rows = conn.execute("SELECT data_movimento, tipo, local_movimento, departamento, categoria, descricao, valor FROM financeiro ORDER BY id DESC").fetchall()
    conn.close()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tesouraria Chicuque"
    ws.append(["Data Movimento", "Tipo", "Conta", "Departamento", "Categoria", "Descrição", "Valor (MT)"])
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    for col in range(1, 8):
        c = ws.cell(row=1, column=col)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append([r['data_movimento'], r['tipo'], r['local_movimento'], r['departamento'], r['categoria'], r['descricao'], r['valor']])
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"Tesouraria_IEAD_Chicuque_{datetime.now().strftime('%Y%m%d')}.xlsx")

@app.route('/exportar/membros')
def exportar_membros():
    if not can_cadastro(): return redirect(url_for('dashboard'))
    conn = get_db()
    rows = conn.execute('''
        SELECT nome, telefone, genero, data_nascimento, faixa_etaria, naturalidade, bairro, 
               filiacao, tipo_documento, numero_documento, ano_conversao, data_batismo, 
               posicao_atual, progressoes, departamento, observacoes
        FROM membros ORDER BY nome ASC
    ''').fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Membresia IEAD Chicuque"
    headers = [
        "Nome Completo", "Contacto", "Género", "Data Nasc.", "Segmento", "Naturalidade", "Bairro",
        "Filiação", "Tipo Doc", "Nº Documento", "Ano Conv.", "Data Batismo",
        "Posição Atual", "Progressões", "Departamento", "Observações"
    ]
    ws.append(headers)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append([r[k] for k in r.keys()])
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 13)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"Membros_Completos_IEAD_{datetime.now().strftime('%Y%m%d')}.xlsx")

if __name__ == '__main__':
    port = int(os.environ.get("PORT", 5000))
    app.run(host='0.0.0.0', port=port)
"""

with open("app.py", "w", encoding="utf-8") as f:
    f.write(app_code)
print("✓ 1/2: app.py atualizado com os novos módulos (Cultos, Convertidos, Património, Escalas e Metas)!")

# =========================================================================
# 2. TEMPLATE DASHBOARD COM TODAS AS NOVAS ABAS E CARDS DE ANIVERSÁRIO
# =========================================================================
html_code = """<!DOCTYPE html>
<html lang="pt">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0, maximum-scale=1.0, user-scalable=no">
    <title>IEAD Chicuque - Gestão Eclesiástica Integrada</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <script src="https://cdn.jsdelivr.net/npm/chart.js"></script>
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
    <link href="https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700;800&display=swap" rel="stylesheet">
    <style>
        body { font-family: 'Plus Jakarta Sans', sans-serif; }
        .tab-content { display: none; }
        .tab-content.active { display: block; }
        .no-scrollbar::-webkit-scrollbar { display: none; }
        .no-scrollbar { -ms-overflow-style: none; scrollbar-width: none; }
        .card-glow { box-shadow: 0 10px 25px -5px rgba(30, 58, 138, 0.08), 0 8px 10px -6px rgba(30, 58, 138, 0.04); }
    </style>
</head>
<body class="bg-gradient-to-br from-slate-100 via-blue-50/50 to-indigo-50/40 text-slate-900 min-h-screen antialiased selection:bg-blue-600 selection:text-white">

    <!-- Topbar -->
    <header class="bg-gradient-to-r from-blue-950 via-indigo-950 to-slate-950 text-white shadow-xl sticky top-0 z-40 border-b border-indigo-500/20">
        <div class="max-w-7xl mx-auto px-4 py-3.5 flex justify-between items-center">
            <div class="flex items-center space-x-3">
                <div class="w-12 h-12 rounded-2xl bg-gradient-to-tr from-white to-blue-50 p-1.5 shadow-lg flex items-center justify-center flex-shrink-0 ring-2 ring-amber-400/40">
                    <img src="/static/logo.svg?v=2026" alt="IEAD" class="w-full h-full object-contain">
                </div>
                <div>
                    <h1 class="text-sm sm:text-base md:text-lg font-extrabold tracking-tight leading-tight uppercase text-white flex items-center gap-1.5">
                        Assembleia de Deus
                        <span class="inline-block w-2 h-2 rounded-full bg-emerald-400 animate-pulse"></span>
                    </h1>
                    <p class="text-xs text-indigo-200 font-semibold tracking-wide">Congregação de Chicuque • Gestão Integrada</p>
                </div>
            </div>

            <div class="flex items-center space-x-2.5">
                <div class="hidden sm:flex flex-col text-right pr-1">
                    <span class="text-xs font-black text-white leading-none tracking-wide">{{ session['usuario'] }}</span>
                    <span class="text-[11px] text-amber-300 font-bold mt-0.5">{{ session['cargo'] }}</span>
                </div>
                <a href="/logout" class="bg-gradient-to-r from-rose-500 to-red-600 hover:from-rose-600 hover:to-red-700 active:scale-95 text-white text-xs font-extrabold px-3.5 py-2 rounded-xl shadow-md transition">
                    Sair
                </a>
            </div>
        </div>

        <!-- Menu Completo de Abas -->
        <div class="max-w-7xl mx-auto px-3 overflow-x-auto no-scrollbar border-t border-white/10 py-2.5 bg-black/10">
            <nav class="flex space-x-2 text-sm font-semibold whitespace-nowrap">
                {% if pode_cadastro %}
                <button onclick="trocarAba('membros')" id="btn-membros" class="tab-btn px-4 py-2 rounded-2xl transition flex items-center gap-1.5 bg-gradient-to-r from-blue-600 to-indigo-600 text-white font-black shadow-md">
                    <span>👤</span> <span>Membros ({{ total_membros }})</span>
                </button>
                <button onclick="trocarAba('cultos')" id="btn-cultos" class="tab-btn px-4 py-2 rounded-2xl transition flex items-center gap-1.5 text-indigo-100 hover:bg-white/10">
                    <span>⛪</span> <span>Cultos & Presenças</span>
                </button>
                <button onclick="trocarAba('convertidos')" id="btn-convertidos" class="tab-btn px-4 py-2 rounded-2xl transition flex items-center gap-1.5 text-indigo-100 hover:bg-white/10">
                    <span>🌱</span> <span>Discipulado</span>
                </button>
                <button onclick="trocarAba('escalas')" id="btn-escalas" class="tab-btn px-4 py-2 rounded-2xl transition flex items-center gap-1.5 text-indigo-100 hover:bg-white/10">
                    <span>📋</span> <span>Escala Púlpito</span>
                </button>
                <button onclick="trocarAba('casamentos')" id="btn-casamentos" class="tab-btn px-4 py-2 rounded-2xl transition flex items-center gap-1.5 text-indigo-100 hover:bg-white/10">
                    <span>💍</span> <span>Casamentos</span>
                </button>
                <button onclick="trocarAba('mortes')" id="btn-mortes" class="tab-btn px-4 py-2 rounded-2xl transition flex items-center gap-1.5 text-indigo-100 hover:bg-white/10">
                    <span>🕊️</span> <span>Óbitos</span>
                </button>
                {% endif %}

                {% if pode_tesouraria %}
                <button onclick="trocarAba('dashboard')" id="btn-dashboard" class="tab-btn px-4 py-2 rounded-2xl transition flex items-center gap-1.5 {% if not pode_cadastro %}bg-gradient-to-r from-blue-600 to-indigo-600 text-white font-black shadow-md{% else %}text-indigo-100 hover:bg-white/10{% endif %}">
                    <span>📊</span> <span>Painel Financeiro</span>
                </button>
                <button onclick="trocarAba('financeiro')" id="btn-financeiro" class="tab-btn px-4 py-2 rounded-2xl transition flex items-center gap-1.5 text-indigo-100 hover:bg-white/10">
                    <span>💰</span> <span>Tesouraria</span>
                </button>
                <button onclick="trocarAba('transferencias')" id="btn-transferencias" class="tab-btn px-4 py-2 rounded-2xl transition flex items-center gap-1.5 text-cyan-300 hover:bg-white/10">
                    <span>🔄</span> <span>Transferir</span>
                </button>
                {% endif %}

                {% if e_admin %}
                <button onclick="trocarAba('patrimonio')" id="btn-patrimonio" class="tab-btn px-4 py-2 rounded-2xl transition flex items-center gap-1.5 text-amber-200 hover:bg-white/10">
                    <span>📦</span> <span>Património</span>
                </button>
                <button onclick="trocarAba('configuracoes')" id="btn-configuracoes" class="tab-btn px-4 py-2 rounded-2xl transition flex items-center gap-1.5 text-amber-300 hover:bg-white/10">
                    <span>⚙️</span> <span>Configurações</span>
                </button>
                <button onclick="trocarAba('usuarios')" id="btn-usuarios" class="tab-btn px-4 py-2 rounded-2xl transition flex items-center gap-1.5 text-emerald-300 hover:bg-white/10">
                    <span>👥</span> <span>Utilizadores</span>
                </button>
                <a href="/financeiro/balancete_pdf" class="tab-btn px-4 py-2 rounded-2xl transition flex items-center gap-1.5 bg-emerald-600 text-white font-extrabold shadow-sm">
                    <span>📑</span> <span>Balancete (PDF)</span>
                </a>
                <a href="/sistema/backup" class="tab-btn px-4 py-2 rounded-2xl transition flex items-center gap-1.5 bg-amber-600 text-white font-extrabold shadow-sm">
                    <span>💾</span> <span>Backup (.ZIP)</span>
                </a>
                {% endif %}
            </nav>
        </div>
    </header>

    <main class="max-w-7xl mx-auto p-3.5 sm:p-5 lg:p-6 space-y-6">

        {% if alerta_duplicado %}
        <div class="p-4 bg-amber-50 border-2 border-amber-300 rounded-2xl flex items-center gap-3 text-amber-900 font-bold text-sm shadow-md">
            <span class="text-2xl">⚠️</span><span>{{ alerta_duplicado }}</span>
        </div>
        {% endif %}

        {% if sucesso_cadastro %}
        <div class="p-4 bg-emerald-50 border-2 border-emerald-300 rounded-2xl flex items-center gap-3 text-emerald-900 font-bold text-sm shadow-md">
            <span class="text-2xl">✅</span><span>{{ sucesso_cadastro }}</span>
        </div>
        {% endif %}

        <!-- BANNER DE ANIVERSARIANTES DO DIA -->
        {% if aniversariantes_dia %}
        <div class="bg-gradient-to-r from-amber-500 via-orange-500 to-amber-600 text-white p-4 sm:p-5 rounded-3xl shadow-lg flex flex-col sm:flex-row justify-between items-start sm:items-center gap-3">
            <div>
                <span class="bg-white/20 text-white text-[10px] font-black uppercase px-2.5 py-0.5 rounded-full">Felicitações Pastorais</span>
                <h3 class="text-base sm:text-lg font-black mt-1">🎂 Aniversariante(s) de Hoje na IEAD Chicuque!</h3>
                <p class="text-xs text-amber-100 font-medium">Lembre-se de parabenizar no culto ou enviar mensagem pastoral.</p>
            </div>
            <div class="flex flex-wrap gap-2">
                {% for a in aniversariantes_dia %}
                <a href="https://wa.me/258{{ a['telefone']|replace(' ', '')|replace('+', '')|replace('-', '') }}?text=Paz%20do%20Senhor%20Irm%C3%A3o(a)%20{{ a['nome'] }},%20a%20IEAD%20Chicuque%20deseja-lhe%20um%20feliz%20anivers%C3%A1rio%20e%20muitas%20b%C3%AAn%C3%A7%C3%A3os%20de%20Deus!" 
                   target="_blank" class="bg-white text-slate-900 hover:bg-amber-50 text-xs font-black px-3.5 py-2 rounded-xl shadow transition flex items-center gap-1.5">
                    <span>💬</span> <span>Felicitar {{ a['nome'].split()[0] }}</span>
                </a>
                {% endfor %}
            </div>
        </div>
        {% endif %}

        {% if pode_cadastro %}
        <!-- ================= ABA 1: MEMBROS (COM OS CARDS E BOTÕES VISÍVEIS) ================= -->
        <section id="aba-membros" class="tab-content {% if pode_cadastro %}active{% endif %} space-y-6">
            <div class="grid grid-cols-1 lg:grid-cols-12 gap-6">
                <!-- Formulário de Cadastro -->
                <div class="lg:col-span-5 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5 sm:p-6">
                    <div class="flex items-center justify-between pb-4 mb-5 border-b border-indigo-50">
                        <div>
                            <h2 class="text-lg font-black text-slate-900 tracking-tight flex items-center gap-2">
                                <span class="text-blue-600">📝</span> Ficha de Membro
                            </h2>
                            <p class="text-xs text-slate-500 font-medium">Registo com fotografia e documentos</p>
                        </div>
                        <span class="bg-blue-600 text-white text-[11px] font-black px-3 py-1 rounded-full shadow-sm">Oficial IEAD</span>
                    </div>

                    <form action="/membros/novo" method="POST" enctype="multipart/form-data" class="space-y-4">
                        <div class="flex flex-col items-center justify-center p-4 bg-blue-50/50 border-2 border-dashed border-indigo-200 rounded-3xl">
                            <label class="cursor-pointer bg-blue-600 hover:bg-blue-700 text-white text-xs font-bold px-4 py-2 rounded-xl shadow transition flex items-center gap-1.5">
                                <span>Tirar Foto / Carregar</span>
                                <input type="file" name="foto" accept="image/*" capture="environment" class="hidden">
                            </label>
                            <p class="text-[11px] text-slate-400 font-semibold mt-1">Fotografia do membro</p>
                        </div>

                        <div>
                            <label class="block text-xs font-bold text-slate-700 mb-1">Nome Completo *</label>
                            <input type="text" name="nome" required placeholder="Nome oficial..." class="w-full h-11 px-3 text-sm border-2 rounded-xl outline-none focus:border-blue-600">
                        </div>

                        <div class="grid grid-cols-2 gap-2">
                            <div>
                                <label class="block text-xs font-bold text-slate-700 mb-1">Contacto Telefónico</label>
                                <input type="text" name="telefone" placeholder="+258 8..." class="w-full h-11 px-3 text-sm border-2 rounded-xl outline-none">
                            </div>
                            <div>
                                <label class="block text-xs font-bold text-slate-700 mb-1">Género</label>
                                <select name="genero" class="w-full h-11 px-2 text-sm border-2 rounded-xl bg-white"><option value="Masculino">Masculino</option><option value="Feminino">Feminino</option></select>
                            </div>
                        </div>

                        <div class="grid grid-cols-2 gap-2">
                            <div>
                                <label class="block text-xs font-bold text-slate-700 mb-1">Data Nascimento</label>
                                <input type="date" name="data_nascimento" class="w-full h-11 px-2 text-sm border-2 rounded-xl bg-white">
                            </div>
                            <div>
                                <label class="block text-xs font-bold text-slate-700 mb-1">Segmento *</label>
                                <select name="faixa_etaria" required class="w-full h-11 px-2 text-sm border-2 rounded-xl bg-white font-bold text-blue-900">
                                    <option value="Adulto">Adulto</option>
                                    <option value="Jovem">Jovem</option>
                                    <option value="Adolescente">Adolescente</option>
                                    <option value="Criança">Criança (Boa Esperança)</option>
                                    <option value="Terceira Idade">Terceira Idade</option>
                                </select>
                            </div>
                        </div>

                        <div class="grid grid-cols-2 gap-2">
                            <div>
                                <label class="block text-xs font-bold text-slate-700 mb-1">Bairro / Zona</label>
                                <select name="bairro" class="w-full h-11 px-2 text-sm border-2 rounded-xl bg-white">
                                    <option value="">-- Selecione --</option>
                                    {% for z in lista_zonas %}<option value="{{ z['nome'] }}">{{ z['nome'] }}</option>{% endfor %}
                                </select>
                            </div>
                            <div>
                                <label class="block text-xs font-bold text-slate-700 mb-1">Naturalidade</label>
                                <input type="text" name="naturalidade" placeholder="Cidade / Província" class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                            </div>
                        </div>

                        <div>
                            <label class="block text-xs font-bold text-slate-700 mb-1">Filiação (Pai & Mãe)</label>
                            <input type="text" name="filiacao" placeholder="Nome dos pais..." class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                        </div>

                        <div class="grid grid-cols-3 gap-2">
                            <div>
                                <label class="block text-xs font-bold text-slate-700 mb-1">Tipo Doc.</label>
                                <select name="tipo_documento" class="w-full h-11 px-2 text-sm border-2 rounded-xl bg-white"><option value="BI">BI</option><option value="Cédula">Cédula</option><option value="Cartão de Eleitor">C. Eleitor</option></select>
                            </div>
                            <div class="col-span-2">
                                <label class="block text-xs font-bold text-slate-700 mb-1">Nº Documento</label>
                                <input type="text" name="numero_documento" placeholder="Número..." class="w-full h-11 px-3 text-sm border-2 rounded-xl font-mono">
                            </div>
                        </div>

                        <div class="grid grid-cols-2 gap-2">
                            <div>
                                <label class="block text-xs font-bold text-slate-700 mb-1">Ano Conversão</label>
                                <input type="number" name="ano_conversao" placeholder="Ex: 2018" class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                            </div>
                            <div>
                                <label class="block text-xs font-bold text-slate-700 mb-1">Data Batismo</label>
                                <input type="date" name="data_batismo" class="w-full h-11 px-2 text-sm border-2 rounded-xl bg-white">
                            </div>
                        </div>

                        <div class="grid grid-cols-2 gap-2">
                            <div>
                                <label class="block text-xs font-bold text-blue-950 mb-1">Cargo *</label>
                                <select name="posicao_atual" required class="w-full h-11 px-2 text-sm border-2 rounded-xl bg-white font-bold text-blue-900">
                                    <option value="Membro em Comunhão">Membro em Comunhão</option>
                                    <option value="Pastor">Pastor</option>
                                    <option value="Presbítero">Presbítero</option>
                                    <option value="Diácono / Diaconisa">Diácono / Diaconisa</option>
                                    <option value="Evangelista">Evangelista</option>
                                    <option value="Obreiro(a)">Obreiro(a)</option>
                                    <option value="Líder de Célula">Líder de Célula</option>
                                    <option value="Membro em Prova">Membro em Prova</option>
                                </select>
                            </div>
                            <div>
                                <label class="block text-xs font-bold text-indigo-950 mb-1">Departamento *</label>
                                <select name="departamento" required class="w-full h-11 px-2 text-sm border-2 rounded-xl bg-white font-bold text-indigo-900">
                                    <option value="Activista">Activista</option>
                                    <option value="Juventude">Juventude</option>
                                    <option value="Mulher (Senhoras)">Mulher (Senhoras)</option>
                                    <option value="Boa Esperança (Crianças)">Boa Esperança (Crianças)</option>
                                    <option value="Homens / Obreiros">Homens / Obreiros</option>
                                    <option value="Louvor / Música">Louvor / Música</option>
                                    <option value="Ação Social">Ação Social</option>
                                    <option value="Geral">Geral</option>
                                </select>
                            </div>
                        </div>

                        <button type="submit" class="w-full h-12 bg-gradient-to-r from-emerald-600 to-teal-700 hover:from-emerald-700 hover:to-teal-800 text-white font-black text-sm rounded-2xl shadow-lg transition">
                            Gravar Ficha de Membro
                        </button>
                    </form>
                </div>

                <!-- CARDS INDIVIDUAIS DOS MEMBROS COM TODOS OS BOTÕES VISÍVEIS -->
                <div class="lg:col-span-7 space-y-4">
                    <div class="bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5 flex justify-between items-center">
                        <div>
                            <h3 class="text-base font-black text-slate-900">Membros Registados ({{ total_membros }})</h3>
                            <p class="text-xs text-slate-500">Cartões em PDF, WhatsApp e Certificados</p>
                        </div>
                        <a href="/exportar/membros" class="h-10 px-4 bg-emerald-600 hover:bg-emerald-700 text-white rounded-xl text-xs font-black shadow transition flex items-center gap-1.5">
                            <span>📥 Baixar Excel</span>
                        </a>
                    </div>

                    <div class="space-y-4 max-h-[750px] overflow-y-auto pr-1">
                        {% for m in todos_membros %}
                        <div class="bg-white/95 rounded-3xl p-5 border border-indigo-100 shadow-md space-y-3">
                            <div class="flex items-start justify-between pb-3 border-b">
                                <div class="flex items-center space-x-3">
                                    {% if m['foto_path'] %}
                                    <img src="{{ m['foto_path'] }}" class="w-12 h-12 object-cover rounded-2xl border-2 border-blue-600 shadow">
                                    {% else %}
                                    <div class="w-12 h-12 rounded-2xl bg-blue-700 text-white flex items-center justify-center font-black text-base shadow">
                                        {{ m['nome'][:2].upper() }}
                                    </div>
                                    {% endif %}
                                    <div>
                                        <h4 class="font-extrabold text-sm sm:text-base text-slate-900">{{ m['nome'] }}</h4>
                                        <div class="flex items-center gap-1 text-xs mt-0.5">
                                            <span class="bg-blue-50 text-blue-900 font-bold px-2 py-0.5 rounded">{{ m['posicao_atual'] }}</span>
                                            <span class="bg-slate-100 text-slate-700 font-medium px-2 py-0.5 rounded">{{ m['departamento'] }}</span>
                                        </div>
                                    </div>
                                </div>
                                {% if e_admin %}
                                <a href="/apagar/membro/{{ m['id'] }}" onclick="return confirm('Eliminar?');" class="text-rose-600 font-bold text-xs p-1">✕</a>
                                {% endif %}
                            </div>

                            <div class="grid grid-cols-2 gap-2 text-xs text-slate-600 bg-slate-50 p-2.5 rounded-xl border">
                                <div><strong>Doc:</strong> {{ m['tipo_documento'] or 'BI' }}: {{ m['numero_documento'] or '-' }}</div>
                                <div><strong>Contacto:</strong> {{ m['telefone'] or 'Sem telefone' }}</div>
                            </div>

                            <!-- BOTÕES VISÍVEIS DIRETAMENTE -->
                            <div class="grid grid-cols-1 sm:grid-cols-2 gap-2">
                                <a href="/membro/cartao_pdf/{{ m['id'] }}" class="h-10 bg-blue-700 hover:bg-blue-800 text-white font-bold text-xs rounded-xl shadow flex items-center justify-center gap-1.5 transition">
                                    <span>🪪</span> <span>Baixar Cartão (PDF)</span>
                                </a>
                                {% if m['telefone'] %}
                                <a href="https://wa.me/258{{ m['telefone']|replace(' ', '')|replace('+', '')|replace('-', '') }}" target="_blank" class="h-10 bg-emerald-600 hover:bg-emerald-700 text-white font-bold text-xs rounded-xl shadow flex items-center justify-center gap-1.5 transition">
                                    <span>💬</span> <span>WhatsApp</span>
                                </a>
                                {% else %}
                                <span class="h-10 bg-slate-100 text-slate-400 font-bold text-xs rounded-xl flex items-center justify-center">Sem Contacto</span>
                                {% endif %}
                            </div>

                            <div class="p-2 bg-indigo-50/60 rounded-xl border border-indigo-100">
                                <span class="text-[10px] font-black text-indigo-950 uppercase tracking-wider block mb-1">Emitir Certificados Pastorais (PDF):</span>
                                <div class="grid grid-cols-3 gap-1">
                                    <a href="/membro/certificado_pdf/{{ m['id'] }}/batismo" class="py-1.5 bg-white border border-blue-200 text-blue-900 text-center font-bold text-xs rounded-lg hover:bg-blue-50">Batismo</a>
                                    <a href="/membro/certificado_pdf/{{ m['id'] }}/apresentacao" class="py-1.5 bg-white border border-emerald-200 text-emerald-900 text-center font-bold text-xs rounded-lg hover:bg-emerald-50">Apresentação</a>
                                    <a href="/membro/certificado_pdf/{{ m['id'] }}/recomendacao" class="py-1.5 bg-white border border-purple-200 text-purple-900 text-center font-bold text-xs rounded-lg hover:bg-purple-50">Recomendação</a>
                                </div>
                            </div>
                        </div>
                        {% endfor %}
                    </div>
                </div>
            </div>
        </section>

        <!-- ================= ABA 2: CULTOS & PRESENÇAS ================= -->
        <section id="aba-cultos" class="tab-content space-y-6">
            <div class="grid grid-cols-1 lg:grid-cols-12 gap-6">
                <div class="lg:col-span-5 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5 sm:p-6">
                    <h3 class="text-base font-black text-slate-900 mb-4 pb-2 border-b">⛪ Registo Rápido de Culto</h3>
                    <form action="/cultos/novo" method="POST" class="space-y-3">
                        <div class="grid grid-cols-2 gap-2">
                            <div><label class="block text-xs font-bold text-slate-700 mb-1">Data *</label><input type="date" name="data_culto" required class="w-full h-11 px-2 text-sm border-2 rounded-xl bg-white"></div>
                            <div>
                                <label class="block text-xs font-bold text-slate-700 mb-1">Culto *</label>
                                <select name="tipo_culto" class="w-full h-11 px-2 text-sm border-2 rounded-xl bg-white font-bold">
                                    <option value="Domingo Manhã">Domingo Manhã</option>
                                    <option value="Domingo Noite">Domingo Noite</option>
                                    <option value="Quarta-feira">Quarta-feira</option>
                                    <option value="Sexta-feira (Vigília)">Sexta-feira (Vigília)</option>
                                    <option value="Culto de Obreiros">Culto de Obreiros</option>
                                </select>
                            </div>
                        </div>
                        <div class="grid grid-cols-3 gap-2">
                            <div><label class="block text-xs font-bold text-slate-700 mb-1">Homens</label><input type="number" name="homens" value="0" class="w-full h-10 px-2 text-sm border-2 rounded-xl"></div>
                            <div><label class="block text-xs font-bold text-slate-700 mb-1">Mulheres</label><input type="number" name="mulheres" value="0" class="w-full h-10 px-2 text-sm border-2 rounded-xl"></div>
                            <div><label class="block text-xs font-bold text-slate-700 mb-1">Jovens</label><input type="number" name="jovens" value="0" class="w-full h-10 px-2 text-sm border-2 rounded-xl"></div>
                        </div>
                        <div class="grid grid-cols-3 gap-2">
                            <div><label class="block text-xs font-bold text-slate-700 mb-1">Crianças</label><input type="number" name="criancas" value="0" class="w-full h-10 px-2 text-sm border-2 rounded-xl"></div>
                            <div><label class="block text-xs font-bold text-slate-700 mb-1">Visitantes</label><input type="number" name="visitantes" value="0" class="w-full h-10 px-2 text-sm border-2 rounded-xl"></div>
                            <div><label class="block text-xs font-bold text-emerald-800 mb-1">Convertidos</label><input type="number" name="novos_convertidos" value="0" class="w-full h-10 px-2 text-sm border-2 border-emerald-300 rounded-xl font-bold"></div>
                        </div>
                        <div><label class="block text-xs font-bold text-slate-700 mb-1">Pregador</label><input type="text" name="pregador" placeholder="Ministro da Palavra" class="w-full h-11 px-3 text-sm border-2 rounded-xl"></div>
                        <div><label class="block text-xs font-bold text-slate-700 mb-1">Tema da Mensagem</label><input type="text" name="tema_mensagem" placeholder="Título do sermão..." class="w-full h-11 px-3 text-sm border-2 rounded-xl"></div>
                        <button type="submit" class="w-full h-12 bg-blue-900 hover:bg-blue-950 text-white font-black text-sm rounded-2xl shadow transition">Gravar Presença no Culto</button>
                    </form>
                </div>
                <div class="lg:col-span-7 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5">
                    <h3 class="text-base font-black text-slate-900 mb-4 pb-2 border-b">Histórico de Cultos</h3>
                    <div class="overflow-x-auto">
                        <table class="w-full text-left text-xs">
                            <thead class="bg-slate-50 border-b"><tr><th class="p-2.5">Data/Culto</th><th class="p-2.5">Pregador</th><th class="p-2.5 text-center">Total Presentes</th><th class="p-2.5 text-center">Apelos</th></tr></thead>
                            <tbody class="divide-y">
                                {% for cu in todos_cultos %}
                                <tr>
                                    <td class="p-2.5 font-bold">{{ cu['data_culto'] }}<br/><span class="text-slate-400 font-normal">{{ cu['tipo_culto'] }}</span></td>
                                    <td class="p-2.5">{{ cu['pregador'] or '-' }}</td>
                                    <td class="p-2.5 text-center font-black text-blue-900 text-sm">{{ cu['total_presentes'] }}</td>
                                    <td class="p-2.5 text-center"><span class="bg-emerald-100 text-emerald-800 font-bold px-2 py-0.5 rounded-full">{{ cu['novos_convertidos'] }}</span></td>
                                </tr>
                                {% endfor %}
                            </tbody>
                        </table>
                    </div>
                </div>
            </div>
        </section>

        <!-- ================= ABA 3: DISCIPULADO & NOVOS CONVERTIDOS ================= -->
        <section id="aba-convertidos" class="tab-content space-y-6">
            <div class="grid grid-cols-1 lg:grid-cols-12 gap-6">
                <div class="lg:col-span-5 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5">
                    <h3 class="text-base font-black text-slate-900 mb-4 pb-2 border-b">🌱 Registar Novo Convertido (Apelo)</h3>
                    <form action="/convertidos/novo" method="POST" class="space-y-3">
                        <input type="text" name="nome" required placeholder="Nome Completo *" class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                        <input type="text" name="telefone" placeholder="Contacto Telefónico" class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                        <div class="grid grid-cols-2 gap-2">
                            <input type="text" name="bairro" placeholder="Bairro / Residência" class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                            <input type="date" name="data_decisao" required class="w-full h-11 px-2 text-sm border-2 rounded-xl bg-white font-semibold">
                        </div>
                        <input type="text" name="quem_convidou" placeholder="Quem convidou / Acompanhante" class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                        <button type="submit" class="w-full h-12 bg-emerald-700 hover:bg-emerald-800 text-white font-black text-sm rounded-2xl shadow transition">Registar para Discipulado</button>
                    </form>
                </div>
                <div class="lg:col-span-7 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5">
                    <h3 class="text-base font-black text-slate-900 mb-4 pb-2 border-b">Acompanhamento Espiritual</h3>
                    <div class="overflow-x-auto">
                        <table class="w-full text-left text-xs">
                            <thead class="bg-slate-50 border-b"><tr><th class="p-2.5">Nome</th><th class="p-2.5">Decisão</th><th class="p-2.5">Contacto</th><th class="p-2.5">Ciclo Espiritual</th></tr></thead>
                            <tbody class="divide-y">
                                {% for nc in todos_convertidos %}
                                <tr>
                                    <td class="p-2.5 font-bold">{{ nc['nome'] }}</td>
                                    <td class="p-2.5">{{ nc['data_decisao'] }}</td>
                                    <td class="p-2.5">{{ nc['telefone'] or '-' }}</td>
                                    <td class="p-2.5">
                                        <form action="/convertidos/atualizar_status/{{ nc['id'] }}" method="POST">
                                            <select name="status_discipulado" onchange="this.form.submit()" class="p-1 border rounded bg-white font-bold text-indigo-900">
                                                <option value="Decisão Inicial" {% if nc['status_discipulado'] == 'Decisão Inicial' %}selected{% endif %}>Decisão Inicial</option>
                                                <option value="Em Discipulado" {% if nc['status_discipulado'] == 'Em Discipulado' %}selected{% endif %}>Em Discipulado</option>
                                                <option value="Pronto p/ Batismo" {% if nc['status_discipulado'] == 'Pronto p/ Batismo' %}selected{% endif %}>Pronto p/ Batismo</option>
                                                <option value="Batizado" {% if nc['status_discipulado'] == 'Batizado' %}selected{% endif %}>Batizado</option>
                                            </select>
                                        </form>
                                    </td>
                                </tr>
                                {% endfor %}
                            </tbody>
                        </table>
                    </div>
                </div>
            </div>
        </section>

        <!-- ================= ABA 4: ESCALAS DE CULTO EM PDF ================= -->
        <section id="aba-escalas" class="tab-content space-y-6">
            <div class="grid grid-cols-1 lg:grid-cols-12 gap-6">
                <div class="lg:col-span-5 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5">
                    <h3 class="text-base font-black text-slate-900 mb-4 pb-2 border-b">📋 Criar Escala de Culto</h3>
                    <form action="/escalas/novo" method="POST" class="space-y-3">
                        <div class="grid grid-cols-2 gap-2">
                            <input type="date" name="data_escala" required class="w-full h-11 px-2 text-sm border-2 rounded-xl bg-white font-bold">
                            <input type="text" name="tipo_culto" required placeholder="Ex: Domingo Manhã" class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                        </div>
                        <input type="text" name="dirigente" placeholder="Dirigente do Culto" class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                        <input type="text" name="pregador" placeholder="Pregador (Palavra)" class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                        <input type="text" name="leitura_palavra" placeholder="Leitor da Bíblia" class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                        <input type="text" name="louvor_grupo" placeholder="Grupo de Louvor" class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                        <input type="text" name="diaconos_servico" placeholder="Diáconos de Porta" class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                        <button type="submit" class="w-full h-12 bg-indigo-900 hover:bg-indigo-950 text-white font-black text-sm rounded-2xl shadow transition">Gravar Escala</button>
                    </form>
                </div>
                <div class="lg:col-span-7 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5">
                    <h3 class="text-base font-black text-slate-900 mb-4 pb-2 border-b">Escalas Emitidas</h3>
                    <div class="overflow-x-auto">
                        <table class="w-full text-left text-xs">
                            <thead class="bg-slate-50 border-b"><tr><th class="p-2.5">Data/Culto</th><th class="p-2.5">Dirigente</th><th class="p-2.5">Pregador</th><th class="p-2.5 text-center">PDF</th></tr></thead>
                            <tbody class="divide-y">
                                {% for esc in todas_escalas %}
                                <tr>
                                    <td class="p-2.5 font-bold">{{ esc['data_escala'] }}<br/><span class="text-slate-400 font-normal">{{ esc['tipo_culto'] }}</span></td>
                                    <td class="p-2.5">{{ esc['dirigente'] or '-' }}</td>
                                    <td class="p-2.5 font-bold text-blue-900">{{ esc['pregador'] or '-' }}</td>
                                    <td class="p-2.5 text-center">
                                        <a href="/escalas/pdf/{{ esc['id'] }}" class="bg-blue-600 hover:bg-blue-700 text-white font-bold px-3 py-1.5 rounded-lg text-xs shadow inline-block">📄 Baixar PDF</a>
                                    </td>
                                </tr>
                                {% endfor %}
                            </tbody>
                        </table>
                    </div>
                </div>
            </div>
        </section>

        <!-- ================= ABA CASAMENTOS & ÓBITOS ================= -->
        <section id="aba-casamentos" class="tab-content space-y-6">
            <div class="grid grid-cols-1 lg:grid-cols-12 gap-6">
                <div class="lg:col-span-5 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5">
                    <h3 class="text-base font-black text-slate-900 mb-3 pb-2 border-b">💍 Registo Matrimonial</h3>
                    <form action="/casamentos/novo" method="POST" class="space-y-3">
                        <input type="text" name="noivo" required placeholder="Nome do Noivo *" class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                        <input type="text" name="noiva" required placeholder="Nome da Noiva *" class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                        <input type="date" name="data_casamento" required class="w-full h-11 px-2 text-sm border-2 rounded-xl bg-white font-bold">
                        <input type="text" name="pastor_oficiante" placeholder="Pastor Oficiante" class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                        <button type="submit" class="w-full h-12 bg-purple-800 text-white font-black text-sm rounded-2xl shadow">Gravar Matrimónio</button>
                    </form>
                </div>
                <div class="lg:col-span-7 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5">
                    <h3 class="text-base font-black text-slate-900 mb-3 pb-2 border-b">Livro de Casamentos</h3>
                    <div class="overflow-x-auto">
                        <table class="w-full text-left text-xs">
                            <thead class="bg-slate-50 border-b"><tr><th class="p-2.5">Data</th><th class="p-2.5">Casal</th><th class="p-2.5">Oficiante</th></tr></thead>
                            <tbody class="divide-y">
                                {% for c in todos_casamentos %}<tr><td class="p-2.5 font-bold">{{ c['data_casamento'] }}</td><td class="p-2.5 font-black text-purple-950">{{ c['noivo'] }} & {{ c['noiva'] }}</td><td class="p-2.5">{{ c['pastor_oficiante'] or '-' }}</td></tr>{% endfor %}
                            </tbody>
                        </table>
                    </div>
                </div>
            </div>
        </section>

        <section id="aba-mortes" class="tab-content space-y-6">
            <div class="grid grid-cols-1 lg:grid-cols-12 gap-6">
                <div class="lg:col-span-5 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5">
                    <h3 class="text-base font-black text-slate-900 mb-3 pb-2 border-b">🕊️ Registo de Óbito</h3>
                    <form action="/mortes/novo" method="POST" class="space-y-3">
                        <input type="text" name="nome_falecido" required placeholder="Nome do Falecido *" class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                        <input type="date" name="data_falecimento" required class="w-full h-11 px-2 text-sm border-2 rounded-xl bg-white font-bold">
                        <input type="text" name="observacoes" placeholder="Notas..." class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                        <button type="submit" class="w-full h-12 bg-slate-900 text-white font-black text-sm rounded-2xl shadow">Gravar Óbito</button>
                    </form>
                </div>
                <div class="lg:col-span-7 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5">
                    <h3 class="text-base font-black text-slate-900 mb-3 pb-2 border-b">Livro de Falecimentos</h3>
                    <div class="overflow-x-auto">
                        <table class="w-full text-left text-xs">
                            <thead class="bg-slate-50 border-b"><tr><th class="p-2.5">Data</th><th class="p-2.5">Falecido</th><th class="p-2.5">Observações</th></tr></thead>
                            <tbody class="divide-y">
                                {% for m in todas_mortes %}<tr><td class="p-2.5 font-bold">{{ m['data_falecimento'] }}</td><td class="p-2.5 font-black">{{ m['nome_falecido'] }}</td><td class="p-2.5">{{ m['observacoes'] or '-' }}</td></tr>{% endfor %}
                            </tbody>
                        </table>
                    </div>
                </div>
            </div>
        </section>
        {% endif %}

        {% if pode_tesouraria %}
        <!-- ================= ABA DASHBOARD COM TERMÓMETRO DE METAS ================= -->
        <section id="aba-dashboard" class="tab-content {% if not pode_cadastro and pode_tesouraria %}active{% endif %} space-y-6">
            
            <!-- Cards de Saldo -->
            <div class="grid grid-cols-1 sm:grid-cols-2 lg:grid-cols-4 gap-4">
                <div class="bg-gradient-to-br from-amber-500 to-orange-600 text-white p-5 rounded-3xl shadow-xl border border-amber-400">
                    <span class="text-xs font-black uppercase text-amber-100 block">💵 Caixa Físico</span>
                    <p class="text-2xl sm:text-3xl font-black text-white mt-2">{{ "{:,.2f}".format(saldo_caixa) }} MT</p>
                </div>
                <div class="bg-gradient-to-br from-blue-600 to-indigo-700 text-white p-5 rounded-3xl shadow-xl border border-blue-400">
                    <span class="text-xs font-black uppercase text-blue-100 block">🏛️ Conta Bancária</span>
                    <p class="text-2xl sm:text-3xl font-black text-white mt-2">{{ "{:,.2f}".format(saldo_banco) }} MT</p>
                </div>
                <div class="bg-gradient-to-br from-emerald-600 to-teal-800 text-white p-5 rounded-3xl shadow-xl border border-emerald-400">
                    <span class="text-xs font-black uppercase text-emerald-100 block">🏦 Saldo Consolidado</span>
                    <p class="text-2xl sm:text-3xl font-black text-white mt-2">{{ "{:,.2f}".format(saldo_total) }} MT</p>
                </div>
                <div class="bg-gradient-to-br from-purple-700 to-indigo-900 text-white p-5 rounded-3xl shadow-xl border border-purple-500">
                    <span class="text-xs font-black uppercase text-purple-200 block">👥 Membresia Registada</span>
                    <p class="text-2xl sm:text-3xl font-black text-white mt-2">{{ total_membros }} membros</p>
                </div>
            </div>

            <!-- TERMÓMETRO VISUAL DE METAS / CAMPANHAS DE CONSTRUÇÃO -->
            {% if campanhas %}
            <div class="bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5 space-y-3">
                <div class="flex justify-between items-center">
                    <h3 class="text-sm font-black uppercase tracking-wider text-slate-900 flex items-center gap-2">
                        <span class="text-xl">🏗️</span> <span>Campanhas de Obras & Metas Financeiras</span>
                    </h3>
                </div>
                {% for c in campanhas %}
                <div class="space-y-1.5 p-3.5 bg-slate-50 rounded-2xl border">
                    <div class="flex justify-between text-xs font-bold">
                        <span class="text-blue-950">{{ c['nome'] }}</span>
                        <span class="text-emerald-700 font-black">{{ "{:,.2f}".format(c['arrecadado']) }} MT de {{ "{:,.2f}".format(c['meta']) }} MT ({{ c['percentual'] }}%)</span>
                    </div>
                    <div class="w-full bg-slate-200 rounded-full h-3 overflow-hidden">
                        <div class="bg-gradient-to-r from-emerald-500 to-teal-600 h-3 rounded-full transition-all duration-500" style="width: {{ c['percentual'] }}%;"></div>
                    </div>
                </div>
                {% endfor %}
            </div>
            {% endif %}

            <!-- Gráficos de Tesouraria -->
            <div class="grid grid-cols-1 lg:grid-cols-2 gap-6">
                <div class="bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5">
                    <h3 class="text-sm font-black uppercase text-slate-800 mb-3">📈 Entradas vs Saídas Mensais</h3>
                    <div class="h-64"><canvas id="graficoEvolucaoMensal"></canvas></div>
                </div>
                <div class="bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5">
                    <h3 class="text-sm font-black uppercase text-slate-800 mb-3">🏛️ Situação por Departamento</h3>
                    <div class="h-64"><canvas id="graficoFinancasDepto"></canvas></div>
                </div>
            </div>

            <div class="grid grid-cols-1 lg:grid-cols-2 gap-6">
                <div class="bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5">
                    <h3 class="text-sm font-black uppercase text-slate-800 mb-3">🪙 Fontes de Fundos (Entradas)</h3>
                    <div class="h-64"><canvas id="graficoFontesFundos"></canvas></div>
                </div>
                <div class="bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5">
                    <h3 class="text-sm font-black uppercase text-slate-800 mb-3">💸 Destino das Saídas (Despesas)</h3>
                    <div class="h-64"><canvas id="graficoSaidasCategorias"></canvas></div>
                </div>
            </div>
        </section>

        <!-- ================= ABA TRANSFERÊNCIAS & TESOURARIA ================= -->
        <section id="aba-transferencias" class="tab-content space-y-6">
            <div class="grid grid-cols-1 lg:grid-cols-12 gap-6">
                <div class="lg:col-span-5 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5">
                    <h3 class="text-base font-black text-slate-900 mb-4 pb-2 border-b">🔄 Transferir Fundos</h3>
                    <form action="/financeiro/transferir" method="POST" class="space-y-3">
                        <input type="date" name="data_movimento" required class="w-full h-11 px-2 text-sm border-2 rounded-xl bg-white font-bold">
                        <div class="p-3 bg-rose-50 rounded-xl space-y-1.5 border">
                            <span class="text-[10px] font-bold text-rose-800 uppercase">Origem</span>
                            <select name="origem_local" class="w-full h-10 px-2 text-xs border rounded-lg bg-white"><option value="Caixa">Caixa Físico</option><option value="Banco">Conta Bancária</option></select>
                            <select name="origem_depto" class="w-full h-10 px-2 text-xs border rounded-lg bg-white"><option value="Geral">Fundo Geral</option>{% for d in lista_deptos %}<option value="{{ d['nome'] }}">{{ d['nome'] }}</option>{% endfor %}</select>
                        </div>
                        <div class="p-3 bg-emerald-50 rounded-xl space-y-1.5 border">
                            <span class="text-[10px] font-bold text-emerald-800 uppercase">Destino</span>
                            <select name="destino_local" class="w-full h-10 px-2 text-xs border rounded-lg bg-white"><option value="Banco">Conta Bancária</option><option value="Caixa">Caixa Físico</option></select>
                            <select name="destino_depto" class="w-full h-10 px-2 text-xs border rounded-lg bg-white"><option value="Geral">Fundo Geral</option>{% for d in lista_deptos %}<option value="{{ d['nome'] }}">{{ d['nome'] }}</option>{% endfor %}</select>
                        </div>
                        <input type="number" step="0.01" name="valor" required placeholder="Valor (MT) *" class="w-full h-11 px-3 text-sm font-bold border-2 rounded-xl">
                        <input type="text" name="motivo" required placeholder="Motivo da transferência *" class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                        <button type="submit" class="w-full h-12 bg-cyan-700 hover:bg-cyan-800 text-white font-black text-sm rounded-2xl shadow">Confirmar Transferência</button>
                    </form>
                </div>
                <div class="lg:col-span-7 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5">
                    <h3 class="text-base font-black text-slate-900 mb-4 pb-2 border-b">Histórico de Transferências</h3>
                    <div class="overflow-x-auto max-h-[500px]">
                        <table class="w-full text-left text-xs">
                            <thead class="bg-slate-50 border-b"><tr><th class="p-2.5">Data</th><th class="p-2.5">De</th><th class="p-2.5">Para</th><th class="p-2.5">Motivo</th><th class="p-2.5 text-right">Valor</th></tr></thead>
                            <tbody class="divide-y">
                                {% for t in ultimas_transferencias %}
                                <tr>
                                    <td class="p-2.5 font-bold">{{ t['data_movimento'] }}</td>
                                    <td class="p-2.5 text-rose-700 font-bold">{{ t['origem_local'] }}</td>
                                    <td class="p-2.5 text-emerald-700 font-bold">{{ t['destino_local'] }}</td>
                                    <td class="p-2.5">{{ t['motivo'] }}</td>
                                    <td class="p-2.5 text-right font-black text-blue-900">{{ "{:,.2f}".format(t['valor']) }} MT</td>
                                </tr>
                                {% endfor %}
                            </tbody>
                        </table>
                    </div>
                </div>
            </div>
        </section>

        <section id="aba-financeiro" class="tab-content space-y-6">
            <div class="grid grid-cols-1 lg:grid-cols-12 gap-6">
                <div class="lg:col-span-5 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5">
                    <h3 class="text-base font-black text-slate-900 mb-4 pb-2 border-b">💰 Lançamento no Caixa</h3>
                    <form action="/financeiro/novo" method="POST" class="space-y-3">
                        <input type="date" name="data_movimento" required class="w-full h-11 px-2 text-sm border-2 rounded-xl bg-white font-bold">
                        <div class="grid grid-cols-2 gap-2">
                            <select name="tipo" id="selectTipoTransacao" onchange="atualizarCategoriasPorTipo()" class="w-full h-11 px-2 text-sm border-2 rounded-xl bg-white font-bold"><option value="Entrada">Entrada (+)</option><option value="Saída">Saída (-)</option></select>
                            <select name="local_movimento" required class="w-full h-11 px-2 text-sm border-2 rounded-xl bg-white font-bold"><option value="Caixa">Caixa Físico</option><option value="Banco">Banco</option></select>
                        </div>
                        <div class="grid grid-cols-2 gap-2">
                            <select name="departamento" class="w-full h-11 px-2 text-sm border-2 rounded-xl bg-white"><option value="Geral">Fundo Geral</option>{% for d in lista_deptos %}<option value="{{ d['nome'] }}">{{ d['nome'] }}</option>{% endfor %}</select>
                            <select name="categoria" id="selectCategoriaTransacao" class="w-full h-11 px-2 text-sm border-2 rounded-xl bg-white"></select>
                        </div>
                        <input type="number" step="0.01" name="valor" required placeholder="Valor (MT) *" class="w-full h-11 px-3 text-sm font-black border-2 rounded-xl">
                        <input type="text" name="descricao" required placeholder="Descrição / Detalhes *" class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                        <button type="submit" class="w-full h-12 bg-emerald-600 hover:bg-emerald-700 text-white font-black text-sm rounded-2xl shadow">Gravar Movimento</button>
                    </form>
                </div>
                <div class="lg:col-span-7 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5">
                    <div class="flex justify-between items-center mb-4 pb-2 border-b">
                        <h3 class="text-base font-black text-slate-900">Extrato Financeiro</h3>
                        <a href="/exportar/financeiro" class="h-9 px-3 bg-emerald-600 text-white rounded-xl text-xs font-bold flex items-center gap-1">📥 Excel</a>
                    </div>
                    <div class="overflow-x-auto max-h-[500px]">
                        <table class="w-full text-left text-xs">
                            <thead class="bg-slate-50 border-b sticky top-0"><tr><th class="p-2.5">Data</th><th class="p-2.5">Conta</th><th class="p-2.5">Categoria</th><th class="p-2.5 text-right">Valor</th><th class="p-2.5 text-center">Recibo</th></tr></thead>
                            <tbody class="divide-y">
                                {% for f in todas_financas %}
                                <tr>
                                    <td class="p-2.5 font-bold">{{ f['data_movimento'] }}</td>
                                    <td class="p-2.5">{{ f['local_movimento'] }}</td>
                                    <td class="p-2.5">{{ f['categoria'] }}</td>
                                    <td class="p-2.5 text-right font-black {{ 'text-emerald-700' if f['tipo'] == 'Entrada' else 'text-rose-700' }}">{{ "{:,.2f}".format(f['valor']) }} MT</td>
                                    <td class="p-2.5 text-center"><a href="/financeiro/recibo_pdf/{{ f['id'] }}" class="bg-slate-100 hover:bg-slate-200 text-slate-700 font-bold px-2 py-1 rounded text-[11px]">📄 PDF</a></td>
                                </tr>
                                {% endfor %}
                            </tbody>
                        </table>
                    </div>
                </div>
            </div>
        </section>
        {% endif %}

        {% if e_admin %}
        <!-- ================= ABA 5: PATRIMÓNIO & INVENTÁRIO DO TEMPLO ================= -->
        <section id="aba-patrimonio" class="tab-content space-y-6">
            <div class="grid grid-cols-1 lg:grid-cols-12 gap-6">
                <div class="lg:col-span-5 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5">
                    <h3 class="text-base font-black text-slate-900 mb-4 pb-2 border-b">📦 Inventariar Bem / Património</h3>
                    <form action="/patrimonio/novo" method="POST" class="space-y-3">
                        <input type="text" name="item" required placeholder="Nome do Bem (Ex: Teclado Yamaha, Mesa de Som) *" class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                        <div class="grid grid-cols-2 gap-2">
                            <select name="departamento" class="w-full h-11 px-2 text-sm border-2 rounded-xl bg-white font-semibold">
                                <option value="Louvor / Som">Louvor / Som</option>
                                <option value="Templo / Mobília">Templo / Mobília</option>
                                <option value="Secretaria">Secretaria</option>
                                <option value="Geral">Geral</option>
                            </select>
                            <input type="number" name="quantidade" value="1" placeholder="Quantidade" class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                        </div>
                        <div class="grid grid-cols-2 gap-2">
                            <select name="estado_conservacao" class="w-full h-11 px-2 text-sm border-2 rounded-xl bg-white font-bold text-slate-700">
                                <option value="Bom">Bom Estado</option>
                                <option value="Em Manutenção">Em Manutenção</option>
                                <option value="Danificado">Danificado</option>
                            </select>
                            <input type="text" name="localizacao" placeholder="Local (Ex: Nave, Altar)" class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                        </div>
                        <input type="text" name="observacoes" placeholder="Notas ou número de série..." class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                        <button type="submit" class="w-full h-12 bg-amber-700 hover:bg-amber-800 text-white font-black text-sm rounded-2xl shadow">Registar no Inventário</button>
                    </form>
                </div>
                <div class="lg:col-span-7 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5">
                    <h3 class="text-base font-black text-slate-900 mb-4 pb-2 border-b">Bens Inventariados da IEAD Chicuque</h3>
                    <div class="overflow-x-auto">
                        <table class="w-full text-left text-xs">
                            <thead class="bg-slate-50 border-b"><tr><th class="p-2.5">Item</th><th class="p-2.5">Depto</th><th class="p-2.5 text-center">Qtd</th><th class="p-2.5">Estado</th><th class="p-2.5 text-center">Ação</th></tr></thead>
                            <tbody class="divide-y">
                                {% for p in todo_patrimonio %}
                                <tr>
                                    <td class="p-2.5 font-bold">{{ p['item'] }}<br/><span class="text-slate-400 font-normal">{{ p['localizacao'] }}</span></td>
                                    <td class="p-2.5">{{ p['departamento'] }}</td>
                                    <td class="p-2.5 text-center font-bold">{{ p['quantidade'] }}</td>
                                    <td class="p-2.5">
                                        <span class="px-2 py-0.5 rounded font-bold text-[11px] {{ 'bg-emerald-100 text-emerald-800' if p['estado_conservacao'] == 'Bom' else 'bg-amber-100 text-amber-800' }}">{{ p['estado_conservacao'] }}</span>
                                    </td>
                                    <td class="p-2.5 text-center"><a href="/apagar/patrimonio/{{ p['id'] }}" onclick="return confirm('Eliminar?');" class="text-rose-600 font-bold">✕</a></td>
                                </tr>
                                {% endfor %}
                            </tbody>
                        </table>
                    </div>
                </div>
            </div>
        </section>

        <!-- ================= ABA 6: CONFIGURAÇÕES & UTILIZADORES ================= -->
        <section id="aba-configuracoes" class="tab-content space-y-6">
            <div class="grid grid-cols-1 md:grid-cols-3 gap-6">
                <div class="bg-white/95 p-5 rounded-3xl card-glow border border-indigo-100">
                    <h3 class="text-sm font-black uppercase text-blue-950 mb-3 pb-2 border-b">Ministérios</h3>
                    <form action="/config/departamento/novo" method="POST" class="flex gap-2 mb-3">
                        <input type="text" name="nome" placeholder="Novo..." required class="h-10 px-2 text-xs border rounded-xl w-full">
                        <button type="submit" class="bg-blue-900 text-white px-3 font-bold rounded-xl">+</button>
                    </form>
                    <ul class="divide-y text-xs max-h-52 overflow-y-auto">
                        {% for d in lista_deptos %}<li class="py-2 flex justify-between"><span>{{ d['nome'] }}</span><a href="/config/departamento/apagar/{{ d['id'] }}" class="text-rose-600 font-bold">✕</a></li>{% endfor %}
                    </ul>
                </div>
                <div class="bg-white/95 p-5 rounded-3xl card-glow border border-indigo-100">
                    <h3 class="text-sm font-black uppercase text-emerald-950 mb-3 pb-2 border-b">Categorias Caixa</h3>
                    <form action="/config/categoria/novo" method="POST" class="space-y-2 mb-3 text-xs">
                        <div class="flex gap-2">
                            <select name="tipo" class="h-10 px-2 border rounded-xl bg-white w-1/3"><option value="Entrada">Entrada</option><option value="Saída">Saída</option></select>
                            <input type="text" name="nome" placeholder="Nome..." required class="h-10 px-2 border rounded-xl w-2/3">
                        </div>
                        <button type="submit" class="w-full h-9 bg-emerald-600 text-white font-bold rounded-xl">+ Guardar</button>
                    </form>
                    <ul class="divide-y text-xs max-h-52 overflow-y-auto">
                        {% for c in lista_categorias %}<li class="py-2 flex justify-between"><span><strong>{{ c['tipo'] }}</strong>: {{ c['nome'] }}</span><a href="/config/categoria/apagar/{{ c['id'] }}" class="text-rose-600 font-bold">✕</a></li>{% endfor %}
                    </ul>
                </div>
                <div class="bg-white/95 p-5 rounded-3xl card-glow border border-indigo-100">
                    <h3 class="text-sm font-black uppercase text-amber-950 mb-3 pb-2 border-b">Zonas & Bairros</h3>
                    <form action="/config/zona/novo" method="POST" class="flex gap-2 mb-3">
                        <input type="text" name="nome" placeholder="Nova zona..." required class="h-10 px-2 text-xs border rounded-xl w-full">
                        <button type="submit" class="bg-amber-600 text-white px-3 font-bold rounded-xl">+</button>
                    </form>
                    <ul class="divide-y text-xs max-h-52 overflow-y-auto">
                        {% for z in lista_zonas %}<li class="py-2 flex justify-between"><span>{{ z['nome'] }}</span><a href="/config/zona/apagar/{{ z['id'] }}" class="text-rose-600 font-bold">✕</a></li>{% endfor %}
                    </ul>
                </div>
            </div>
        </section>

        <section id="aba-usuarios" class="tab-content space-y-6">
            <div class="grid grid-cols-1 md:grid-cols-12 gap-6">
                <div class="md:col-span-5 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5">
                    <h3 class="text-base font-black text-slate-900 mb-4 pb-2 border-b">👤 Criar Novo Acesso</h3>
                    <form action="/usuarios/novo" method="POST" class="space-y-3">
                        <input type="text" name="usuario" required placeholder="Login *" class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                        <input type="password" name="senha" required placeholder="Palavra-passe *" class="w-full h-11 px-3 text-sm border-2 rounded-xl">
                        <select name="cargo" required class="w-full h-11 px-2 text-sm border-2 rounded-xl bg-white font-bold text-blue-950">
                            <option value="Secretário">Secretário (Cadastro & Atos)</option>
                            <option value="Líder">Líder de Ministério</option>
                            <option value="Tesoureiro">Tesoureiro (Tesouraria & Dashboard)</option>
                            <option value="Pastor">Pastor (Administrador Geral)</option>
                        </select>
                        <button type="submit" class="w-full h-12 bg-emerald-700 text-white font-black text-sm rounded-2xl shadow">Criar Utilizador</button>
                    </form>
                </div>
                <div class="md:col-span-7 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5">
                    <h3 class="text-base font-black text-slate-900 mb-4 pb-2 border-b">Utilizadores Registados</h3>
                    <table class="w-full text-left text-xs">
                        <thead class="bg-slate-50 border-b"><tr><th class="p-2.5">Login</th><th class="p-2.5">Permissão</th><th class="p-2.5 text-center">Ação</th></tr></thead>
                        <tbody class="divide-y">
                            {% for u in lista_usuarios %}
                            <tr>
                                <td class="p-2.5 font-bold">{{ u['usuario'] }}</td>
                                <td class="p-2.5"><span class="px-2 py-0.5 rounded font-bold bg-blue-100 text-blue-900">{{ u['cargo'] }}</span></td>
                                <td class="p-2.5 text-center">{% if u['usuario'] != 'admin' %}<a href="/usuarios/apagar/{{ u['id'] }}" onclick="return confirm('Eliminar?');" class="text-rose-600 font-bold">✕</a>{% else %}<span class="text-slate-400">Principal</span>{% endif %}</td>
                            </tr>
                            {% endfor %}
                        </tbody>
                    </table>
                </div>
            </div>
        </section>
        {% endif %}

    </main>

    <script>
        const todasCategorias = {{ categorias_json | safe }};

        function atualizarCategoriasPorTipo() {
            const el = document.getElementById('selectTipoTransacao');
            if (!el) return;
            const tipo = el.value;
            const selectCat = document.getElementById('selectCategoriaTransacao');
            selectCat.innerHTML = '';
            const filtradas = todasCategorias.filter(c => c.tipo === tipo);
            filtradas.forEach(c => {
                const opt = document.createElement('option');
                opt.value = c.nome;
                opt.textContent = c.nome;
                selectCat.appendChild(opt);
            });
        }
        document.addEventListener('DOMContentLoaded', atualizarCategoriasPorTipo);

        function trocarAba(abaId) {
            document.querySelectorAll('.tab-content').forEach(el => el.classList.remove('active'));
            document.querySelectorAll('.tab-btn').forEach(btn => {
                btn.className = 'tab-btn px-4 py-2 rounded-2xl transition flex items-center gap-1.5 text-indigo-100 hover:bg-white/10';
            });
            const abaAlvo = document.getElementById('aba-' + abaId);
            if (abaAlvo) abaAlvo.classList.add('active');
            const activeBtn = document.getElementById('btn-' + abaId);
            if (activeBtn) activeBtn.className = 'tab-btn px-4 py-2 rounded-2xl transition flex items-center gap-1.5 bg-gradient-to-r from-blue-600 to-indigo-600 text-white font-black shadow-md';
        }

        {% if pode_tesouraria %}
        const paletaCores = ['#059669', '#2563eb', '#d97706', '#7c3aed', '#db2777', '#0891b2', '#ea580c', '#4f46e5'];
        const paletaSaidas = ['#e11d48', '#f97316', '#dc2626', '#9333ea', '#c026d3', '#b91c1c', '#64748b'];

        new Chart(document.getElementById('graficoEvolucaoMensal').getContext('2d'), {
            type: 'line',
            data: {
                labels: ['Jan', 'Fev', 'Mar', 'Abr', 'Mai', 'Jun', 'Jul', 'Ago', 'Set', 'Out', 'Nov', 'Dez'],
                datasets: [
                    { label: 'Entradas', data: {{ evolucao_entradas | safe }}, borderColor: '#059669', backgroundColor: 'rgba(5, 150, 105, 0.1)', borderWidth: 3, tension: 0.35, fill: true },
                    { label: 'Saídas', data: {{ evolucao_saidas | safe }}, borderColor: '#e11d48', backgroundColor: 'rgba(225, 29, 72, 0.05)', borderWidth: 3, tension: 0.35, fill: true }
                ]
            },
            options: { responsive: true, maintainAspectRatio: false }
        });

        new Chart(document.getElementById('graficoFinancasDepto').getContext('2d'), {
            type: 'bar',
            data: {
                labels: {{ depto_fin_labels | safe }},
                datasets: [
                    { label: 'Entradas', data: {{ depto_fin_entradas | safe }}, backgroundColor: '#059669', borderRadius: 6 },
                    { label: 'Saídas', data: {{ depto_fin_saidas | safe }}, backgroundColor: '#e11d48', borderRadius: 6 }
                ]
            },
            options: { responsive: true, maintainAspectRatio: false }
        });

        new Chart(document.getElementById('graficoFontesFundos').getContext('2d'), {
            type: 'doughnut',
            data: {
                labels: {{ fontes_labels | safe }},
                datasets: [{ data: {{ fontes_valores | safe }}, backgroundColor: paletaCores, borderWidth: 2, borderColor: '#ffffff' }]
            },
            options: { responsive: true, maintainAspectRatio: false, plugins: { legend: { position: 'bottom', labels: { boxWidth: 10, font: { size: 10 } } } } }
        });

        new Chart(document.getElementById('graficoSaidasCategorias').getContext('2d'), {
            type: 'doughnut',
            data: {
                labels: {{ saidas_labels | safe }},
                datasets: [{ data: {{ saidas_valores | safe }}, backgroundColor: paletaSaidas, borderWidth: 2, borderColor: '#ffffff' }]
            },
            options: { responsive: true, maintainAspectRatio: false, plugins: { legend: { position: 'bottom', labels: { boxWidth: 10, font: { size: 10 } } } } }
        });
        {% endif %}
    </script>
</body>
</html>
"""

with open(os.path.join("templates", "dashboard.html"), "w", encoding="utf-8") as f:
    f.write(html_code)

print("✓ 2/2: templates/dashboard.html atualizado com todas as abas e botões visíveis!")