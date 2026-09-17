import os

# 1. Garante todas as pastas necessárias
os.makedirs("templates", exist_ok=True)
os.makedirs("static", exist_ok=True)
os.makedirs(os.path.join("static", "uploads"), exist_ok=True)

# 2. requirements.txt atualizado
reqs = """flask
gunicorn
openpyxl
werkzeug
"""
with open("requirements.txt", "w", encoding="utf-8") as f:
    f.write(reqs)
print("✓ 1/4: requirements.txt atualizado!")

# 3. templates/dashboard.html com cadastro completo e foto
html_code = """<!DOCTYPE html>
<html lang="pt">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0, maximum-scale=1.0, user-scalable=no">
    <title>IEAD Chicuque - Gestão Integrada</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <script src="https://cdn.jsdelivr.net/npm/chart.js"></script>
    <style>
        .tab-content { display: none; }
        .tab-content.active { display: block; }
    </style>
</head>
<body class="bg-slate-100 text-slate-800 font-sans min-h-screen">

    <header class="bg-gradient-to-r from-blue-950 via-blue-900 to-indigo-900 text-white shadow sticky top-0 z-40">
        <div class="max-w-7xl mx-auto px-4 py-3 flex flex-wrap justify-between items-center gap-3">
            <div class="flex items-center space-x-3">
                <img src="/static/logo.svg?v=2026" alt="IEAD" class="w-10 h-10 object-contain bg-white rounded-full p-1 shadow" onerror="this.style.display='none'">
                <div>
                    <h1 class="text-xs sm:text-sm md:text-base font-bold uppercase tracking-tight">Igreja Evangélica Assembleia de Deus</h1>
                    <p class="text-[11px] text-blue-200">Chicuque • Membresia & Gestão Integrada</p>
                </div>
            </div>
            <div class="flex items-center space-x-3">
                <span class="text-xs bg-blue-800/80 px-3 py-1.5 rounded-lg border border-blue-700 text-blue-100">
                    {{ session['usuario'] }} ({{ session['cargo'] }})
                </span>
                <a href="/logout" class="bg-rose-600 hover:bg-rose-700 text-white text-xs px-3 py-1.5 rounded-lg font-bold transition">Sair</a>
            </div>
        </div>

        <nav class="max-w-7xl mx-auto px-4 flex overflow-x-auto space-x-2 border-t border-blue-800/60 pt-2 pb-1 text-xs">
            <button onclick="trocarAba('membros')" id="btn-membros" class="tab-btn px-4 py-2 font-bold rounded-t-lg bg-white text-blue-900 shadow">👤 Membresia ({{ total_membros }})</button>
            <button onclick="trocarAba('dashboard')" id="btn-dashboard" class="tab-btn px-4 py-2 font-bold rounded-t-lg text-blue-200 hover:bg-blue-800/50">📊 Painel Geral</button>
            <button onclick="trocarAba('financeiro')" id="btn-financeiro" class="tab-btn px-4 py-2 font-bold rounded-t-lg text-blue-200 hover:bg-blue-800/50">💰 Tesouraria</button>
            <button onclick="trocarAba('transferencias')" id="btn-transferencias" class="tab-btn px-4 py-2 font-bold rounded-t-lg text-cyan-300 hover:bg-blue-800/50">🔄 Transferir</button>
            <button onclick="trocarAba('casamentos')" id="btn-casamentos" class="tab-btn px-4 py-2 font-bold rounded-t-lg text-blue-200 hover:bg-blue-800/50">💍 Casamentos</button>
            <button onclick="trocarAba('mortes')" id="btn-mortes" class="tab-btn px-4 py-2 font-bold rounded-t-lg text-blue-200 hover:bg-blue-800/50">🕊️ Óbitos</button>
            <button onclick="trocarAba('configuracoes')" id="btn-configuracoes" class="tab-btn px-4 py-2 font-bold rounded-t-lg text-amber-300 hover:bg-blue-800/50">⚙️ Zonas & Deptos</button>
            <button onclick="trocarAba('usuarios')" id="btn-usuarios" class="tab-btn px-4 py-2 font-bold rounded-t-lg text-emerald-300 hover:bg-blue-800/50">👥 Utilizadores</button>
        </nav>
    </header>

    <main class="max-w-7xl mx-auto p-4 space-y-6">

        <!-- ABA MEMBROS -->
        <section id="aba-membros" class="tab-content active space-y-6">
            <div class="grid grid-cols-1 lg:grid-cols-3 gap-6">
                
                <!-- Formulário Completo -->
                <div class="bg-white p-5 rounded-2xl shadow-sm border border-slate-200">
                    <div class="flex items-center justify-between pb-3 mb-4 border-b">
                        <h3 class="text-sm font-bold uppercase text-blue-900 flex items-center gap-2">
                            <span>📝</span> Ficha Completa de Membro
                        </h3>
                        <span class="text-[10px] bg-blue-100 text-blue-800 font-bold px-2 py-0.5 rounded">Câmara / Upload</span>
                    </div>

                    <form action="/membros/novo" method="POST" enctype="multipart/form-data" class="space-y-4 text-xs">
                        
                        <!-- Upload de Fotografia -->
                        <div class="p-3 bg-slate-50 border border-dashed border-slate-300 rounded-xl text-center">
                            <label class="block font-bold text-slate-700 mb-1">Fotografia do Membro (Câmara / Galeria)</label>
                            <input type="file" name="foto" accept="image/*" capture="environment" 
                                   class="w-full text-[11px] text-slate-500 file:mr-2 file:py-1.5 file:px-3 file:rounded-lg file:border-0 file:text-xs file:font-semibold file:bg-blue-900 file:text-white hover:file:bg-blue-800">
                        </div>

                        <!-- 1. Dados Pessoais -->
                        <div class="space-y-2">
                            <span class="font-bold text-slate-500 uppercase text-[10px] block border-b pb-1">1. Dados Pessoais</span>
                            <div>
                                <label class="block font-bold text-slate-700 mb-1">Nome Completo *</label>
                                <input type="text" name="nome" required placeholder="Nome oficial..." class="w-full p-2.5 border rounded-lg outline-none font-medium">
                            </div>
                            <div class="grid grid-cols-2 gap-2">
                                <div>
                                    <label class="block font-medium text-slate-700 mb-1">Contacto Telefónico</label>
                                    <input type="text" name="telefone" placeholder="+258 8..." class="w-full p-2 border rounded-lg outline-none">
                                </div>
                                <div>
                                    <label class="block font-medium text-slate-700 mb-1">Género</label>
                                    <select name="genero" class="w-full p-2 border rounded-lg bg-white outline-none">
                                        <option value="Masculino">Masculino</option>
                                        <option value="Feminino">Feminino</option>
                                    </select>
                                </div>
                            </div>
                            <div class="grid grid-cols-2 gap-2">
                                <div>
                                    <label class="block font-medium text-slate-700 mb-1">Data Nascimento</label>
                                    <input type="date" name="data_nascimento" class="w-full p-2 border rounded-lg outline-none">
                                </div>
                                <div>
                                    <label class="block font-medium text-slate-700 mb-1">Segmento Etário *</label>
                                    <select name="faixa_etaria" required class="w-full p-2 border rounded-lg bg-white outline-none font-bold text-blue-900">
                                        <option value="Adulto">Adulto</option>
                                        <option value="Jovem">Jovem</option>
                                        <option value="Adolescente">Adolescente</option>
                                        <option value="Criança">Criança (Boa Esperança)</option>
                                        <option value="Terceira Idade">Terceira Idade</option>
                                    </select>
                                </div>
                            </div>
                            <div class="grid grid-cols-2 gap-2">
                                <div>
                                    <label class="block font-medium text-slate-700 mb-1">Naturalidade</label>
                                    <input type="text" name="naturalidade" placeholder="Província / Distrito" class="w-full p-2 border rounded-lg outline-none">
                                </div>
                                <div>
                                    <label class="block font-medium text-slate-700 mb-1">Bairro / Zona</label>
                                    <select name="bairro" class="w-full p-2 border rounded-lg bg-white outline-none">
                                        <option value="">-- Selecione --</option>
                                        {% for z in lista_zonas %}<option value="{{ z['nome'] }}">{{ z['nome'] }}</option>{% endfor %}
                                    </select>
                                </div>
                            </div>
                            <div>
                                <label class="block font-medium text-slate-700 mb-1">Filiação (Pai & Mãe)</label>
                                <input type="text" name="filiacao" placeholder="Nome dos pais..." class="w-full p-2 border rounded-lg outline-none">
                            </div>
                        </div>

                        <!-- 2. Documentação -->
                        <div class="space-y-2 pt-1">
                            <span class="font-bold text-slate-500 uppercase text-[10px] block border-b pb-1">2. Documento Civil</span>
                            <div class="grid grid-cols-3 gap-2">
                                <div>
                                    <label class="block font-medium text-slate-700 mb-1">Tipo</label>
                                    <select name="tipo_documento" class="w-full p-2 border rounded-lg bg-white outline-none font-semibold">
                                        <option value="BI">B.I.</option>
                                        <option value="Cédula">Cédula</option>
                                        <option value="Cartão de Eleitor">C. Eleitor</option>
                                        <option value="Passaporte">Passaporte</option>
                                        <option value="Nenhum">Sem Doc.</option>
                                    </select>
                                </div>
                                <div class="col-span-2">
                                    <label class="block font-medium text-slate-700 mb-1">Número do Documento</label>
                                    <input type="text" name="numero_documento" placeholder="Número..." class="w-full p-2 border rounded-lg outline-none font-mono">
                                </div>
                            </div>
                        </div>

                        <!-- 3. Eclesiástica e Departamentos -->
                        <div class="space-y-2 pt-1">
                            <span class="font-bold text-slate-500 uppercase text-[10px] block border-b pb-1">3. Vida Eclesiástica & Cargos</span>
                            <div class="grid grid-cols-2 gap-2">
                                <div>
                                    <label class="block font-medium text-slate-700 mb-1">Ano Conversão</label>
                                    <input type="number" name="ano_conversao" placeholder="Ex: 2018" min="1930" max="2035" class="w-full p-2 border rounded-lg outline-none">
                                </div>
                                <div>
                                    <label class="block font-medium text-slate-700 mb-1">Data Batismo Águas</label>
                                    <input type="date" name="data_batismo" class="w-full p-2 border rounded-lg outline-none">
                                </div>
                            </div>
                            <div class="grid grid-cols-2 gap-2">
                                <div>
                                    <label class="block font-bold text-blue-950 mb-1">Posição / Cargo *</label>
                                    <select name="posicao_atual" required class="w-full p-2 border rounded-lg bg-white outline-none font-semibold">
                                        <option value="Membro em Comunhão">Membro em Comunhão</option>
                                        <option value="Pastor">Pastor</option>
                                        <option value="Presbítero">Presbítero</option>
                                        <option value="Diácono / Diaconisa">Diácono / Diaconisa</option>
                                        <option value="Evangelista">Evangelista</option>
                                        <option value="Obreiro(a)">Obreiro(a)</option>
                                        <option value="Líder de Célula">Líder de Célula</option>
                                        <option value="Membro em Prova">Membro em Prova</option>
                                    </select>
                                </div>
                                <div>
                                    <label class="block font-bold text-blue-950 mb-1">Departamento *</label>
                                    <select name="departamento" required class="w-full p-2 border rounded-lg bg-white outline-none font-bold text-indigo-900">
                                        <option value="Activista">Activista</option>
                                        <option value="Juventude">Juventude</option>
                                        <option value="Mulher (Senhoras)">Mulher (Senhoras)</option>
                                        <option value="Boa Esperança (Crianças)">Boa Esperança (Crianças)</option>
                                        <option value="Homens / Obreiros">Homens / Obreiros</option>
                                        <option value="Louvor / Música">Louvor / Música</option>
                                        <option value="Ação Social">Ação Social</option>
                                        <option value="Geral">Geral</option>
                                    </select>
                                </div>
                            </div>
                            <div>
                                <label class="block font-medium text-slate-700 mb-1">Espaço de Progressões Eclesiásticas</label>
                                <textarea name="progressoes" rows="2" placeholder="Histórico de consagrações e cargos anteriores..." class="w-full p-2 border rounded-lg outline-none text-[11px]"></textarea>
                            </div>
                            <div>
                                <label class="block font-medium text-slate-700 mb-1">Observações Gerais</label>
                                <input type="text" name="observacoes" placeholder="Notas pastorais..." class="w-full p-2 border rounded-lg outline-none">
                            </div>
                        </div>

                        <button type="submit" class="w-full bg-blue-900 hover:bg-blue-950 text-white font-bold py-3 rounded-xl shadow-md transition">
                            Gravar Membro Completo
                        </button>
                    </form>
                </div>

                <!-- Tabela de Membros -->
                <div class="lg:col-span-2 bg-white p-5 rounded-2xl shadow-sm border border-slate-200 flex flex-col justify-between">
                    <div>
                        <div class="flex flex-col sm:flex-row justify-between items-stretch sm:items-center gap-3 mb-4 pb-3 border-b">
                            <input type="text" id="filtroMembros" onkeyup="filtrarTabela('filtroMembros', 'tabelaMembros')" 
                                   placeholder="🔍 Pesquisar por nome, BI, cargo, departamento..." 
                                   class="text-xs p-2.5 border rounded-lg outline-none w-full sm:w-80">
                            
                            <a href="/exportar/membros" class="text-xs bg-emerald-600 hover:bg-emerald-700 text-white px-3 py-2.5 rounded-lg font-bold flex items-center justify-center gap-1 transition">
                                📥 Exportar Excel (.xlsx)
                            </a>
                        </div>

                        <div class="overflow-x-auto max-h-[620px]">
                            <table id="tabelaMembros" class="w-full text-left text-xs min-w-[650px]">
                                <thead class="bg-slate-50 text-slate-600 border-b sticky top-0 z-10">
                                    <tr>
                                        <th class="p-2.5">Foto</th>
                                        <th class="p-2.5">Nome & Segmento</th>
                                        <th class="p-2.5">Cargo</th>
                                        <th class="p-2.5">Departamento</th>
                                        <th class="p-2.5">Documento</th>
                                        <th class="p-2.5 text-center">Ações</th>
                                    </tr>
                                </thead>
                                <tbody class="divide-y divide-slate-100">
                                    {% for m in todos_membros %}
                                    <tr class="hover:bg-slate-50/80">
                                        <td class="p-2.5">
                                            {% if m['foto_path'] %}
                                            <img src="{{ m['foto_path'] }}" class="w-10 h-10 object-cover rounded-full border border-slate-300 shadow-sm">
                                            {% else %}
                                            <div class="w-10 h-10 rounded-full bg-blue-100 text-blue-900 flex items-center justify-center font-bold text-xs">
                                                {{ m['nome'][:2].upper() }}
                                            </div>
                                            {% endif %}
                                        </td>
                                        <td class="p-2.5">
                                            <span class="font-bold text-slate-800 block">{{ m['nome'] }}</span>
                                            <span class="text-[10px] text-slate-400 block">{{ m['faixa_etaria'] or 'Adulto' }} • {{ m['telefone'] or 'Sem tel' }}</span>
                                        </td>
                                        <td class="p-2.5">
                                            <span class="bg-blue-50 text-blue-800 px-2 py-0.5 rounded font-bold text-[11px] block w-fit">
                                                {{ m['posicao_atual'] or 'Membro' }}
                                            </span>
                                            {% if m['progressoes'] %}
                                            <span class="text-[9px] text-indigo-600 block mt-0.5 truncate max-w-[130px]" title="{{ m['progressoes'] }}">📈 {{ m['progressoes'] }}</span>
                                            {% endif %}
                                        </td>
                                        <td class="p-2.5">
                                            <span class="px-2 py-0.5 rounded text-[10px] font-bold 
                                                {% if m['departamento'] == 'Activista' %}bg-amber-100 text-amber-800
                                                {% elif m['departamento'] == 'Juventude' %}bg-purple-100 text-purple-800
                                                {% elif m['departamento'] == 'Mulher (Senhoras)' %}bg-rose-100 text-rose-800
                                                {% elif m['departamento'] == 'Boa Esperança (Crianças)' %}bg-emerald-100 text-emerald-800
                                                {% else %}bg-slate-100 text-slate-700{% endif %}">
                                                {{ m['departamento'] or 'Geral' }}
                                            </span>
                                        </td>
                                        <td class="p-2.5 font-mono text-[11px] text-slate-600">
                                            {{ m['tipo_documento'] or 'BI' }}: {{ m['numero_documento'] or '-' }}
                                        </td>
                                        <td class="p-2.5 text-center whitespace-nowrap">
                                            <button onclick="verCartaoMembro({{ m['id'] }})" class="bg-blue-100 hover:bg-blue-200 text-blue-800 px-2 py-1 rounded text-[11px] font-bold mr-1">Ficha</button>
                                            <a href="/apagar/membro/{{ m['id'] }}" onclick="return confirm('Eliminar o registo?');" class="text-rose-600 hover:underline font-bold text-[11px]">Apagar</a>
                                        </td>
                                    </tr>
                                    {% endfor %}
                                </tbody>
                            </table>
                        </div>
                    </div>
                </div>

            </div>
        </section>

        <!-- ABA DASHBOARD -->
        <section id="aba-dashboard" class="tab-content space-y-6">
            <div class="grid grid-cols-1 sm:grid-cols-2 lg:grid-cols-4 gap-4">
                <div class="bg-white p-4 rounded-xl shadow-sm border border-slate-200">
                    <p class="text-[11px] font-bold uppercase text-amber-700">💵 Saldo Caixa Físico</p>
                    <p class="text-2xl font-bold text-amber-800 mt-1">{{ "{:,.2f}".format(saldo_caixa) }} MT</p>
                </div>
                <div class="bg-white p-4 rounded-xl shadow-sm border border-slate-200">
                    <p class="text-[11px] font-bold uppercase text-indigo-700">🏛️ Saldo Conta Bancária</p>
                    <p class="text-2xl font-bold text-indigo-800 mt-1">{{ "{:,.2f}".format(saldo_banco) }} MT</p>
                </div>
                <div class="bg-white p-4 rounded-xl shadow-sm border border-slate-200 bg-gradient-to-br from-white to-blue-50">
                    <p class="text-[11px] font-bold uppercase text-blue-900">🏦 Saldo Total Geral</p>
                    <p class="text-2xl font-black text-blue-800 mt-1">{{ "{:,.2f}".format(saldo_total) }} MT</p>
                </div>
                <div class="bg-white p-4 rounded-xl shadow-sm border border-slate-200">
                    <p class="text-[11px] font-bold uppercase text-slate-500">👥 Total Membros</p>
                    <p class="text-2xl font-bold text-slate-800 mt-1">{{ total_membros }}</p>
                </div>
            </div>
            <div class="grid grid-cols-1 lg:grid-cols-2 gap-6">
                <div class="bg-white p-5 rounded-xl shadow-sm border border-slate-200">
                    <h3 class="text-xs font-bold uppercase text-slate-600 mb-3">📈 Finanças Mensais</h3>
                    <div class="h-64"><canvas id="graficoEvolucaoMensal"></canvas></div>
                </div>
                <div class="bg-white p-5 rounded-xl shadow-sm border border-slate-200">
                    <h3 class="text-xs font-bold uppercase text-slate-600 mb-3">👥 Distribuição por Departamento</h3>
                    <div class="h-64"><canvas id="graficoDeptos"></canvas></div>
                </div>
            </div>
        </section>

        <!-- ABA TRANSFERENCIAS -->
        <section id="aba-transferencias" class="tab-content space-y-6">
            <div class="grid grid-cols-1 lg:grid-cols-3 gap-6">
                <div class="bg-white p-5 rounded-xl shadow-sm border border-slate-200 h-fit">
                    <h3 class="text-sm font-bold uppercase text-slate-800 mb-3 pb-2 border-b">🔄 Efetuar Transferência</h3>
                    <form action="/financeiro/transferir" method="POST" class="space-y-3 text-xs">
                        <div><label class="block font-bold mb-1">Data *</label><input type="date" name="data_movimento" required class="w-full p-2 border rounded-lg"></div>
                        <div class="p-3 bg-rose-50 rounded-lg space-y-2">
                            <span class="font-bold text-rose-800 text-[10px] uppercase">Origem</span>
                            <select name="origem_local" class="w-full p-2 border rounded-lg bg-white"><option value="Caixa">Caixa</option><option value="Banco">Banco</option></select>
                            <select name="origem_depto" class="w-full p-2 border rounded-lg bg-white"><option value="Geral">Geral</option>{% for d in lista_deptos %}<option value="{{ d['nome'] }}">{{ d['nome'] }}</option>{% endfor %}</select>
                        </div>
                        <div class="p-3 bg-emerald-50 rounded-lg space-y-2">
                            <span class="font-bold text-emerald-800 text-[10px] uppercase">Destino</span>
                            <select name="destino_local" class="w-full p-2 border rounded-lg bg-white"><option value="Banco">Banco</option><option value="Caixa">Caixa</option></select>
                            <select name="destino_depto" class="w-full p-2 border rounded-lg bg-white"><option value="Geral">Geral</option>{% for d in lista_deptos %}<option value="{{ d['nome'] }}">{{ d['nome'] }}</option>{% endfor %}</select>
                        </div>
                        <div><label class="block font-bold mb-1">Valor (MT) *</label><input type="number" step="0.01" name="valor" required class="w-full p-2 border rounded-lg font-bold"></div>
                        <div><label class="block font-bold mb-1">Motivo *</label><input type="text" name="motivo" required class="w-full p-2 border rounded-lg"></div>
                        <button type="submit" class="w-full bg-cyan-700 text-white font-bold py-2.5 rounded-lg">Confirmar Transferência</button>
                    </form>
                </div>
                <div class="lg:col-span-2 bg-white p-5 rounded-xl shadow-sm border border-slate-200">
                    <h3 class="text-sm font-bold uppercase mb-3 pb-2 border-b">Histórico de Transferências</h3>
                    <div class="overflow-x-auto max-h-[500px]">
                        <table class="w-full text-left text-xs">
                            <thead class="bg-slate-50 border-b"><tr><th class="p-2">Data</th><th class="p-2">De</th><th class="p-2">Para</th><th class="p-2">Motivo</th><th class="p-2 text-right">Valor</th></tr></thead>
                            <tbody class="divide-y divide-slate-100">
                                {% for t in ultimas_transferencias %}
                                <tr><td class="p-2">{{ t['data_movimento'] }}</td><td class="p-2 font-bold text-rose-700">{{ t['origem_local'] }} ({{ t['origem_depto'] }})</td><td class="p-2 font-bold text-emerald-700">{{ t['destino_local'] }} ({{ t['destino_depto'] }})</td><td class="p-2">{{ t['motivo'] }}</td><td class="p-2 text-right font-bold">{{ "{:,.2f}".format(t['valor']) }} MT</td></tr>
                                {% endfor %}
                            </tbody>
                        </table>
                    </div>
                </div>
            </div>
        </section>

        <!-- ABA TESOURARIA -->
        <section id="aba-financeiro" class="tab-content space-y-6">
            <div class="grid grid-cols-1 lg:grid-cols-3 gap-6">
                <div class="bg-white p-5 rounded-xl shadow-sm border border-slate-200 h-fit">
                    <h3 class="text-sm font-bold uppercase text-emerald-700 mb-3 pb-2 border-b">Lançamento de Caixa</h3>
                    <form action="/financeiro/novo" method="POST" class="space-y-3 text-xs">
                        <div><label class="block font-bold mb-1">Data *</label><input type="date" name="data_movimento" required class="w-full p-2 border rounded-lg"></div>
                        <div class="grid grid-cols-2 gap-2">
                            <div><label class="block font-bold mb-1">Tipo</label><select name="tipo" id="selectTipoTransacao" onchange="atualizarCategoriasPorTipo()" class="w-full p-2 border rounded-lg bg-white"><option value="Entrada">Entrada (+)</option><option value="Saída">Saída (-)</option></select></div>
                            <div><label class="block font-bold mb-1">Conta *</label><select name="local_movimento" required class="w-full p-2 border rounded-lg bg-white font-semibold"><option value="Caixa">Caixa</option><option value="Banco">Banco</option></select></div>
                        </div>
                        <div class="grid grid-cols-2 gap-2">
                            <div><label class="block font-bold mb-1">Departamento</label><select name="departamento" class="w-full p-2 border rounded-lg bg-white"><option value="Geral">Geral</option>{% for d in lista_deptos %}<option value="{{ d['nome'] }}">{{ d['nome'] }}</option>{% endfor %}</select></div>
                            <div><label class="block font-bold mb-1">Categoria</label><select name="categoria" id="selectCategoriaTransacao" class="w-full p-2 border rounded-lg bg-white"></select></div>
                        </div>
                        <div><label class="block font-bold mb-1">Valor (MT) *</label><input type="number" step="0.01" name="valor" required class="w-full p-2 border rounded-lg font-bold"></div>
                        <div><label class="block font-bold mb-1">Descrição *</label><input type="text" name="descricao" required class="w-full p-2 border rounded-lg"></div>
                        <button type="submit" class="w-full bg-emerald-600 text-white font-bold py-2.5 rounded-lg">Gravar Movimento</button>
                    </form>
                </div>
                <div class="lg:col-span-2 bg-white p-5 rounded-xl shadow-sm border border-slate-200">
                    <h3 class="text-sm font-bold uppercase mb-3 pb-2 border-b">Movimentos Registados</h3>
                    <div class="overflow-x-auto max-h-[500px]">
                        <table class="w-full text-left text-xs">
                            <thead class="bg-slate-50 border-b sticky top-0"><tr><th class="p-2">Data</th><th class="p-2">Tipo</th><th class="p-2">Conta</th><th class="p-2">Fundo</th><th class="p-2">Categoria</th><th class="p-2">Descrição</th><th class="p-2 text-right">Valor</th><th class="p-2 text-center">Ação</th></tr></thead>
                            <tbody class="divide-y divide-slate-100">
                                {% for f in todas_financas %}
                                <tr>
                                    <td class="p-2">{{ f['data_movimento'] }}</td>
                                    <td class="p-2"><span class="px-2 py-0.5 rounded text-[10px] font-bold {{ 'bg-emerald-100 text-emerald-700' if f['tipo'] == 'Entrada' else 'bg-rose-100 text-rose-700' }}">{{ f['tipo'] }}</span></td>
                                    <td class="p-2 font-bold">{{ f['local_movimento'] }}</td>
                                    <td class="p-2">{{ f['departamento'] or 'Geral' }}</td>
                                    <td class="p-2">{{ f['categoria'] }}</td>
                                    <td class="p-2 text-slate-600">{{ f['descricao'] }}</td>
                                    <td class="p-2 text-right font-bold {{ 'text-emerald-600' if f['tipo'] == 'Entrada' else 'text-rose-600' }}">{{ "{:,.2f}".format(f['valor']) }} MT</td>
                                    <td class="p-2 text-center"><a href="/apagar/financeiro/{{ f['id'] }}" onclick="return confirm('Eliminar?');" class="text-rose-600 font-bold">X</a></td>
                                </tr>
                                {% endfor %}
                            </tbody>
                        </table>
                    </div>
                </div>
            </div>
        </section>

        <!-- ABA CASAMENTOS -->
        <section id="aba-casamentos" class="tab-content space-y-6">
            <div class="grid grid-cols-1 lg:grid-cols-3 gap-6">
                <div class="bg-white p-5 rounded-xl shadow-sm border border-slate-200 h-fit">
                    <h3 class="text-sm font-bold uppercase text-indigo-900 mb-3 pb-2 border-b">Registo Matrimonial</h3>
                    <form action="/casamentos/novo" method="POST" class="space-y-3 text-xs">
                        <input type="text" name="noivo" required placeholder="Noivo *" class="w-full p-2 border rounded-lg">
                        <input type="text" name="noiva" required placeholder="Noiva *" class="w-full p-2 border rounded-lg">
                        <input type="date" name="data_casamento" required class="w-full p-2 border rounded-lg">
                        <input type="text" name="pastor_oficiante" placeholder="Pastor Oficiante" class="w-full p-2 border rounded-lg">
                        <button type="submit" class="w-full bg-indigo-800 text-white font-bold py-2.5 rounded-lg">Gravar Casamento</button>
                    </form>
                </div>
                <div class="lg:col-span-2 bg-white p-5 rounded-xl shadow-sm border border-slate-200">
                    <h3 class="text-sm font-bold uppercase mb-3 pb-2 border-b">Livro de Casamentos</h3>
                    <div class="overflow-x-auto">
                        <table class="w-full text-left text-xs">
                            <thead class="bg-slate-50 border-b"><tr><th class="p-2">Data</th><th class="p-2">Noivos</th><th class="p-2">Oficiante</th><th class="p-2 text-center">Ação</th></tr></thead>
                            <tbody class="divide-y divide-slate-100">
                                {% for c in todos_casamentos %}
                                <tr><td class="p-2">{{ c['data_casamento'] }}</td><td class="p-2 font-bold">{{ c['noivo'] }} & {{ c['noiva'] }}</td><td class="p-2">{{ c['pastor_oficiante'] or '-' }}</td><td class="p-2 text-center"><a href="/apagar/casamento/{{ c['id'] }}" onclick="return confirm('Eliminar?');" class="text-rose-600 font-bold">X</a></td></tr>
                                {% endfor %}
                            </tbody>
                        </table>
                    </div>
                </div>
            </div>
        </section>

        <!-- ABA OBITOS -->
        <section id="aba-mortes" class="tab-content space-y-6">
            <div class="grid grid-cols-1 lg:grid-cols-3 gap-6">
                <div class="bg-white p-5 rounded-xl shadow-sm border border-slate-200 h-fit">
                    <h3 class="text-sm font-bold uppercase mb-3 pb-2 border-b">Registo de Óbito</h3>
                    <form action="/mortes/novo" method="POST" class="space-y-3 text-xs">
                        <input type="text" name="nome_falecido" required placeholder="Nome Falecido *" class="w-full p-2 border rounded-lg">
                        <input type="date" name="data_falecimento" required class="w-full p-2 border rounded-lg">
                        <input type="text" name="observacoes" placeholder="Observações..." class="w-full p-2 border rounded-lg">
                        <button type="submit" class="w-full bg-slate-800 text-white font-bold py-2.5 rounded-lg">Gravar Óbito</button>
                    </form>
                </div>
                <div class="lg:col-span-2 bg-white p-5 rounded-xl shadow-sm border border-slate-200">
                    <h3 class="text-sm font-bold uppercase mb-3 pb-2 border-b">Livro de Óbitos</h3>
                    <div class="overflow-x-auto">
                        <table class="w-full text-left text-xs">
                            <thead class="bg-slate-50 border-b"><tr><th class="p-2">Data</th><th class="p-2">Nome</th><th class="p-2">Notas</th><th class="p-2 text-center">Ação</th></tr></thead>
                            <tbody class="divide-y divide-slate-100">
                                {% for m in todas_mortes %}
                                <tr><td class="p-2">{{ m['data_falecimento'] }}</td><td class="p-2 font-bold">{{ m['nome_falecido'] }}</td><td class="p-2">{{ m['observacoes'] or '-' }}</td><td class="p-2 text-center"><a href="/apagar/morte/{{ m['id'] }}" onclick="return confirm('Eliminar?');" class="text-rose-600 font-bold">X</a></td></tr>
                                {% endfor %}
                            </tbody>
                        </table>
                    </div>
                </div>
            </div>
        </section>

        <!-- ABA CONFIGURACOES -->
        <section id="aba-configuracoes" class="tab-content space-y-6">
            <div class="grid grid-cols-1 md:grid-cols-3 gap-6">
                <div class="bg-white p-5 rounded-xl shadow-sm border border-slate-200">
                    <h3 class="text-xs font-bold uppercase text-blue-900 mb-3 pb-2 border-b">1. Ministérios</h3>
                    <form action="/config/departamento/novo" method="POST" class="flex gap-2 mb-3">
                        <input type="text" name="nome" placeholder="Novo..." required class="text-xs p-2 border rounded-lg w-full">
                        <button type="submit" class="bg-blue-800 text-white px-3 font-bold rounded-lg">+</button>
                    </form>
                    <ul class="divide-y text-xs max-h-52 overflow-y-auto">
                        {% for d in lista_deptos %}<li class="py-2 flex justify-between"><span>{{ d['nome'] }}</span><a href="/config/departamento/apagar/{{ d['id'] }}" class="text-rose-600 font-bold">X</a></li>{% endfor %}
                    </ul>
                </div>
                <div class="bg-white p-5 rounded-xl shadow-sm border border-slate-200">
                    <h3 class="text-xs font-bold uppercase text-emerald-800 mb-3 pb-2 border-b">2. Categorias</h3>
                    <form action="/config/categoria/novo" method="POST" class="space-y-2 mb-3 text-xs">
                        <div class="flex gap-2">
                            <select name="tipo" class="p-2 border rounded-lg bg-white w-1/3"><option value="Entrada">Entrada</option><option value="Saída">Saída</option></select>
                            <input type="text" name="nome" placeholder="Nome..." required class="p-2 border rounded-lg w-2/3">
                        </div>
                        <button type="submit" class="w-full bg-emerald-600 text-white py-2 font-bold rounded-lg">+ Guardar</button>
                    </form>
                    <ul class="divide-y text-xs max-h-52 overflow-y-auto">
                        {% for c in lista_categorias %}<li class="py-2 flex justify-between"><span>{{ c['tipo'] }} - {{ c['nome'] }}</span><a href="/config/categoria/apagar/{{ c['id'] }}" class="text-rose-600 font-bold">X</a></li>{% endfor %}
                    </ul>
                </div>
                <div class="bg-white p-5 rounded-xl shadow-sm border border-slate-200">
                    <h3 class="text-xs font-bold uppercase text-amber-800 mb-3 pb-2 border-b">3. Zonas / Bairros</h3>
                    <form action="/config/zona/novo" method="POST" class="flex gap-2 mb-3">
                        <input type="text" name="nome" placeholder="Novo..." required class="text-xs p-2 border rounded-lg w-full">
                        <button type="submit" class="bg-amber-600 text-white px-3 font-bold rounded-lg">+</button>
                    </form>
                    <ul class="divide-y text-xs max-h-52 overflow-y-auto">
                        {% for z in lista_zonas %}<li class="py-2 flex justify-between"><span>{{ z['nome'] }}</span><a href="/config/zona/apagar/{{ z['id'] }}" class="text-rose-600 font-bold">X</a></li>{% endfor %}
                    </ul>
                </div>
            </div>
        </section>

        <!-- ABA USUARIOS -->
        <section id="aba-usuarios" class="tab-content space-y-6">
            <div class="grid grid-cols-1 md:grid-cols-3 gap-6">
                <div class="bg-white p-5 rounded-xl shadow-sm border border-slate-200 h-fit">
                    <h3 class="text-sm font-bold uppercase text-emerald-800 mb-3 pb-2 border-b">👤 Criar Utilizador</h3>
                    <form action="/usuarios/novo" method="POST" class="space-y-3 text-xs">
                        <input type="text" name="usuario" required placeholder="Utilizador *" class="w-full p-2 border rounded-lg">
                        <input type="password" name="senha" required placeholder="Palavra-passe *" class="w-full p-2 border rounded-lg">
                        <select name="cargo" required class="w-full p-2 border rounded-lg bg-white">
                            <option value="Pastor">Pastor</option>
                            <option value="Tesoureiro">Tesoureiro</option>
                            <option value="Secretário">Secretário</option>
                            <option value="Líder">Líder de Ministério</option>
                        </select>
                        <button type="submit" class="w-full bg-emerald-700 text-white font-bold py-2.5 rounded-lg">Criar Utilizador</button>
                    </form>
                </div>
                <div class="md:col-span-2 bg-white p-5 rounded-xl shadow-sm border border-slate-200">
                    <h3 class="text-sm font-bold uppercase mb-3 pb-2 border-b">Utilizadores Registados</h3>
                    <div class="overflow-x-auto">
                        <table class="w-full text-left text-xs">
                            <thead class="bg-slate-50 border-b"><tr><th class="p-2">Login</th><th class="p-2">Cargo</th><th class="p-2 text-center">Ação</th></tr></thead>
                            <tbody class="divide-y divide-slate-100">
                                {% for u in lista_usuarios %}
                                <tr>
                                    <td class="p-2 font-bold">{{ u['usuario'] }}</td>
                                    <td class="p-2">{{ u['cargo'] }}</td>
                                    <td class="p-2 text-center">
                                        {% if u['usuario'] != 'admin' %}<a href="/usuarios/apagar/{{ u['id'] }}" onclick="return confirm('Eliminar?');" class="text-rose-600 font-bold">X</a>
                                        {% else %}<span class="text-slate-400">Principal</span>{% endif %}
                                    </td>
                                </tr>
                                {% endfor %}
                            </tbody>
                        </table>
                    </div>
                </div>
            </div>
        </section>

    </main>

    <!-- Modal Cartão de Membro -->
    <div id="modalMembro" class="fixed inset-0 bg-slate-900/60 backdrop-blur-sm hidden items-center justify-center p-4 z-50">
        <div class="bg-white rounded-2xl max-w-lg w-full p-6 shadow-2xl relative border border-slate-200">
            <button onclick="fecharModalMembro()" class="absolute top-4 right-4 text-slate-400 hover:text-slate-600 font-bold text-lg">&times;</button>
            <div id="conteudoModalMembro" class="space-y-4"></div>
        </div>
    </div>

    <script>
        const membrosDados = {{ membros_json | safe }};

        function verCartaoMembro(id) {
            const m = membrosDados.find(x => x.id === id);
            if (!m) return;

            const fotoHtml = m.foto_path 
                ? `<img src="${m.foto_path}" class="w-24 h-24 object-cover rounded-2xl mx-auto border-2 border-blue-900 shadow">`
                : `<div class="w-24 h-24 rounded-2xl bg-blue-100 text-blue-900 flex items-center justify-center text-2xl font-black mx-auto border-2 border-blue-200">${m.nome.substring(0,2).toUpperCase()}</div>`;

            document.getElementById('conteudoModalMembro').innerHTML = `
                <div class="text-center pb-3 border-b">
                    ${fotoHtml}
                    <h3 class="text-base font-bold text-slate-900 mt-2">${m.nome}</h3>
                    <span class="inline-block bg-blue-900 text-white text-[10px] font-bold px-2.5 py-0.5 rounded-full mt-1">${m.posicao_atual || 'Membro'}</span>
                    <span class="inline-block bg-indigo-50 text-indigo-700 text-[10px] font-bold px-2.5 py-0.5 rounded-full mt-1">${m.departamento || 'Geral'}</span>
                </div>
                <div class="grid grid-cols-2 gap-2 text-xs text-slate-600">
                    <div><strong>Segmento:</strong> ${m.faixa_etaria || '-'}</div>
                    <div><strong>Contacto:</strong> ${m.telefone || '-'}</div>
                    <div><strong>Documento:</strong> ${m.tipo_documento || 'BI'}: ${m.numero_documento || '-'}</div>
                    <div><strong>Naturalidade:</strong> ${m.naturalidade || '-'}</div>
                    <div><strong>Bairro/Zona:</strong> ${m.bairro || '-'}</div>
                    <div><strong>Ano Conversão:</strong> ${m.ano_conversao || '-'}</div>
                    <div><strong>Data Batismo:</strong> ${m.data_batismo || '-'}</div>
                    <div><strong>Filiação:</strong> ${m.filiacao || '-'}</div>
                </div>
                <div class="p-2.5 bg-slate-50 rounded-lg text-xs border">
                    <strong class="text-blue-950 block mb-1">Histórico de Progressões Eclesiásticas:</strong>
                    <p class="text-slate-700 text-[11px] whitespace-pre-line">${m.progressoes || 'Nenhuma alteração de cargo registada.'}</p>
                </div>
                <button onclick="window.print()" class="w-full bg-slate-800 text-white font-bold py-2 rounded-lg text-xs hover:bg-slate-900 transition">🖨️ Imprimir Ficha</button>
            `;
            document.getElementById('modalMembro').classList.remove('hidden');
            document.getElementById('modalMembro').classList.add('flex');
        }

        function fecharModalMembro() {
            document.getElementById('modalMembro').classList.add('hidden');
            document.getElementById('modalMembro').classList.remove('flex');
        }

        const todasCategorias = {{ categorias_json | safe }};

        function atualizarCategoriasPorTipo() {
            const tipo = document.getElementById('selectTipoTransacao').value;
            const selectCat = document.getElementById('selectCategoriaTransacao');
            selectCat.innerHTML = '';
            const filtradas = todasCategorias.filter(c => c.tipo === tipo);
            filtradas.forEach(c => {
                const opt = document.createElement('option');
                opt.value = c.nome;
                opt.textContent = c.nome;
                selectCat.appendChild(opt);
            });
        }
        document.addEventListener('DOMContentLoaded', atualizarCategoriasPorTipo);

        function trocarAba(abaId) {
            document.querySelectorAll('.tab-content').forEach(el => el.classList.remove('active'));
            document.querySelectorAll('.tab-btn').forEach(btn => {
                btn.className = 'tab-btn px-4 py-2 font-bold rounded-t-lg text-blue-200 hover:bg-blue-800/50';
            });
            document.getElementById('aba-' + abaId).classList.add('active');
            const activeBtn = document.getElementById('btn-' + abaId);
            if (activeBtn) activeBtn.className = 'tab-btn px-4 py-2 font-bold rounded-t-lg bg-white text-blue-900 shadow';
        }

        function filtrarTabela(inputId, tabelaId) {
            const filtro = document.getElementById(inputId).value.toLowerCase();
            const tr = document.getElementById(tabelaId).getElementsByTagName('tr');
            for (let i = 1; i < tr.length; i++) {
                tr[i].style.display = tr[i].innerText.toLowerCase().includes(filtro) ? '' : 'none';
            }
        }

        new Chart(document.getElementById('graficoEvolucaoMensal').getContext('2d'), {
            type: 'line',
            data: {
                labels: ['Jan', 'Fev', 'Mar', 'Abr', 'Mai', 'Jun', 'Jul', 'Ago', 'Set', 'Out', 'Nov', 'Dez'],
                datasets: [
                    { label: 'Entradas', data: {{ evolucao_entradas | safe }}, borderColor: '#059669', fill: false },
                    { label: 'Saídas', data: {{ evolucao_saidas | safe }}, borderColor: '#e11d48', fill: false }
                ]
            },
            options: { responsive: true, maintainAspectRatio: false }
        });

        new Chart(document.getElementById('graficoDeptos').getContext('2d'), {
            type: 'bar',
            data: {
                labels: {{ depto_labels | safe }},
                datasets: [{ label: 'Membros', data: {{ depto_valores | safe }}, backgroundColor: '#1d4ed8' }]
            },
            options: { responsive: true, maintainAspectRatio: false }
        });
    </script>
</body>
</html>
"""
with open(os.path.join("templates", "dashboard.html"), "w", encoding="utf-8") as f:
    f.write(html_code)
print("✓ 2/4: templates/dashboard.html criado com sucesso!")

# 4. app.py completo com todas as novas colunas e migração automática
app_code = """from flask import Flask, render_template, request, redirect, url_for, session, send_file
import sqlite3
from datetime import datetime
import json
import io
import os
from werkzeug.utils import secure_filename
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

app = Flask(__name__)
app.secret_key = "iead_chicuque_chave_super_segura_2026"
DB_NAME = "gestao_chicuque.db"
UPLOAD_FOLDER = os.path.join("static", "uploads")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

def init_db():
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    
    # Utilizadores
    c.execute('''CREATE TABLE IF NOT EXISTS usuarios (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        usuario TEXT UNIQUE NOT NULL,
        senha TEXT NOT NULL,
        cargo TEXT NOT NULL
    )''')
    if c.execute("SELECT COUNT(*) FROM usuarios").fetchone()[0] == 0:
        c.execute("INSERT INTO usuarios (usuario, senha, cargo) VALUES (?, ?, ?)",
                  ('admin', 'chicuque123', 'Pastor Presidente'))

    # Membros completo
    c.execute('''CREATE TABLE IF NOT EXISTS membros (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nome TEXT NOT NULL,
        telefone TEXT,
        genero TEXT,
        data_nascimento TEXT,
        faixa_etaria TEXT,
        naturalidade TEXT,
        bairro TEXT,
        filiacao TEXT,
        tipo_documento TEXT,
        numero_documento TEXT,
        ano_conversao INTEGER,
        data_batismo TEXT,
        posicao_atual TEXT,
        progressoes TEXT,
        departamento TEXT,
        foto_path TEXT,
        observacoes TEXT,
        data_registo TEXT
    )''')

    # Migração das colunas
    c.execute("PRAGMA table_info(membros)")
    cols = [col[1] for col in c.fetchall()]
    novas = [
        ('genero', 'TEXT'), ('data_nascimento', 'TEXT'), ('faixa_etaria', 'TEXT'),
        ('naturalidade', 'TEXT'), ('filiacao', 'TEXT'), ('tipo_documento', 'TEXT'),
        ('numero_documento', 'TEXT'), ('ano_conversao', 'INTEGER'), ('posicao_atual', 'TEXT'),
        ('progressoes', 'TEXT'), ('foto_path', 'TEXT')
    ]
    for n, t in novas:
        if n not in cols:
            c.execute(f"ALTER TABLE membros ADD COLUMN {n} {t}")

    # Financeiro
    c.execute('''CREATE TABLE IF NOT EXISTS financeiro (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        tipo TEXT CHECK(tipo IN ('Entrada', 'Saída')),
        local_movimento TEXT CHECK(local_movimento IN ('Caixa', 'Banco')) DEFAULT 'Caixa',
        departamento TEXT DEFAULT 'Geral',
        categoria TEXT NOT NULL,
        valor REAL NOT NULL,
        data_movimento TEXT NOT NULL,
        dia INTEGER,
        mes INTEGER,
        ano INTEGER,
        data_registo TEXT NOT NULL,
        membro_id INTEGER,
        descricao TEXT,
        FOREIGN KEY (membro_id) REFERENCES membros (id)
    )''')

    # Transferências
    c.execute('''CREATE TABLE IF NOT EXISTS transferencias (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        data_movimento TEXT NOT NULL,
        origem_local TEXT NOT NULL,
        origem_depto TEXT NOT NULL,
        destino_local TEXT NOT NULL,
        destino_depto TEXT NOT NULL,
        valor REAL NOT NULL,
        motivo TEXT NOT NULL,
        data_registo TEXT NOT NULL
    )''')

    # Casamentos e Mortes
    c.execute('''CREATE TABLE IF NOT EXISTS casamentos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        noivo TEXT NOT NULL,
        noiva TEXT NOT NULL,
        data_casamento TEXT NOT NULL,
        pastor_oficiante TEXT,
        data_registo TEXT
    )''')

    c.execute('''CREATE TABLE IF NOT EXISTS mortes (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nome_falecido TEXT NOT NULL,
        data_falecimento TEXT NOT NULL,
        observacoes TEXT,
        data_registo TEXT
    )''')

    # Configurações
    c.execute('''CREATE TABLE IF NOT EXISTS departamentos_lista (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nome TEXT UNIQUE NOT NULL
    )''')
    c.execute('''CREATE TABLE IF NOT EXISTS categorias_financeiras (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        tipo TEXT CHECK(tipo IN ('Entrada', 'Saída')) NOT NULL,
        nome TEXT NOT NULL
    )''')
    c.execute('''CREATE TABLE IF NOT EXISTS zonas_lista (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nome TEXT UNIQUE NOT NULL
    )''')

    deptos = [
        ('Activista',), ('Juventude',), ('Mulher (Senhoras)',), 
        ('Boa Esperança (Crianças)',), ('Homens / Obreiros',), 
        ('Louvor / Música',), ('Ação Social',), ('Construção',)
    ]
    for d in deptos:
        c.execute("INSERT OR IGNORE INTO departamentos_lista (nome) VALUES (?)", d)

    if c.execute("SELECT COUNT(*) FROM categorias_financeiras").fetchone()[0] == 0:
        padroes = [
            ('Entrada', 'Dízimo'), ('Entrada', 'Oferta'), ('Entrada', 'Doação / Voto'), ('Entrada', 'Campanha de Construção'),
            ('Saída', 'Manutenção do Templo'), ('Saída', 'Ação Social / Ajudas'), ('Saída', 'Combustível / Transporte'),
            ('Saída', 'Água e Eletricidade'), ('Saída', 'Material de Culto')
        ]
        c.executemany("INSERT INTO categorias_financeiras (tipo, nome) VALUES (?, ?)", padroes)

    if c.execute("SELECT COUNT(*) FROM zonas_lista").fetchone()[0] == 0:
        c.executemany("INSERT INTO zonas_lista (nome) VALUES (?)",
                      [('Chicuque Sede',), ('Maxixe Cidade',), ('Nhacoongo',), ('Conguiana',), ('Bairro 1',)])

    conn.commit()
    conn.close()

init_db()

def get_db():
    conn = sqlite3.connect(DB_NAME)
    conn.row_factory = sqlite3.Row
    return conn

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        usuario = request.form.get('usuario')
        senha = request.form.get('senha')
        conn = get_db()
        user = conn.execute("SELECT * FROM usuarios WHERE usuario = ? AND senha = ?", (usuario, senha)).fetchone()
        conn.close()
        if user:
            session['usuario'] = user['usuario']
            session['cargo'] = user['cargo']
            return redirect(url_for('dashboard'))
        return render_template('login.html', erro="Credenciais incorretas.")
    return render_template('login.html', erro=None)

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.route('/')
def dashboard():
    if 'usuario' not in session:
        return redirect(url_for('login'))
        
    conn = get_db()
    total_membros = conn.execute("SELECT COUNT(*) FROM membros").fetchone()[0]
    total_casamentos = conn.execute("SELECT COUNT(*) FROM casamentos").fetchone()[0]
    total_mortes = conn.execute("SELECT COUNT(*) FROM mortes").fetchone()[0]

    entrada_caixa = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Entrada' AND (local_movimento = 'Caixa' OR local_movimento IS NULL)").fetchone()[0] or 0.0
    saida_caixa = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Saída' AND (local_movimento = 'Caixa' OR local_movimento IS NULL)").fetchone()[0] or 0.0
    saldo_caixa = entrada_caixa - saida_caixa

    entrada_banco = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Entrada' AND local_movimento = 'Banco'").fetchone()[0] or 0.0
    saida_banco = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Saída' AND local_movimento = 'Banco'").fetchone()[0] or 0.0
    saldo_banco = entrada_banco - saida_banco

    saldo_total = (entrada_caixa + entrada_banco) - (saida_caixa + saida_banco)

    ano_atual = datetime.now().year
    evolucao_entradas, evolucao_saidas = [], []
    for m in range(1, 13):
        e_mes = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Entrada' AND mes = ? AND ano = ?", (m, ano_atual)).fetchone()[0] or 0.0
        s_mes = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Saída' AND mes = ? AND ano = ?", (m, ano_atual)).fetchone()[0] or 0.0
        evolucao_entradas.append(e_mes)
        evolucao_saidas.append(s_mes)

    todos_membros = conn.execute("SELECT * FROM membros ORDER BY id DESC").fetchall()
    todas_financas = conn.execute("SELECT * FROM financeiro ORDER BY id DESC").fetchall()
    todos_casamentos = conn.execute("SELECT * FROM casamentos ORDER BY id DESC").fetchall()
    todas_mortes = conn.execute("SELECT * FROM mortes ORDER BY id DESC").fetchall()
    lista_usuarios = conn.execute("SELECT id, usuario, cargo FROM usuarios ORDER BY id ASC").fetchall()
    ultimas_transferencias = conn.execute("SELECT * FROM transferencias ORDER BY id DESC LIMIT 20").fetchall()

    lista_deptos = conn.execute("SELECT * FROM departamentos_lista ORDER BY nome ASC").fetchall()
    lista_categorias = conn.execute("SELECT * FROM categorias_financeiras ORDER BY tipo, nome ASC").fetchall()
    lista_zonas = conn.execute("SELECT * FROM zonas_lista ORDER BY nome ASC").fetchall()

    categorias_json = json.dumps([{'tipo': c['tipo'], 'nome': c['nome']} for c in lista_categorias])
    membros_json = json.dumps([dict(m) for m in todos_membros])

    dept_rows = conn.execute("SELECT COALESCE(NULLIF(TRIM(departamento), ''), 'Geral') as dep, COUNT(*) as qtd FROM membros GROUP BY dep").fetchall()
    depto_labels = [r['dep'] for r in dept_rows] if dept_rows else ['Geral']
    depto_valores = [r['qtd'] for r in dept_rows] if dept_rows else [0]

    conn.close()

    return render_template('dashboard.html',
                           total_membros=total_membros,
                           total_casamentos=total_casamentos,
                           total_mortes=total_mortes,
                           saldo_caixa=saldo_caixa,
                           saldo_banco=saldo_banco,
                           saldo_total=saldo_total,
                           todos_membros=todos_membros,
                           todas_financas=todas_financas,
                           todos_casamentos=todos_casamentos,
                           todas_mortes=todas_mortes,
                           lista_deptos=lista_deptos,
                           lista_categorias=lista_categorias,
                           lista_zonas=lista_zonas,
                           lista_usuarios=lista_usuarios,
                           ultimas_transferencias=ultimas_transferencias,
                           categorias_json=categorias_json,
                           membros_json=membros_json,
                           evolucao_entradas=json.dumps(evolucao_entradas),
                           evolucao_saidas=json.dumps(evolucao_saidas),
                           depto_labels=json.dumps(depto_labels),
                           depto_valores=json.dumps(depto_valores))

@app.route('/membros/novo', methods=['POST'])
def novo_membro():
    if 'usuario' not in session: return redirect(url_for('login'))
    
    foto_path = ""
    if 'foto' in request.files:
        foto = request.files['foto']
        if foto and foto.filename:
            ext = foto.filename.rsplit('.', 1)[-1].lower()
            if ext in ['png', 'jpg', 'jpeg', 'webp']:
                nome_foto = f"membro_{int(datetime.now().timestamp())}.{ext}"
                salvar_em = os.path.join(app.config['UPLOAD_FOLDER'], nome_foto)
                foto.save(salvar_em)
                foto_path = f"/static/uploads/{nome_foto}"

    ano_conv = request.form.get('ano_conversao')
    ano_conv_val = int(ano_conv) if ano_conv and ano_conv.isdigit() else None

    conn = get_db()
    conn.execute('''
        INSERT INTO membros (
            nome, telefone, genero, data_nascimento, faixa_etaria,
            naturalidade, bairro, filiacao, tipo_documento, numero_documento,
            ano_conversao, data_batismo, posicao_atual, progressoes,
            departamento, foto_path, observacoes, data_registo
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', (
        request.form['nome'].strip(),
        request.form.get('telefone', '').strip(),
        request.form.get('genero', 'Masculino'),
        request.form.get('data_nascimento', ''),
        request.form.get('faixa_etaria', 'Adulto'),
        request.form.get('naturalidade', '').strip(),
        request.form.get('bairro', ''),
        request.form.get('filiacao', '').strip(),
        request.form.get('tipo_documento', 'BI'),
        request.form.get('numero_documento', '').strip(),
        ano_conv_val,
        request.form.get('data_batismo', ''),
        request.form.get('posicao_atual', 'Membro em Comunhão'),
        request.form.get('progressoes', '').strip(),
        request.form.get('departamento', 'Geral'),
        foto_path,
        request.form.get('observacoes', '').strip(),
        datetime.now().strftime("%d/%m/%Y")
    ))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/financeiro/transferir', methods=['POST'])
def transferir_fundos():
    if 'usuario' not in session: return redirect(url_for('login'))
    data_raw = request.form.get('data_movimento') or datetime.now().strftime("%Y-%m-%d")
    dt_obj = datetime.strptime(data_raw, "%Y-%m-%d")
    origem_local, origem_depto = request.form['origem_local'], request.form['origem_depto']
    destino_local, destino_depto = request.form['destino_local'], request.form['destino_depto']
    valor = float(request.form['valor'])
    motivo = request.form['motivo']
    agora = datetime.now().strftime("%d/%m/%Y %H:%M")

    conn = get_db()
    conn.execute('''INSERT INTO financeiro (tipo, local_movimento, departamento, categoria, valor, data_movimento, dia, mes, ano, data_registo, descricao)
                    VALUES ('Saída', ?, ?, 'Transferência de Fundos', ?, ?, ?, ?, ?, ?, ?)''',
                 (origem_local, origem_depto, valor, dt_obj.strftime("%d/%m/%Y"), dt_obj.day, dt_obj.month, dt_obj.year, agora, f"Transf. p/ [{destino_local} - {destino_depto}]: {motivo}"))
    conn.execute('''INSERT INTO financeiro (tipo, local_movimento, departamento, categoria, valor, data_movimento, dia, mes, ano, data_registo, descricao)
                    VALUES ('Entrada', ?, ?, 'Transferência de Fundos', ?, ?, ?, ?, ?, ?, ?)''',
                 (destino_local, destino_depto, valor, dt_obj.strftime("%d/%m/%Y"), dt_obj.day, dt_obj.month, dt_obj.year, agora, f"Transf. de [{origem_local} - {origem_depto}]: {motivo}"))
    conn.execute('''INSERT INTO transferencias (data_movimento, origem_local, origem_depto, destino_local, destino_depto, valor, motivo, data_registo)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)''',
                 (dt_obj.strftime("%d/%m/%Y"), origem_local, origem_depto, destino_local, destino_depto, valor, motivo, agora))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/financeiro/novo', methods=['POST'])
def novo_financeiro():
    if 'usuario' not in session: return redirect(url_for('login'))
    data_raw = request.form.get('data_movimento') or datetime.now().strftime("%Y-%m-%d")
    dt_obj = datetime.strptime(data_raw, "%Y-%m-%d")
    conn = get_db()
    conn.execute('''INSERT INTO financeiro (tipo, local_movimento, departamento, categoria, valor, data_movimento, dia, mes, ano, data_registo, membro_id, descricao)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
                 (request.form['tipo'], request.form.get('local_movimento', 'Caixa'), request.form.get('departamento', 'Geral'),
                  request.form['categoria'], float(request.form['valor']), dt_obj.strftime("%d/%m/%Y"), dt_obj.day, dt_obj.month, dt_obj.year,
                  datetime.now().strftime("%d/%m/%Y %H:%M"), request.form.get('membro_id') or None, request.form['descricao']))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/usuarios/novo', methods=['POST'])
def novo_usuario():
    if 'usuario' not in session: return redirect(url_for('login'))
    user, senha, cargo = request.form['usuario'].strip(), request.form['senha'].strip(), request.form['cargo'].strip()
    if user and senha:
        conn = get_db()
        try:
            conn.execute("INSERT INTO usuarios (usuario, senha, cargo) VALUES (?, ?, ?)", (user, senha, cargo))
            conn.commit()
        except sqlite3.IntegrityError: pass
        conn.close()
    return redirect(url_for('dashboard'))

@app.route('/usuarios/apagar/<int:id>')
def apagar_usuario(id):
    if 'usuario' not in session: return redirect(url_for('login'))
    conn = get_db()
    conn.execute("DELETE FROM usuarios WHERE id = ? AND id != 1", (id,))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/casamentos/novo', methods=['POST'])
def novo_casamento():
    if 'usuario' not in session: return redirect(url_for('login'))
    conn = get_db()
    conn.execute('''INSERT INTO casamentos (noivo, noiva, data_casamento, pastor_oficiante, data_registo)
                    VALUES (?, ?, ?, ?, ?)''',
                 (request.form['noivo'], request.form['noiva'], request.form['data_casamento'],
                  request.form.get('pastor_oficiante', ''), datetime.now().strftime("%d/%m/%Y")))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/mortes/novo', methods=['POST'])
def novo_morte():
    if 'usuario' not in session: return redirect(url_for('login'))
    conn = get_db()
    conn.execute('''INSERT INTO mortes (nome_falecido, data_falecimento, observacoes, data_registo)
                    VALUES (?, ?, ?, ?)''',
                 (request.form['nome_falecido'], request.form['data_falecimento'],
                  request.form.get('observacoes', ''), datetime.now().strftime("%d/%m/%Y")))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/apagar/<tabela>/<int:id>')
def apagar_registo(tabela, id):
    if 'usuario' not in session: return redirect(url_for('login'))
    mapa = {'membro': ('membros', 'id'), 'financeiro': ('financeiro', 'id'), 'casamento': ('casamentos', 'id'), 'morte': ('mortes', 'id')}
    if tabela in mapa:
        tab, col = mapa[tabela]
        conn = get_db()
        conn.execute(f"DELETE FROM {tab} WHERE {col} = ?", (id,))
        conn.commit()
        conn.close()
    return redirect(url_for('dashboard'))

@app.route('/config/departamento/novo', methods=['POST'])
def novo_depto():
    if 'usuario' not in session: return redirect(url_for('login'))
    nome = request.form.get('nome', '').strip()
    if nome:
        conn = get_db()
        try:
            conn.execute("INSERT INTO departamentos_lista (nome) VALUES (?)", (nome,))
            conn.commit()
        except sqlite3.IntegrityError: pass
        conn.close()
    return redirect(url_for('dashboard'))

@app.route('/config/departamento/apagar/<int:id>')
def apagar_depto(id):
    if 'usuario' not in session: return redirect(url_for('login'))
    conn = get_db()
    conn.execute("DELETE FROM departamentos_lista WHERE id = ?", (id,))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/config/categoria/novo', methods=['POST'])
def nova_categoria():
    if 'usuario' not in session: return redirect(url_for('login'))
    tipo, nome = request.form.get('tipo'), request.form.get('nome', '').strip()
    if tipo and nome:
        conn = get_db()
        conn.execute("INSERT INTO categorias_financeiras (tipo, nome) VALUES (?, ?)", (tipo, nome))
        conn.commit()
        conn.close()
    return redirect(url_for('dashboard'))

@app.route('/config/categoria/apagar/<int:id>')
def apagar_categoria(id):
    if 'usuario' not in session: return redirect(url_for('login'))
    conn = get_db()
    conn.execute("DELETE FROM categorias_financeiras WHERE id = ?", (id,))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/config/zona/novo', methods=['POST'])
def nova_zona():
    if 'usuario' not in session: return redirect(url_for('login'))
    nome = request.form.get('nome', '').strip()
    if nome:
        conn = get_db()
        try:
            conn.execute("INSERT INTO zonas_lista (nome) VALUES (?)", (nome,))
            conn.commit()
        except sqlite3.IntegrityError: pass
        conn.close()
    return redirect(url_for('dashboard'))

@app.route('/config/zona/apagar/<int:id>')
def apagar_zona(id):
    if 'usuario' not in session: return redirect(url_for('login'))
    conn = get_db()
    conn.execute("DELETE FROM zonas_lista WHERE id = ?", (id,))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/exportar/financeiro')
def exportar_financeiro():
    if 'usuario' not in session: return redirect(url_for('login'))
    conn = get_db()
    rows = conn.execute("SELECT data_movimento, tipo, local_movimento, departamento, categoria, descricao, valor FROM financeiro ORDER BY id DESC").fetchall()
    conn.close()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tesouraria Chicuque"
    ws.append(["Data Movimento", "Tipo", "Conta", "Departamento", "Categoria", "Descrição", "Valor (MT)"])
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    for col in range(1, 8):
        c = ws.cell(row=1, column=col)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append([r['data_movimento'], r['tipo'], r['local_movimento'], r['departamento'], r['categoria'], r['descricao'], r['valor']])
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"Tesouraria_IEAD_Chicuque_{datetime.now().strftime('%Y%m%d')}.xlsx")

@app.route('/exportar/membros')
def exportar_membros():
    if 'usuario' not in session: return redirect(url_for('login'))
    conn = get_db()
    rows = conn.execute('''
        SELECT nome, telefone, genero, data_nascimento, faixa_etaria, naturalidade, bairro, 
               filiacao, tipo_documento, numero_documento, ano_conversao, data_batismo, 
               posicao_atual, progressoes, departamento, observacoes
        FROM membros ORDER BY nome ASC
    ''').fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Membresia IEAD Chicuque"

    headers = [
        "Nome Completo", "Contacto", "Género", "Data Nasc.", "Segmento", "Naturalidade", "Bairro",
        "Filiação", "Tipo Doc", "Nº Documento", "Ano Conv.", "Data Batismo",
        "Posição Atual", "Progressões", "Departamento", "Observações"
    ]
    ws.append(headers)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center")

    for r in rows:
        ws.append([r[k] for k in r.keys()])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 13)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"Membros_Completos_IEAD_{datetime.now().strftime('%Y%m%d')}.xlsx")

if __name__ == '__main__':
    port = int(os.environ.get("PORT", 5000))
    app.run(host='0.0.0.0', port=port)
"""
with open("app.py", "w", encoding="utf-8") as f:
    f.write(app_code)
print("✓ 3/4: app.py atualizado com sucesso!")

# 5. Procfile
with open("Procfile", "w", encoding="utf-8") as f:
    f.write("web: gunicorn app:app\n")
print("✓ 4/4: Procfile garantido!")