import os

# -------------------------------------------------------------
# 1. ATUALIZAÇÃO DO BACKEND COM AS CONSULTAS ANALÍTICAS
# -------------------------------------------------------------
app_code = """from flask import Flask, render_template, request, redirect, url_for, session, send_file
import sqlite3
from datetime import datetime
import json
import io
import os
from werkzeug.utils import secure_filename
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

app = Flask(__name__)
app.secret_key = "iead_chicuque_chave_super_segura_2026"
DB_NAME = "gestao_chicuque.db"
UPLOAD_FOLDER = os.path.join("static", "uploads")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

def init_db():
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    
    c.execute('''CREATE TABLE IF NOT EXISTS usuarios (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        usuario TEXT UNIQUE NOT NULL,
        senha TEXT NOT NULL,
        cargo TEXT NOT NULL
    )''')
    if c.execute("SELECT COUNT(*) FROM usuarios").fetchone()[0] == 0:
        c.execute("INSERT INTO usuarios (usuario, senha, cargo) VALUES (?, ?, ?)",
                  ('admin', 'chicuque123', 'Pastor Presidente'))

    c.execute('''CREATE TABLE IF NOT EXISTS membros (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nome TEXT NOT NULL,
        telefone TEXT,
        genero TEXT,
        data_nascimento TEXT,
        faixa_etaria TEXT,
        naturalidade TEXT,
        bairro TEXT,
        filiacao TEXT,
        tipo_documento TEXT,
        numero_documento TEXT,
        ano_conversao INTEGER,
        data_batismo TEXT,
        posicao_atual TEXT,
        progressoes TEXT,
        departamento TEXT,
        foto_path TEXT,
        observacoes TEXT,
        data_registo TEXT
    )''')

    c.execute("PRAGMA table_info(membros)")
    cols = [col[1] for col in c.fetchall()]
    novas = [
        ('genero', 'TEXT'), ('data_nascimento', 'TEXT'), ('faixa_etaria', 'TEXT'),
        ('naturalidade', 'TEXT'), ('filiacao', 'TEXT'), ('tipo_documento', 'TEXT'),
        ('numero_documento', 'TEXT'), ('ano_conversao', 'INTEGER'), ('posicao_atual', 'TEXT'),
        ('progressoes', 'TEXT'), ('foto_path', 'TEXT')
    ]
    for n, t in novas:
        if n not in cols:
            c.execute(f"ALTER TABLE membros ADD COLUMN {n} {t}")

    c.execute('''CREATE TABLE IF NOT EXISTS financeiro (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        tipo TEXT CHECK(tipo IN ('Entrada', 'Saída')),
        local_movimento TEXT CHECK(local_movimento IN ('Caixa', 'Banco')) DEFAULT 'Caixa',
        departamento TEXT DEFAULT 'Geral',
        categoria TEXT NOT NULL,
        valor REAL NOT NULL,
        data_movimento TEXT NOT NULL,
        dia INTEGER,
        mes INTEGER,
        ano INTEGER,
        data_registo TEXT NOT NULL,
        membro_id INTEGER,
        descricao TEXT,
        FOREIGN KEY (membro_id) REFERENCES membros (id)
    )''')

    c.execute('''CREATE TABLE IF NOT EXISTS transferencias (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        data_movimento TEXT NOT NULL,
        origem_local TEXT NOT NULL,
        origem_depto TEXT NOT NULL,
        destino_local TEXT NOT NULL,
        destino_depto TEXT NOT NULL,
        valor REAL NOT NULL,
        motivo TEXT NOT NULL,
        data_registo TEXT NOT NULL
    )''')

    c.execute('''CREATE TABLE IF NOT EXISTS casamentos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        noivo TEXT NOT NULL,
        noiva TEXT NOT NULL,
        data_casamento TEXT NOT NULL,
        pastor_oficiante TEXT,
        data_registo TEXT
    )''')

    c.execute('''CREATE TABLE IF NOT EXISTS mortes (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nome_falecido TEXT NOT NULL,
        data_falecimento TEXT NOT NULL,
        observacoes TEXT,
        data_registo TEXT
    )''')

    c.execute('''CREATE TABLE IF NOT EXISTS departamentos_lista (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nome TEXT UNIQUE NOT NULL
    )''')
    c.execute('''CREATE TABLE IF NOT EXISTS categorias_financeiras (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        tipo TEXT CHECK(tipo IN ('Entrada', 'Saída')) NOT NULL,
        nome TEXT NOT NULL
    )''')
    c.execute('''CREATE TABLE IF NOT EXISTS zonas_lista (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nome TEXT UNIQUE NOT NULL
    )''')

    deptos = [
        ('Activista',), ('Juventude',), ('Mulher (Senhoras)',), 
        ('Boa Esperança (Crianças)',), ('Homens / Obreiros',), 
        ('Louvor / Música',), ('Ação Social',), ('Construção',)
    ]
    for d in deptos:
        c.execute("INSERT OR IGNORE INTO departamentos_lista (nome) VALUES (?)", d)

    if c.execute("SELECT COUNT(*) FROM categorias_financeiras").fetchone()[0] == 0:
        padroes = [
            ('Entrada', 'Dízimo'), ('Entrada', 'Oferta'), ('Entrada', 'Doação / Voto'), ('Entrada', 'Campanha de Construção'),
            ('Saída', 'Manutenção do Templo'), ('Saída', 'Ação Social / Ajudas'), ('Saída', 'Combustível / Transporte'),
            ('Saída', 'Água e Eletricidade'), ('Saída', 'Material de Culto')
        ]
        c.executemany("INSERT INTO categorias_financeiras (tipo, nome) VALUES (?, ?)", padroes)

    if c.execute("SELECT COUNT(*) FROM zonas_lista").fetchone()[0] == 0:
        c.executemany("INSERT INTO zonas_lista (nome) VALUES (?)",
                      [('Chicuque Sede',), ('Maxixe Cidade',), ('Nhacoongo',), ('Conguiana',), ('Bairro 1',)])

    conn.commit()
    conn.close()

init_db()

def get_db():
    conn = sqlite3.connect(DB_NAME)
    conn.row_factory = sqlite3.Row
    return conn

def is_admin():
    cargo = session.get('cargo', '')
    return session.get('usuario') == 'admin' or 'Pastor' in cargo

def can_cadastro():
    cargo = session.get('cargo', '')
    return is_admin() or 'Secretário' in cargo or 'Líder' in cargo

def can_tesouraria():
    cargo = session.get('cargo', '')
    return is_admin() or 'Tesoureiro' in cargo

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        usuario = request.form.get('usuario')
        senha = request.form.get('senha')
        conn = get_db()
        user = conn.execute("SELECT * FROM usuarios WHERE usuario = ? AND senha = ?", (usuario, senha)).fetchone()
        conn.close()
        if user:
            session['usuario'] = user['usuario']
            session['cargo'] = user['cargo']
            return redirect(url_for('dashboard'))
        return render_template('login.html', erro="Credenciais incorretas.")
    return render_template('login.html', erro=None)

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.route('/')
def dashboard():
    if 'usuario' not in session:
        return redirect(url_for('login'))
        
    conn = get_db()
    total_membros = conn.execute("SELECT COUNT(*) FROM membros").fetchone()[0]
    total_casamentos = conn.execute("SELECT COUNT(*) FROM casamentos").fetchone()[0]
    total_mortes = conn.execute("SELECT COUNT(*) FROM mortes").fetchone()[0]

    entrada_caixa = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Entrada' AND (local_movimento = 'Caixa' OR local_movimento IS NULL)").fetchone()[0] or 0.0
    saida_caixa = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Saída' AND (local_movimento = 'Caixa' OR local_movimento IS NULL)").fetchone()[0] or 0.0
    saldo_caixa = entrada_caixa - saida_caixa

    entrada_banco = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Entrada' AND local_movimento = 'Banco'").fetchone()[0] or 0.0
    saida_banco = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Saída' AND local_movimento = 'Banco'").fetchone()[0] or 0.0
    saldo_banco = entrada_banco - saida_banco

    saldo_total = (entrada_caixa + entrada_banco) - (saida_caixa + saida_banco)

    # 1. Evolução Mensal Geral
    ano_atual = datetime.now().year
    evolucao_entradas, evolucao_saidas = [], []
    for m in range(1, 13):
        e_mes = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Entrada' AND mes = ? AND ano = ?", (m, ano_atual)).fetchone()[0] or 0.0
        s_mes = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Saída' AND mes = ? AND ano = ?", (m, ano_atual)).fetchone()[0] or 0.0
        evolucao_entradas.append(e_mes)
        evolucao_saidas.append(s_mes)

    # 2. Situação Financeira por Departamento (Entradas vs Saídas)
    deptos_lista = [d['nome'] for d in conn.execute("SELECT nome FROM departamentos_lista ORDER BY nome ASC").fetchall()]
    deptos_financeiro = ['Geral'] + [d for d in deptos_lista if d != 'Geral']
    depto_fin_labels, depto_fin_entradas, depto_fin_saidas = [], [], []
    for d in deptos_financeiro:
        ent = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Entrada' AND departamento = ?", (d,)).fetchone()[0] or 0.0
        sai = conn.execute("SELECT SUM(valor) FROM financeiro WHERE tipo = 'Saída' AND departamento = ?", (d,)).fetchone()[0] or 0.0
        if ent > 0 or sai > 0 or d in ['Geral', 'Activista', 'Juventude', 'Mulher (Senhoras)', 'Boa Esperança (Crianças)']:
            depto_fin_labels.append(d)
            depto_fin_entradas.append(ent)
            depto_fin_saidas.append(sai)

    # 3. Fontes de Fundos (Entradas por Categoria)
    fontes_rows = conn.execute('''
        SELECT categoria, SUM(valor) as total 
        FROM financeiro 
        WHERE tipo = 'Entrada' 
        GROUP BY categoria 
        ORDER BY total DESC
    ''').fetchall()
    fontes_labels = [r['categoria'] for r in fontes_rows] if fontes_rows else ['Sem Entradas']
    fontes_valores = [r['total'] for r in fontes_rows] if fontes_rows else [0]

    # 4. Destino de Saídas (Despesas por Categoria)
    saidas_rows = conn.execute('''
        SELECT categoria, SUM(valor) as total 
        FROM financeiro 
        WHERE tipo = 'Saída' 
        GROUP BY categoria 
        ORDER BY total DESC
    ''').fetchall()
    saidas_labels = [r['categoria'] for r in saidas_rows] if saidas_rows else ['Sem Saídas']
    saidas_valores = [r['total'] for r in saidas_rows] if saidas_rows else [0]

    # Membros por departamento
    dept_rows = conn.execute("SELECT COALESCE(NULLIF(TRIM(departamento), ''), 'Geral') as dep, COUNT(*) as qtd FROM membros GROUP BY dep").fetchall()
    depto_labels = [r['dep'] for r in dept_rows] if dept_rows else ['Geral']
    depto_valores = [r['qtd'] for r in dept_rows] if dept_rows else [0]

    todos_membros = conn.execute("SELECT * FROM membros ORDER BY id DESC").fetchall()
    todas_financas = conn.execute("SELECT * FROM financeiro ORDER BY id DESC").fetchall()
    todos_casamentos = conn.execute("SELECT * FROM casamentos ORDER BY id DESC").fetchall()
    todas_mortes = conn.execute("SELECT * FROM mortes ORDER BY id DESC").fetchall()
    lista_usuarios = conn.execute("SELECT id, usuario, cargo FROM usuarios ORDER BY id ASC").fetchall()
    ultimas_transferencias = conn.execute("SELECT * FROM transferencias ORDER BY id DESC LIMIT 20").fetchall()

    lista_deptos = conn.execute("SELECT * FROM departamentos_lista ORDER BY nome ASC").fetchall()
    lista_categorias = conn.execute("SELECT * FROM categorias_financeiras ORDER BY tipo, nome ASC").fetchall()
    lista_zonas = conn.execute("SELECT * FROM zonas_lista ORDER BY nome ASC").fetchall()

    categorias_json = json.dumps([{'tipo': c['tipo'], 'nome': c['nome']} for c in lista_categorias])
    membros_json = json.dumps([dict(m) for m in todos_membros])

    conn.close()

    return render_template('dashboard.html',
                           pode_cadastro=can_cadastro(),
                           pode_tesouraria=can_tesouraria(),
                           e_admin=is_admin(),
                           total_membros=total_membros,
                           total_casamentos=total_casamentos,
                           total_mortes=total_mortes,
                           saldo_caixa=saldo_caixa,
                           saldo_banco=saldo_banco,
                           saldo_total=saldo_total,
                           todos_membros=todos_membros,
                           todas_financas=todas_financas,
                           todos_casamentos=todos_casamentos,
                           todas_mortes=todas_mortes,
                           lista_deptos=lista_deptos,
                           lista_categorias=lista_categorias,
                           lista_zonas=lista_zonas,
                           lista_usuarios=lista_usuarios,
                           ultimas_transferencias=ultimas_transferencias,
                           categorias_json=categorias_json,
                           membros_json=membros_json,
                           evolucao_entradas=json.dumps(evolucao_entradas),
                           evolucao_saidas=json.dumps(evolucao_saidas),
                           depto_fin_labels=json.dumps(depto_fin_labels),
                           depto_fin_entradas=json.dumps(depto_fin_entradas),
                           depto_fin_saidas=json.dumps(depto_fin_saidas),
                           fontes_labels=json.dumps(fontes_labels),
                           fontes_valores=json.dumps(fontes_valores),
                           saidas_labels=json.dumps(saidas_labels),
                           saidas_valores=json.dumps(saidas_valores),
                           depto_labels=json.dumps(depto_labels),
                           depto_valores=json.dumps(depto_valores))

@app.route('/membros/novo', methods=['POST'])
def novo_membro():
    if not can_cadastro(): return redirect(url_for('dashboard'))
    foto_path = ""
    if 'foto' in request.files:
        foto = request.files['foto']
        if foto and foto.filename:
            ext = foto.filename.rsplit('.', 1)[-1].lower()
            if ext in ['png', 'jpg', 'jpeg', 'webp']:
                nome_foto = f"membro_{int(datetime.now().timestamp())}.{ext}"
                salvar_em = os.path.join(app.config['UPLOAD_FOLDER'], nome_foto)
                foto.save(salvar_em)
                foto_path = f"/static/uploads/{nome_foto}"

    ano_conv = request.form.get('ano_conversao')
    ano_conv_val = int(ano_conv) if ano_conv and ano_conv.isdigit() else None

    conn = get_db()
    conn.execute('''
        INSERT INTO membros (
            nome, telefone, genero, data_nascimento, faixa_etaria,
            naturalidade, bairro, filiacao, tipo_documento, numero_documento,
            ano_conversao, data_batismo, posicao_atual, progressoes,
            departamento, foto_path, observacoes, data_registo
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', (
        request.form['nome'].strip(),
        request.form.get('telefone', '').strip(),
        request.form.get('genero', 'Masculino'),
        request.form.get('data_nascimento', ''),
        request.form.get('faixa_etaria', 'Adulto'),
        request.form.get('naturalidade', '').strip(),
        request.form.get('bairro', ''),
        request.form.get('filiacao', '').strip(),
        request.form.get('tipo_documento', 'BI'),
        request.form.get('numero_documento', '').strip(),
        ano_conv_val,
        request.form.get('data_batismo', ''),
        request.form.get('posicao_atual', 'Membro em Comunhão'),
        request.form.get('progressoes', '').strip(),
        request.form.get('departamento', 'Geral'),
        foto_path,
        request.form.get('observacoes', '').strip(),
        datetime.now().strftime("%d/%m/%Y")
    ))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/casamentos/novo', methods=['POST'])
def novo_casamento():
    if not can_cadastro(): return redirect(url_for('dashboard'))
    conn = get_db()
    conn.execute('''INSERT INTO casamentos (noivo, noiva, data_casamento, pastor_oficiante, data_registo)
                    VALUES (?, ?, ?, ?, ?)''',
                 (request.form['noivo'], request.form['noiva'], request.form['data_casamento'],
                  request.form.get('pastor_oficiante', ''), datetime.now().strftime("%d/%m/%Y")))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/mortes/novo', methods=['POST'])
def novo_morte():
    if not can_cadastro(): return redirect(url_for('dashboard'))
    conn = get_db()
    conn.execute('''INSERT INTO mortes (nome_falecido, data_falecimento, observacoes, data_registo)
                    VALUES (?, ?, ?, ?)''',
                 (request.form['nome_falecido'], request.form['data_falecimento'],
                  request.form.get('observacoes', ''), datetime.now().strftime("%d/%m/%Y")))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/financeiro/novo', methods=['POST'])
def novo_financeiro():
    if not can_tesouraria(): return redirect(url_for('dashboard'))
    data_raw = request.form.get('data_movimento') or datetime.now().strftime("%Y-%m-%d")
    dt_obj = datetime.strptime(data_raw, "%Y-%m-%d")
    conn = get_db()
    conn.execute('''INSERT INTO financeiro (tipo, local_movimento, departamento, categoria, valor, data_movimento, dia, mes, ano, data_registo, membro_id, descricao)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
                 (request.form['tipo'], request.form.get('local_movimento', 'Caixa'), request.form.get('departamento', 'Geral'),
                  request.form['categoria'], float(request.form['valor']), dt_obj.strftime("%d/%m/%Y"), dt_obj.day, dt_obj.month, dt_obj.year,
                  datetime.now().strftime("%d/%m/%Y %H:%M"), request.form.get('membro_id') or None, request.form['descricao']))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/financeiro/transferir', methods=['POST'])
def transferir_fundos():
    if not can_tesouraria(): return redirect(url_for('dashboard'))
    data_raw = request.form.get('data_movimento') or datetime.now().strftime("%Y-%m-%d")
    dt_obj = datetime.strptime(data_raw, "%Y-%m-%d")
    origem_local, origem_depto = request.form['origem_local'], request.form['origem_depto']
    destino_local, destino_depto = request.form['destino_local'], request.form['destino_depto']
    valor = float(request.form['valor'])
    motivo = request.form['motivo']
    agora = datetime.now().strftime("%d/%m/%Y %H:%M")

    conn = get_db()
    conn.execute('''INSERT INTO financeiro (tipo, local_movimento, departamento, categoria, valor, data_movimento, dia, mes, ano, data_registo, descricao)
                    VALUES ('Saída', ?, ?, 'Transferência de Fundos', ?, ?, ?, ?, ?, ?, ?)''',
                 (origem_local, origem_depto, valor, dt_obj.strftime("%d/%m/%Y"), dt_obj.day, dt_obj.month, dt_obj.year, agora, f"Transf. p/ [{destino_local} - {destino_depto}]: {motivo}"))
    conn.execute('''INSERT INTO financeiro (tipo, local_movimento, departamento, categoria, valor, data_movimento, dia, mes, ano, data_registo, descricao)
                    VALUES ('Entrada', ?, ?, 'Transferência de Fundos', ?, ?, ?, ?, ?, ?, ?)''',
                 (destino_local, destino_depto, valor, dt_obj.strftime("%d/%m/%Y"), dt_obj.day, dt_obj.month, dt_obj.year, agora, f"Transf. de [{origem_local} - {origem_depto}]: {motivo}"))
    conn.execute('''INSERT INTO transferencias (data_movimento, origem_local, origem_depto, destino_local, destino_depto, valor, motivo, data_registo)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)''',
                 (dt_obj.strftime("%d/%m/%Y"), origem_local, origem_depto, destino_local, destino_depto, valor, motivo, agora))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/usuarios/novo', methods=['POST'])
def novo_usuario():
    if not is_admin(): return redirect(url_for('dashboard'))
    user, senha, cargo = request.form['usuario'].strip(), request.form['senha'].strip(), request.form['cargo'].strip()
    if user and senha:
        conn = get_db()
        try:
            conn.execute("INSERT INTO usuarios (usuario, senha, cargo) VALUES (?, ?, ?)", (user, senha, cargo))
            conn.commit()
        except sqlite3.IntegrityError: pass
        conn.close()
    return redirect(url_for('dashboard'))

@app.route('/usuarios/apagar/<int:id>')
def apagar_usuario(id):
    if not is_admin(): return redirect(url_for('dashboard'))
    conn = get_db()
    conn.execute("DELETE FROM usuarios WHERE id = ? AND id != 1", (id,))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/apagar/<tabela>/<int:id>')
def apagar_registo(tabela, id):
    if not is_admin(): return redirect(url_for('dashboard'))
    mapa = {'membro': ('membros', 'id'), 'financeiro': ('financeiro', 'id'), 'casamento': ('casamentos', 'id'), 'morte': ('mortes', 'id')}
    if tabela in mapa:
        tab, col = mapa[tabela]
        conn = get_db()
        conn.execute(f"DELETE FROM {tab} WHERE {col} = ?", (id,))
        conn.commit()
        conn.close()
    return redirect(url_for('dashboard'))

@app.route('/config/departamento/novo', methods=['POST'])
def novo_depto():
    if not is_admin(): return redirect(url_for('dashboard'))
    nome = request.form.get('nome', '').strip()
    if nome:
        conn = get_db()
        try:
            conn.execute("INSERT INTO departamentos_lista (nome) VALUES (?)", (nome,))
            conn.commit()
        except sqlite3.IntegrityError: pass
        conn.close()
    return redirect(url_for('dashboard'))

@app.route('/config/departamento/apagar/<int:id>')
def apagar_depto(id):
    if not is_admin(): return redirect(url_for('dashboard'))
    conn = get_db()
    conn.execute("DELETE FROM departamentos_lista WHERE id = ?", (id,))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/config/categoria/novo', methods=['POST'])
def nova_categoria():
    if not is_admin(): return redirect(url_for('dashboard'))
    tipo, nome = request.form.get('tipo'), request.form.get('nome', '').strip()
    if tipo and nome:
        conn = get_db()
        conn.execute("INSERT INTO categorias_financeiras (tipo, nome) VALUES (?, ?)", (tipo, nome))
        conn.commit()
        conn.close()
    return redirect(url_for('dashboard'))

@app.route('/config/categoria/apagar/<int:id>')
def apagar_categoria(id):
    if not is_admin(): return redirect(url_for('dashboard'))
    conn = get_db()
    conn.execute("DELETE FROM categorias_financeiras WHERE id = ?", (id,))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/config/zona/novo', methods=['POST'])
def nova_zona():
    if not is_admin(): return redirect(url_for('dashboard'))
    nome = request.form.get('nome', '').strip()
    if nome:
        conn = get_db()
        try:
            conn.execute("INSERT INTO zonas_lista (nome) VALUES (?)", (nome,))
            conn.commit()
        except sqlite3.IntegrityError: pass
        conn.close()
    return redirect(url_for('dashboard'))

@app.route('/config/zona/apagar/<int:id>')
def apagar_zona(id):
    if not is_admin(): return redirect(url_for('dashboard'))
    conn = get_db()
    conn.execute("DELETE FROM zonas_lista WHERE id = ?", (id,))
    conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))

@app.route('/exportar/financeiro')
def exportar_financeiro():
    if not can_tesouraria(): return redirect(url_for('dashboard'))
    conn = get_db()
    rows = conn.execute("SELECT data_movimento, tipo, local_movimento, departamento, categoria, descricao, valor FROM financeiro ORDER BY id DESC").fetchall()
    conn.close()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tesouraria Chicuque"
    ws.append(["Data Movimento", "Tipo", "Conta", "Departamento", "Categoria", "Descrição", "Valor (MT)"])
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    for col in range(1, 8):
        c = ws.cell(row=1, column=col)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append([r['data_movimento'], r['tipo'], r['local_movimento'], r['departamento'], r['categoria'], r['descricao'], r['valor']])
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"Tesouraria_IEAD_Chicuque_{datetime.now().strftime('%Y%m%d')}.xlsx")

@app.route('/exportar/membros')
def exportar_membros():
    if not can_cadastro(): return redirect(url_for('dashboard'))
    conn = get_db()
    rows = conn.execute('''
        SELECT nome, telefone, genero, data_nascimento, faixa_etaria, naturalidade, bairro, 
               filiacao, tipo_documento, numero_documento, ano_conversao, data_batismo, 
               posicao_atual, progressoes, departamento, observacoes
        FROM membros ORDER BY nome ASC
    ''').fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Membresia IEAD Chicuque"

    headers = [
        "Nome Completo", "Contacto", "Género", "Data Nasc.", "Segmento", "Naturalidade", "Bairro",
        "Filiação", "Tipo Doc", "Nº Documento", "Ano Conv.", "Data Batismo",
        "Posição Atual", "Progressões", "Departamento", "Observações"
    ]
    ws.append(headers)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center")

    for r in rows:
        ws.append([r[k] for k in r.keys()])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 13)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"Membros_Completos_IEAD_{datetime.now().strftime('%Y%m%d')}.xlsx")

if __name__ == '__main__':
    port = int(os.environ.get("PORT", 5000))
    app.run(host='0.0.0.0', port=port)
"""

with open("app.py", "w", encoding="utf-8") as f:
    f.write(app_code)
print("✓ 1/2: app.py atualizado com agregações analíticas de departamentos, fontes e saídas!")

# -------------------------------------------------------------
# 2. ATUALIZAÇÃO DO TEMPLATE COM OS 4 GRÁFICOS NO PAINEL
# -------------------------------------------------------------
html_code = """<!DOCTYPE html>
<html lang="pt">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0, maximum-scale=1.0, user-scalable=no">
    <title>IEAD Chicuque - Gestão Eclesiástica</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <script src="https://cdn.jsdelivr.net/npm/chart.js"></script>
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
    <link href="https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700;800&display=swap" rel="stylesheet">
    <style>
        body { font-family: 'Plus Jakarta Sans', sans-serif; }
        .tab-content { display: none; }
        .tab-content.active { display: block; }
        .no-scrollbar::-webkit-scrollbar { display: none; }
        .no-scrollbar { -ms-overflow-style: none; scrollbar-width: none; }
        .card-glow { box-shadow: 0 10px 25px -5px rgba(30, 58, 138, 0.08), 0 8px 10px -6px rgba(30, 58, 138, 0.04); }
    </style>
</head>
<body class="bg-gradient-to-br from-slate-100 via-blue-50/50 to-indigo-50/40 text-slate-900 min-h-screen antialiased selection:bg-blue-600 selection:text-white">

    <header class="bg-gradient-to-r from-blue-950 via-indigo-950 to-slate-950 text-white shadow-xl sticky top-0 z-40 border-b border-indigo-500/20">
        <div class="max-w-7xl mx-auto px-4 py-3.5 flex justify-between items-center">
            <div class="flex items-center space-x-3">
                <div class="w-12 h-12 rounded-2xl bg-gradient-to-tr from-white to-blue-50 p-1.5 shadow-lg flex items-center justify-center flex-shrink-0 ring-2 ring-amber-400/40">
                    <img src="/static/logo.svg?v=2026" alt="IEAD" class="w-full h-full object-contain">
                </div>
                <div>
                    <h1 class="text-sm sm:text-base md:text-lg font-extrabold tracking-tight leading-tight uppercase text-white flex items-center gap-1.5">
                        Assembleia de Deus
                        <span class="inline-block w-2 h-2 rounded-full bg-emerald-400 animate-pulse"></span>
                    </h1>
                    <p class="text-xs text-indigo-200 font-semibold tracking-wide">Congregação de Chicuque • Gestão Integrada</p>
                </div>
            </div>

            <div class="flex items-center space-x-2.5">
                <div class="hidden sm:flex flex-col text-right pr-1">
                    <span class="text-xs font-black text-white leading-none tracking-wide">{{ session['usuario'] }}</span>
                    <span class="text-[11px] text-amber-300 font-bold mt-0.5">{{ session['cargo'] }}</span>
                </div>
                <div class="sm:hidden px-2.5 py-1 bg-indigo-900/80 rounded-xl border border-indigo-400/30 text-[11px] font-bold text-amber-300">
                    {{ session['usuario'] }}
                </div>
                <a href="/logout" class="bg-gradient-to-r from-rose-500 to-red-600 hover:from-rose-600 hover:to-red-700 active:scale-95 text-white text-xs font-extrabold px-3.5 py-2 rounded-xl shadow-md shadow-rose-900/30 transition duration-150 flex items-center gap-1">
                    <span>Sair</span>
                </a>
            </div>
        </div>

        <div class="max-w-7xl mx-auto px-3 overflow-x-auto no-scrollbar border-t border-white/10 py-2.5 bg-black/10">
            <nav class="flex space-x-2 text-sm font-semibold whitespace-nowrap">
                {% if pode_cadastro %}
                <button onclick="trocarAba('membros')" id="btn-membros" class="tab-btn px-4 py-2.5 rounded-2xl transition flex items-center gap-2 bg-gradient-to-r from-blue-600 to-indigo-600 text-white font-black shadow-md shadow-blue-900/40">
                    <span>👤</span> <span>Membros ({{ total_membros }})</span>
                </button>
                <button onclick="trocarAba('casamentos')" id="btn-casamentos" class="tab-btn px-4 py-2.5 rounded-2xl transition flex items-center gap-2 text-indigo-100 hover:bg-white/10 hover:text-white">
                    <span>💍</span> <span>Casamentos</span>
                </button>
                <button onclick="trocarAba('mortes')" id="btn-mortes" class="tab-btn px-4 py-2.5 rounded-2xl transition flex items-center gap-2 text-indigo-100 hover:bg-white/10 hover:text-white">
                    <span>🕊️</span> <span>Óbitos</span>
                </button>
                {% endif %}

                {% if pode_tesouraria %}
                <button onclick="trocarAba('dashboard')" id="btn-dashboard" class="tab-btn px-4 py-2.5 rounded-2xl transition flex items-center gap-2 {% if not pode_cadastro %}bg-gradient-to-r from-blue-600 to-indigo-600 text-white font-black shadow-md{% else %}text-indigo-100 hover:bg-white/10 hover:text-white{% endif %}">
                    <span>📊</span> <span>Painel Financeiro</span>
                </button>
                <button onclick="trocarAba('financeiro')" id="btn-financeiro" class="tab-btn px-4 py-2.5 rounded-2xl transition flex items-center gap-2 text-indigo-100 hover:bg-white/10 hover:text-white">
                    <span>💰</span> <span>Tesouraria</span>
                </button>
                <button onclick="trocarAba('transferencias')" id="btn-transferencias" class="tab-btn px-4 py-2.5 rounded-2xl transition flex items-center gap-2 text-cyan-300 hover:bg-white/10">
                    <span>🔄</span> <span>Transferências</span>
                </button>
                {% endif %}

                {% if e_admin %}
                <button onclick="trocarAba('configuracoes')" id="btn-configuracoes" class="tab-btn px-4 py-2.5 rounded-2xl transition flex items-center gap-2 text-amber-300 hover:bg-white/10">
                    <span>⚙️</span> <span>Zonas & Deptos</span>
                </button>
                <button onclick="trocarAba('usuarios')" id="btn-usuarios" class="tab-btn px-4 py-2.5 rounded-2xl transition flex items-center gap-2 text-emerald-300 hover:bg-white/10">
                    <span>👥</span> <span>Utilizadores</span>
                </button>
                {% endif %}
            </nav>
        </div>
    </header>

    <main class="max-w-7xl mx-auto p-3.5 sm:p-5 lg:p-6 space-y-6">

        {% if pode_cadastro %}
        <!-- ================= ABA MEMBROS ================= -->
        <section id="aba-membros" class="tab-content {% if pode_cadastro %}active{% endif %} space-y-6">
            <div class="grid grid-cols-1 lg:grid-cols-12 gap-6">
                <!-- Formulário -->
                <div class="lg:col-span-5 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5 sm:p-6">
                    <div class="flex items-center justify-between pb-4 mb-5 border-b border-indigo-50">
                        <div>
                            <h2 class="text-lg font-black text-slate-900 tracking-tight flex items-center gap-2">
                                <span class="text-blue-600">📝</span> Ficha de Membro
                            </h2>
                            <p class="text-xs text-slate-500 font-medium">Registo completo com foto e dados</p>
                        </div>
                        <span class="bg-gradient-to-r from-blue-500 to-indigo-600 text-white text-[11px] font-black px-3 py-1 rounded-full shadow-sm">Oficial IEAD</span>
                    </div>

                    <form action="/membros/novo" method="POST" enctype="multipart/form-data" class="space-y-5">
                        <div class="flex flex-col items-center justify-center p-5 bg-gradient-to-b from-blue-50/60 to-indigo-50/40 border-2 border-dashed border-indigo-200 rounded-3xl relative group hover:border-blue-500 transition duration-200">
                            <div class="relative w-24 h-24 mb-3">
                                <img id="previewFoto" src="" alt="" class="hidden w-24 h-24 rounded-full object-cover shadow-lg ring-4 ring-white border-2 border-blue-600">
                                <div id="placeholderFoto" class="w-24 h-24 rounded-full bg-gradient-to-tr from-blue-600 to-indigo-700 text-white flex flex-col items-center justify-center shadow-lg ring-4 ring-blue-100">
                                    <svg class="w-9 h-9" fill="none" stroke="currentColor" viewBox="0 0 24 24"><path stroke-linecap="round" stroke-linejoin="round" stroke-width="2" d="M3 9a2 2 0 012-2h.93a2 2 0 001.664-.89l.812-1.22A2 2 0 0110.07 4h3.86a2 2 0 011.664.89l.812 1.22A2 2 0 0018.07 7H19a2 2 0 012 2v9a2 2 0 01-2 2H5a2 2 0 01-2-2V9z"></path><path stroke-linecap="round" stroke-linejoin="round" stroke-width="2" d="M15 13a3 3 0 11-6 0 3 3 0 016 0z"></path></svg>
                                </div>
                            </div>
                            <label class="cursor-pointer bg-gradient-to-r from-blue-600 to-indigo-600 hover:from-blue-700 hover:to-indigo-700 text-white text-xs font-bold px-5 py-2.5 rounded-xl shadow-md shadow-blue-600/30 active:scale-95 transition flex items-center gap-1.5">
                                <span>Tirar Foto / Carregar</span>
                                <input type="file" name="foto" id="inputFoto" accept="image/*" capture="environment" class="hidden" onchange="mostrarPreview(event)">
                            </label>
                            <p class="text-[11px] text-slate-500 font-semibold mt-2">Fotografia do membro para a ficha</p>
                        </div>

                        <div class="space-y-3 pt-1">
                            <span class="text-xs font-black uppercase tracking-wider text-blue-900 bg-blue-50 border border-blue-200 px-3 py-1 rounded-lg inline-block">1. Identificação Pessoal</span>
                            <div>
                                <label class="block text-sm font-bold text-slate-800 mb-1.5">Nome Completo *</label>
                                <input type="text" name="nome" required placeholder="Digite o nome completo" class="w-full h-12 px-4 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none focus:border-blue-600 font-medium">
                            </div>
                            <div class="grid grid-cols-1 sm:grid-cols-2 gap-3">
                                <div>
                                    <label class="block text-sm font-bold text-slate-800 mb-1.5">Contacto Telefónico</label>
                                    <input type="text" name="telefone" placeholder="+258 8..." class="w-full h-12 px-4 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none focus:border-blue-600">
                                </div>
                                <div>
                                    <label class="block text-sm font-bold text-slate-800 mb-1.5">Género</label>
                                    <select name="genero" class="w-full h-12 px-3.5 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none bg-white font-semibold">
                                        <option value="Masculino">Masculino</option>
                                        <option value="Feminino">Feminino</option>
                                    </select>
                                </div>
                            </div>
                            <div class="grid grid-cols-1 sm:grid-cols-2 gap-3">
                                <div>
                                    <label class="block text-sm font-bold text-slate-800 mb-1.5">Data de Nascimento</label>
                                    <input type="date" name="data_nascimento" class="w-full h-12 px-3 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none bg-white">
                                </div>
                                <div>
                                    <label class="block text-sm font-bold text-slate-800 mb-1.5">Segmento Etário *</label>
                                    <select name="faixa_etaria" required class="w-full h-12 px-3.5 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none bg-white font-black text-blue-900">
                                        <option value="Adulto">Adulto</option>
                                        <option value="Jovem">Jovem</option>
                                        <option value="Adolescente">Adolescente</option>
                                        <option value="Criança">Criança (Boa Esperança)</option>
                                        <option value="Terceira Idade">Terceira Idade</option>
                                    </select>
                                </div>
                            </div>
                            <div class="grid grid-cols-1 sm:grid-cols-2 gap-3">
                                <div>
                                    <label class="block text-sm font-bold text-slate-800 mb-1.5">Naturalidade</label>
                                    <input type="text" name="naturalidade" placeholder="Cidade / Província" class="w-full h-12 px-4 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none">
                                </div>
                                <div>
                                    <label class="block text-sm font-bold text-slate-800 mb-1.5">Bairro / Zona</label>
                                    <select name="bairro" class="w-full h-12 px-3.5 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none bg-white">
                                        <option value="">-- Selecione a Zona --</option>
                                        {% for z in lista_zonas %}<option value="{{ z['nome'] }}">{{ z['nome'] }}</option>{% endfor %}
                                    </select>
                                </div>
                            </div>
                            <div>
                                <label class="block text-sm font-bold text-slate-800 mb-1.5">Filiação (Pai & Mãe)</label>
                                <input type="text" name="filiacao" placeholder="Nome dos pais" class="w-full h-12 px-4 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none">
                            </div>
                        </div>

                        <div class="space-y-3 pt-2">
                            <span class="text-xs font-black uppercase tracking-wider text-amber-900 bg-amber-50 border border-amber-200 px-3 py-1 rounded-lg inline-block">2. Documentação Civil</span>
                            <div class="grid grid-cols-1 sm:grid-cols-3 gap-3">
                                <div>
                                    <label class="block text-sm font-bold text-slate-800 mb-1.5">Tipo Doc.</label>
                                    <select name="tipo_documento" class="w-full h-12 px-3 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none bg-white font-bold">
                                        <option value="BI">B.I.</option>
                                        <option value="Cédula">Cédula</option>
                                        <option value="Cartão de Eleitor">C. Eleitor</option>
                                        <option value="Passaporte">Passaporte</option>
                                        <option value="Nenhum">Sem Documento</option>
                                    </select>
                                </div>
                                <div class="sm:col-span-2">
                                    <label class="block text-sm font-bold text-slate-800 mb-1.5">Número do Documento</label>
                                    <input type="text" name="numero_documento" placeholder="Número do BI ou Cédula" class="w-full h-12 px-4 text-sm sm:text-base font-mono font-bold border-2 border-slate-200 rounded-2xl outline-none">
                                </div>
                            </div>
                        </div>

                        <div class="space-y-3 pt-2">
                            <span class="text-xs font-black uppercase tracking-wider text-indigo-900 bg-indigo-50 border border-indigo-200 px-3 py-1 rounded-lg inline-block">3. Vida Eclesiástica & Cargos</span>
                            <div class="grid grid-cols-1 sm:grid-cols-2 gap-3">
                                <div>
                                    <label class="block text-sm font-bold text-slate-800 mb-1.5">Ano de Conversão</label>
                                    <input type="number" name="ano_conversao" placeholder="Ex: 2016" min="1930" max="2035" class="w-full h-12 px-4 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none font-medium">
                                </div>
                                <div>
                                    <label class="block text-sm font-bold text-slate-800 mb-1.5">Data Batismo Águas</label>
                                    <input type="date" name="data_batismo" class="w-full h-12 px-3 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none font-medium">
                                </div>
                            </div>
                            <div class="grid grid-cols-1 sm:grid-cols-2 gap-3">
                                <div>
                                    <label class="block text-sm font-bold text-blue-950 mb-1.5">Posição / Cargo *</label>
                                    <select name="posicao_atual" required class="w-full h-12 px-3 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none bg-white font-black text-blue-900">
                                        <option value="Membro em Comunhão">Membro em Comunhão</option>
                                        <option value="Pastor">Pastor</option>
                                        <option value="Presbítero">Presbítero</option>
                                        <option value="Diácono / Diaconisa">Diácono / Diaconisa</option>
                                        <option value="Evangelista">Evangelista</option>
                                        <option value="Obreiro(a)">Obreiro(a)</option>
                                        <option value="Líder de Célula">Líder de Célula</option>
                                        <option value="Membro em Prova">Membro em Prova</option>
                                    </select>
                                </div>
                                <div>
                                    <label class="block text-sm font-bold text-indigo-950 mb-1.5">Departamento *</label>
                                    <select name="departamento" required class="w-full h-12 px-3 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none bg-white font-black text-indigo-900">
                                        <option value="Activista">Activista</option>
                                        <option value="Juventude">Juventude</option>
                                        <option value="Mulher (Senhoras)">Mulher (Senhoras)</option>
                                        <option value="Boa Esperança (Crianças)">Boa Esperança (Crianças)</option>
                                        <option value="Homens / Obreiros">Homens / Obreiros</option>
                                        <option value="Louvor / Música">Louvor / Música</option>
                                        <option value="Ação Social">Ação Social</option>
                                        <option value="Geral">Geral</option>
                                    </select>
                                </div>
                            </div>
                            <div>
                                <label class="block text-sm font-bold text-slate-800 mb-1.5">Espaço para Progressões Eclesiásticas</label>
                                <textarea name="progressoes" rows="2" placeholder="Histórico de consagrações ministeriais..." class="w-full p-3.5 text-sm border-2 border-slate-200 rounded-2xl outline-none"></textarea>
                            </div>
                            <div>
                                <label class="block text-sm font-bold text-slate-800 mb-1.5">Observações Gerais</label>
                                <input type="text" name="observacoes" placeholder="Notas pastorais..." class="w-full h-12 px-4 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none">
                            </div>
                        </div>

                        <button type="submit" class="w-full h-14 bg-gradient-to-r from-emerald-600 via-teal-600 to-emerald-700 hover:from-emerald-700 hover:to-teal-800 active:scale-[0.99] text-white font-black text-base rounded-2xl shadow-xl shadow-emerald-600/30 transition duration-200 flex items-center justify-center space-x-2">
                            <span>Gravar Ficha de Membro</span>
                            <svg class="w-5 h-5" fill="none" stroke="currentColor" viewBox="0 0 24 24"><path stroke-linecap="round" stroke-linejoin="round" stroke-width="2.5" d="M5 13l4 4L19 7"></path></svg>
                        </button>
                    </form>
                </div>

                <!-- Tabela -->
                <div class="lg:col-span-7 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5 sm:p-6 flex flex-col justify-between">
                    <div>
                        <div class="flex flex-col sm:flex-row justify-between items-stretch sm:items-center gap-3 mb-5 pb-4 border-b border-indigo-50">
                            <div>
                                <h3 class="text-lg font-black text-slate-900">Membros Registados</h3>
                                <p class="text-xs text-slate-500 font-medium">Arquivo eclesiástico digital</p>
                            </div>
                            <div>
                                <a href="/exportar/membros" class="h-11 px-4 bg-gradient-to-r from-emerald-500 to-teal-600 hover:from-emerald-600 hover:to-teal-700 text-white rounded-xl text-xs sm:text-sm font-black shadow-md shadow-emerald-600/20 transition flex items-center justify-center gap-2 active:scale-95">
                                    <span>Baixar Excel</span>
                                </a>
                            </div>
                        </div>

                        <div class="mb-4">
                            <input type="text" id="filtroMembros" onkeyup="filtrarTabela('filtroMembros', 'tabelaMembros')" 
                                   placeholder="Pesquisar por nome, BI, cargo, departamento..." 
                                   class="w-full h-12 px-4 text-sm sm:text-base border-2 border-indigo-100 rounded-2xl outline-none focus:border-blue-600 transition bg-slate-50/50 hover:bg-white focus:bg-white font-medium">
                        </div>

                        <div class="overflow-x-auto max-h-[640px] rounded-2xl border border-indigo-50">
                            <table id="tabelaMembros" class="w-full text-left text-sm min-w-[620px]">
                                <thead class="bg-gradient-to-r from-blue-900 to-indigo-950 text-white font-black text-xs uppercase tracking-wider sticky top-0 z-10">
                                    <tr>
                                        <th class="p-3.5">Membro</th>
                                        <th class="p-3.5">Posição / Cargo</th>
                                        <th class="p-3.5">Departamento</th>
                                        <th class="p-3.5">Documento</th>
                                        <th class="p-3.5 text-center">Ação</th>
                                    </tr>
                                </thead>
                                <tbody class="divide-y divide-slate-100 text-slate-700 bg-white">
                                    {% for m in todos_membros %}
                                    <tr class="hover:bg-blue-50/60 transition">
                                        <td class="p-3.5 flex items-center space-x-3">
                                            {% if m['foto_path'] %}
                                            <img src="{{ m['foto_path'] }}" class="w-12 h-12 object-cover rounded-2xl border-2 border-white shadow-md ring-2 ring-blue-600/30">
                                            {% else %}
                                            <div class="w-12 h-12 rounded-2xl bg-gradient-to-tr from-blue-600 to-indigo-700 text-white flex items-center justify-center font-black text-sm shadow-md ring-2 ring-blue-600/20">
                                                {{ m['nome'][:2].upper() }}
                                            </div>
                                            {% endif %}
                                            <div>
                                                <span class="font-extrabold text-slate-900 block text-sm sm:text-base leading-tight">{{ m['nome'] }}</span>
                                                <span class="text-xs text-slate-500 font-semibold block mt-0.5">{{ m['faixa_etaria'] or 'Adulto' }} • {{ m['telefone'] or 'Sem tel' }}</span>
                                            </div>
                                        </td>
                                        <td class="p-3.5">
                                            <span class="bg-blue-50 border border-blue-200 text-blue-900 font-black px-3 py-1 rounded-xl text-xs inline-block shadow-sm">
                                                {{ m['posicao_atual'] or 'Membro' }}
                                            </span>
                                        </td>
                                        <td class="p-3.5">
                                            <span class="px-3 py-1 rounded-xl text-xs font-black inline-block shadow-sm bg-slate-100 text-slate-800 border">
                                                {{ m['departamento'] or 'Geral' }}
                                            </span>
                                        </td>
                                        <td class="p-3.5 font-mono text-xs text-slate-700 font-bold">
                                            {{ m['tipo_documento'] or 'BI' }}: {{ m['numero_documento'] or '-' }}
                                        </td>
                                        <td class="p-3.5 text-center whitespace-nowrap">
                                            <button onclick="verCartaoMembro({{ m['id'] }})" class="bg-gradient-to-r from-blue-600 to-indigo-700 text-white font-extrabold px-3.5 py-1.5 rounded-xl text-xs shadow-md transition active:scale-95 mr-1">
                                                Ficha
                                            </button>
                                            {% if e_admin %}
                                            <a href="/apagar/membro/{{ m['id'] }}" onclick="return confirm('Eliminar?');" class="text-rose-600 hover:text-rose-800 font-black text-xs p-1">✕</a>
                                            {% endif %}
                                        </td>
                                    </tr>
                                    {% endfor %}
                                </tbody>
                            </table>
                        </div>
                    </div>
                </div>

            </div>
        </section>

        <!-- ================= ABA CASAMENTOS ================= -->
        <section id="aba-casamentos" class="tab-content space-y-6">
            <div class="grid grid-cols-1 lg:grid-cols-12 gap-6">
                <div class="lg:col-span-5 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5 sm:p-6">
                    <h3 class="text-base sm:text-lg font-black text-slate-900 mb-4 pb-3 border-b border-indigo-50 flex items-center gap-2">
                        <span class="text-purple-600">💍</span> Registo Matrimonial
                    </h3>
                    <form action="/casamentos/novo" method="POST" class="space-y-4">
                        <input type="text" name="noivo" required placeholder="Nome do noivo *" class="w-full h-12 px-4 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none focus:border-purple-600">
                        <input type="text" name="noiva" required placeholder="Nome da noiva *" class="w-full h-12 px-4 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none focus:border-purple-600">
                        <input type="date" name="data_casamento" required class="w-full h-12 px-3 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none focus:border-purple-600 font-semibold">
                        <input type="text" name="pastor_oficiante" placeholder="Pastor Oficiante" class="w-full h-12 px-4 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none focus:border-purple-600">
                        <button type="submit" class="w-full h-12 bg-gradient-to-r from-purple-700 to-indigo-800 text-white font-black rounded-2xl shadow-lg transition">Gravar Matrimónio</button>
                    </form>
                </div>
                <div class="lg:col-span-7 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5 sm:p-6">
                    <h3 class="text-base sm:text-lg font-black text-slate-900 mb-4 pb-3 border-b border-indigo-50">Livro de Casamentos</h3>
                    <div class="overflow-x-auto rounded-2xl border border-indigo-50">
                        <table class="w-full text-left text-sm">
                            <thead class="bg-gradient-to-r from-purple-900 to-indigo-950 text-white font-black text-xs uppercase">
                                <tr><th class="p-3">Data</th><th class="p-3">Casal</th><th class="p-3">Oficiante</th>{% if e_admin %}<th class="p-3 text-center">Ação</th>{% endif %}</tr>
                            </thead>
                            <tbody class="divide-y divide-slate-100 bg-white">
                                {% for c in todos_casamentos %}
                                <tr class="hover:bg-purple-50/50 transition">
                                    <td class="p-3 font-bold">{{ c['data_casamento'] }}</td>
                                    <td class="p-3 font-extrabold text-purple-950">{{ c['noivo'] }} & {{ c['noiva'] }}</td>
                                    <td class="p-3 text-slate-600">{{ c['pastor_oficiante'] or '-' }}</td>
                                    {% if e_admin %}<td class="p-3 text-center"><a href="/apagar/casamento/{{ c['id'] }}" onclick="return confirm('Eliminar?');" class="text-rose-600 font-bold">✕</a></td>{% endif %}
                                </tr>
                                {% endfor %}
                            </tbody>
                        </table>
                    </div>
                </div>
            </div>
        </section>

        <!-- ================= ABA ÓBITOS ================= -->
        <section id="aba-mortes" class="tab-content space-y-6">
            <div class="grid grid-cols-1 lg:grid-cols-12 gap-6">
                <div class="lg:col-span-5 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5 sm:p-6">
                    <h3 class="text-base sm:text-lg font-black text-slate-900 mb-4 pb-3 border-b border-indigo-50 flex items-center gap-2">
                        <span class="text-slate-600">🕊️</span> Registo de Óbito
                    </h3>
                    <form action="/mortes/novo" method="POST" class="space-y-4">
                        <input type="text" name="nome_falecido" required placeholder="Nome do Falecido *" class="w-full h-12 px-4 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none focus:border-slate-600">
                        <input type="date" name="data_falecimento" required class="w-full h-12 px-3 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none focus:border-slate-600 font-semibold">
                        <input type="text" name="observacoes" placeholder="Observações..." class="w-full h-12 px-4 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none focus:border-slate-600">
                        <button type="submit" class="w-full h-12 bg-gradient-to-r from-slate-800 to-slate-950 text-white font-black rounded-2xl shadow-lg transition">Gravar Óbito</button>
                    </form>
                </div>
                <div class="lg:col-span-7 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5 sm:p-6">
                    <h3 class="text-base sm:text-lg font-black text-slate-900 mb-4 pb-3 border-b border-indigo-50">Livro de Falecimentos</h3>
                    <div class="overflow-x-auto rounded-2xl border border-indigo-50">
                        <table class="w-full text-left text-sm">
                            <thead class="bg-gradient-to-r from-slate-800 to-slate-950 text-white font-black text-xs uppercase">
                                <tr><th class="p-3">Data</th><th class="p-3">Falecido</th><th class="p-3">Observações</th>{% if e_admin %}<th class="p-3 text-center">Ação</th>{% endif %}</tr>
                            </thead>
                            <tbody class="divide-y divide-slate-100 bg-white">
                                {% for m in todas_mortes %}
                                <tr>
                                    <td class="p-3 font-bold">{{ m['data_falecimento'] }}</td>
                                    <td class="p-3 font-black text-slate-900">{{ m['nome_falecido'] }}</td>
                                    <td class="p-3 text-slate-500">{{ m['observacoes'] or '-' }}</td>
                                    {% if e_admin %}<td class="p-3 text-center"><a href="/apagar/morte/{{ m['id'] }}" onclick="return confirm('Eliminar?');" class="text-rose-600 font-bold">✕</a></td>{% endif %}
                                </tr>
                                {% endfor %}
                            </tbody>
                        </table>
                    </div>
                </div>
            </div>
        </section>
        {% endif %}

        {% if pode_tesouraria %}
        <!-- ================= ABA DASHBOARD FINANCEIRO COM OS NOVOS GRÁFICOS ================= -->
        <section id="aba-dashboard" class="tab-content {% if not pode_cadastro and pode_tesouraria %}active{% endif %} space-y-6">
            
            <!-- Indicadores Principais -->
            <div class="grid grid-cols-1 sm:grid-cols-2 lg:grid-cols-4 gap-4">
                <div class="bg-gradient-to-br from-amber-500 to-orange-600 text-white p-5 rounded-3xl shadow-xl shadow-amber-500/20 border border-amber-400">
                    <span class="text-xs font-black uppercase text-amber-100 block">💵 Caixa Físico</span>
                    <p class="text-2xl sm:text-3xl font-black text-white mt-2 tracking-tight">{{ "{:,.2f}".format(saldo_caixa) }} <span class="text-xs font-bold text-amber-200">MT</span></p>
                </div>
                <div class="bg-gradient-to-br from-blue-600 to-indigo-700 text-white p-5 rounded-3xl shadow-xl shadow-blue-600/20 border border-blue-400">
                    <span class="text-xs font-black uppercase text-blue-100 block">🏛️ Conta Bancária</span>
                    <p class="text-2xl sm:text-3xl font-black text-white mt-2 tracking-tight">{{ "{:,.2f}".format(saldo_banco) }} <span class="text-xs font-bold text-blue-200">MT</span></p>
                </div>
                <div class="bg-gradient-to-br from-emerald-600 to-teal-800 text-white p-5 rounded-3xl shadow-xl shadow-emerald-600/25 border border-emerald-400">
                    <span class="text-xs font-black uppercase text-emerald-100 block">🏦 Saldo Consolidado</span>
                    <p class="text-2xl sm:text-3xl font-black text-white mt-2 tracking-tight">{{ "{:,.2f}".format(saldo_total) }} <span class="text-xs font-bold text-emerald-200">MT</span></p>
                </div>
                <div class="bg-gradient-to-br from-purple-700 to-indigo-900 text-white p-5 rounded-3xl shadow-xl shadow-purple-700/20 border border-purple-500">
                    <span class="text-xs font-black uppercase text-purple-200 block">👥 Membresia Registada</span>
                    <p class="text-2xl sm:text-3xl font-black text-white mt-2 tracking-tight">{{ total_membros }} <span class="text-xs font-bold text-purple-200">membros</span></p>
                </div>
            </div>

            <!-- LINHA 1 DE GRÁFICOS: EVOLUÇÃO GERAL E SITUAÇÃO DOS DEPARTAMENTOS -->
            <div class="grid grid-cols-1 lg:grid-cols-2 gap-6">
                <!-- 1. Evolução Mensal -->
                <div class="bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5 sm:p-6">
                    <div class="flex items-center justify-between mb-4">
                        <h3 class="text-sm font-black uppercase tracking-wider text-slate-800 flex items-center gap-2">
                            <span class="w-3 h-3 rounded-full bg-emerald-500"></span> Evolução Mensal (Entradas vs Saídas)
                        </h3>
                        <span class="text-[11px] font-bold text-slate-400">Ano Corrente</span>
                    </div>
                    <div class="h-64 sm:h-72"><canvas id="graficoEvolucaoMensal"></canvas></div>
                </div>

                <!-- 2. Situação Financeira dos Departamentos -->
                <div class="bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5 sm:p-6">
                    <div class="flex items-center justify-between mb-4">
                        <h3 class="text-sm font-black uppercase tracking-wider text-slate-800 flex items-center gap-2">
                            <span class="w-3 h-3 rounded-full bg-blue-600"></span> Situação por Departamento (Fundos)
                        </h3>
                        <span class="text-[11px] font-bold text-slate-400">Entradas vs Saídas</span>
                    </div>
                    <div class="h-64 sm:h-72"><canvas id="graficoFinancasDepto"></canvas></div>
                </div>
            </div>

            <!-- LINHA 2 DE GRÁFICOS: FONTES DE FUNDOS (RECEITAS) E DESTINO DE SAÍDAS (DESPESAS) -->
            <div class="grid grid-cols-1 lg:grid-cols-2 gap-6">
                <!-- 3. Fontes de Fundos -->
                <div class="bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5 sm:p-6">
                    <div class="flex items-center justify-between mb-4">
                        <h3 class="text-sm font-black uppercase tracking-wider text-slate-800 flex items-center gap-2">
                            <span class="w-3 h-3 rounded-full bg-teal-500"></span> Fontes de Fundos (Entradas)
                        </h3>
                        <span class="text-[11px] font-bold text-teal-600 bg-teal-50 px-2 py-0.5 rounded-full border border-teal-200">Receitas</span>
                    </div>
                    <div class="h-64 sm:h-72"><canvas id="graficoFontesFundos"></canvas></div>
                </div>

                <!-- 4. Destino de Saídas -->
                <div class="bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5 sm:p-6">
                    <div class="flex items-center justify-between mb-4">
                        <h3 class="text-sm font-black uppercase tracking-wider text-slate-800 flex items-center gap-2">
                            <span class="w-3 h-3 rounded-full bg-rose-500"></span> Destino de Saídas (Despesas)
                        </h3>
                        <span class="text-[11px] font-bold text-rose-600 bg-rose-50 px-2 py-0.5 rounded-full border border-rose-200">Despesas</span>
                    </div>
                    <div class="h-64 sm:h-72"><canvas id="graficoSaidasCategorias"></canvas></div>
                </div>
            </div>
        </section>

        <!-- ================= ABA TRANSFERÊNCIAS ================= -->
        <section id="aba-transferencias" class="tab-content space-y-6">
            <div class="grid grid-cols-1 lg:grid-cols-12 gap-6">
                <div class="lg:col-span-5 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5 sm:p-6">
                    <h3 class="text-base sm:text-lg font-black text-slate-900 mb-4 pb-3 border-b border-indigo-50 flex items-center gap-2">
                        <span class="text-cyan-600">🔄</span> Transferência entre Contas
                    </h3>
                    <form action="/financeiro/transferir" method="POST" class="space-y-4">
                        <div>
                            <label class="block text-sm font-bold mb-1">Data *</label>
                            <input type="date" name="data_movimento" required class="w-full h-12 px-3 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none font-semibold">
                        </div>
                        <div class="p-4 bg-rose-50 border-2 border-rose-200 rounded-2xl space-y-2">
                            <span class="text-xs font-black text-rose-800 uppercase block">Conta de Origem</span>
                            <select name="origem_local" class="w-full h-11 px-3 text-sm border-2 border-rose-200 rounded-xl bg-white font-bold"><option value="Caixa">💵 Caixa Físico</option><option value="Banco">🏛️ Conta Bancária</option></select>
                            <select name="origem_depto" class="w-full h-11 px-3 text-sm border-2 border-rose-200 rounded-xl bg-white"><option value="Geral">Fundo Geral</option>{% for d in lista_deptos %}<option value="{{ d['nome'] }}">{{ d['nome'] }}</option>{% endfor %}</select>
                        </div>
                        <div class="p-4 bg-emerald-50 border-2 border-emerald-200 rounded-2xl space-y-2">
                            <span class="text-xs font-black text-emerald-800 uppercase block">Conta de Destino</span>
                            <select name="destino_local" class="w-full h-11 px-3 text-sm border-2 border-emerald-200 rounded-xl bg-white font-bold"><option value="Banco">🏛️ Conta Bancária</option><option value="Caixa">💵 Caixa Físico</option></select>
                            <select name="destino_depto" class="w-full h-11 px-3 text-sm border-2 border-emerald-200 rounded-xl bg-white"><option value="Geral">Fundo Geral</option>{% for d in lista_deptos %}<option value="{{ d['nome'] }}">{{ d['nome'] }}</option>{% endfor %}</select>
                        </div>
                        <div>
                            <label class="block text-sm font-bold mb-1">Valor da Transferência (MT) *</label>
                            <input type="number" step="0.01" name="valor" required placeholder="0.00" class="w-full h-12 px-4 text-base font-black border-2 border-slate-200 rounded-2xl outline-none">
                        </div>
                        <div>
                            <label class="block text-sm font-bold mb-1">Motivo *</label>
                            <input type="text" name="motivo" required placeholder="Justificação..." class="w-full h-12 px-4 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none">
                        </div>
                        <button type="submit" class="w-full h-14 bg-gradient-to-r from-cyan-600 via-blue-600 to-indigo-700 text-white font-black rounded-2xl shadow-lg transition duration-200">
                            Efetuar Transferência de Fundos
                        </button>
                    </form>
                </div>
                <div class="lg:col-span-7 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5 sm:p-6">
                    <h3 class="text-base sm:text-lg font-black text-slate-900 mb-4 pb-3 border-b border-indigo-50">Histórico de Transferências</h3>
                    <div class="overflow-x-auto max-h-[520px] rounded-2xl border border-indigo-50">
                        <table class="w-full text-left text-sm">
                            <thead class="bg-gradient-to-r from-cyan-900 to-indigo-950 text-white font-black text-xs uppercase">
                                <tr><th class="p-3">Data</th><th class="p-3">De</th><th class="p-3">Para</th><th class="p-3">Motivo</th><th class="p-3 text-right">Valor</th></tr>
                            </thead>
                            <tbody class="divide-y divide-slate-100 bg-white">
                                {% for t in ultimas_transferencias %}
                                <tr class="hover:bg-cyan-50/40 transition">
                                    <td class="p-3 font-bold">{{ t['data_movimento'] }}</td>
                                    <td class="p-3 font-bold text-rose-700">{{ t['origem_local'] }}</td>
                                    <td class="p-3 font-bold text-emerald-700">{{ t['destino_local'] }}</td>
                                    <td class="p-3 text-slate-600">{{ t['motivo'] }}</td>
                                    <td class="p-3 text-right font-black text-blue-900">{{ "{:,.2f}".format(t['valor']) }} MT</td>
                                </tr>
                                {% endfor %}
                            </tbody>
                        </table>
                    </div>
                </div>
            </div>
        </section>

        <!-- ================= ABA TESOURARIA ================= -->
        <section id="aba-financeiro" class="tab-content space-y-6">
            <div class="grid grid-cols-1 lg:grid-cols-12 gap-6">
                <div class="lg:col-span-5 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5 sm:p-6">
                    <h3 class="text-base sm:text-lg font-black text-slate-900 mb-4 pb-3 border-b border-indigo-50 flex items-center gap-2">
                        <span class="text-emerald-600">💰</span> Lançamento Financeiro
                    </h3>
                    <form action="/financeiro/novo" method="POST" class="space-y-4">
                        <div>
                            <label class="block text-sm font-bold mb-1">Data *</label>
                            <input type="date" name="data_movimento" required class="w-full h-12 px-3 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none font-semibold">
                        </div>
                        <div class="grid grid-cols-2 gap-3">
                            <div>
                                <label class="block text-sm font-bold mb-1">Tipo</label>
                                <select name="tipo" id="selectTipoTransacao" onchange="atualizarCategoriasPorTipo()" class="w-full h-12 px-3 text-sm sm:text-base border-2 border-slate-200 rounded-2xl bg-white font-black">
                                    <option value="Entrada">Entrada (+)</option>
                                    <option value="Saída">Saída (-)</option>
                                </select>
                            </div>
                            <div>
                                <label class="block text-sm font-bold mb-1">Conta *</label>
                                <select name="local_movimento" required class="w-full h-12 px-3 text-sm sm:text-base border-2 border-slate-200 rounded-2xl bg-white font-black">
                                    <option value="Caixa">💵 Caixa</option>
                                    <option value="Banco">🏛️ Banco</option>
                                </select>
                            </div>
                        </div>
                        <div class="grid grid-cols-2 gap-3">
                            <div>
                                <label class="block text-sm font-bold mb-1">Fundo / Depto</label>
                                <select name="departamento" class="w-full h-12 px-3 text-sm border-2 border-slate-200 rounded-2xl bg-white font-semibold">
                                    <option value="Geral">Geral</option>
                                    {% for d in lista_deptos %}<option value="{{ d['nome'] }}">{{ d['nome'] }}</option>{% endfor %}
                                </select>
                            </div>
                            <div>
                                <label class="block text-sm font-bold mb-1">Categoria</label>
                                <select name="categoria" id="selectCategoriaTransacao" class="w-full h-12 px-3 text-sm border-2 border-slate-200 rounded-2xl bg-white font-semibold"></select>
                            </div>
                        </div>
                        <div>
                            <label class="block text-sm font-bold mb-1">Valor (MT) *</label>
                            <input type="number" step="0.01" name="valor" required placeholder="0.00" class="w-full h-12 px-4 text-base font-black border-2 border-slate-200 rounded-2xl outline-none">
                        </div>
                        <div>
                            <label class="block text-sm font-bold mb-1">Descrição</label>
                            <input type="text" name="descricao" required placeholder="Detalhes do movimento" class="w-full h-12 px-4 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none">
                        </div>
                        <button type="submit" class="w-full h-14 bg-gradient-to-r from-emerald-600 to-teal-700 hover:from-emerald-700 hover:to-teal-800 text-white font-black rounded-2xl shadow-xl transition">
                            Gravar Lançamento no Caixa
                        </button>
                    </form>
                </div>
                <div class="lg:col-span-7 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5 sm:p-6">
                    <div class="flex flex-col sm:flex-row justify-between items-start sm:items-center gap-3 mb-4 pb-3 border-b border-indigo-50">
                        <h3 class="text-base sm:text-lg font-black text-slate-900">Extrato de Movimentos</h3>
                        <a href="/exportar/financeiro" class="h-10 px-4 bg-gradient-to-r from-emerald-500 to-teal-600 text-white rounded-xl text-xs font-black shadow-md transition flex items-center gap-1.5">
                            <span>📥 Exportar Excel</span>
                        </a>
                    </div>
                    <div class="overflow-x-auto max-h-[520px] rounded-2xl border border-indigo-50">
                        <table class="w-full text-left text-sm">
                            <thead class="bg-gradient-to-r from-emerald-900 to-teal-950 text-white font-black text-xs uppercase sticky top-0">
                                <tr><th class="p-3">Data</th><th class="p-3">Tipo</th><th class="p-3">Conta</th><th class="p-3">Categoria</th><th class="p-3">Descrição</th><th class="p-3 text-right">Valor</th>{% if e_admin %}<th class="p-3 text-center">Ação</th>{% endif %}</tr>
                            </thead>
                            <tbody class="divide-y divide-slate-100 bg-white">
                                {% for f in todas_financas %}
                                <tr class="hover:bg-emerald-50/40 transition">
                                    <td class="p-3 font-bold">{{ f['data_movimento'] }}</td>
                                    <td class="p-3"><span class="px-2.5 py-1 rounded-full text-xs font-black {{ 'bg-emerald-100 text-emerald-800 border border-emerald-300' if f['tipo'] == 'Entrada' else 'bg-rose-100 text-rose-800 border border-rose-300' }}">{{ f['tipo'] }}</span></td>
                                    <td class="p-3 font-extrabold text-slate-800">{{ f['local_movimento'] }}</td>
                                    <td class="p-3 font-semibold text-slate-700">{{ f['categoria'] }}</td>
                                    <td class="p-3 text-slate-500 text-xs">{{ f['descricao'] }}</td>
                                    <td class="p-3 text-right font-black {{ 'text-emerald-700' if f['tipo'] == 'Entrada' else 'text-rose-700' }}">{{ "{:,.2f}".format(f['valor']) }} MT</td>
                                    {% if e_admin %}<td class="p-3 text-center"><a href="/apagar/financeiro/{{ f['id'] }}" onclick="return confirm('Eliminar?');" class="text-rose-600 font-bold">✕</a></td>{% endif %}
                                </tr>
                                {% endfor %}
                            </tbody>
                        </table>
                    </div>
                </div>
            </div>
        </section>
        {% endif %}

        {% if e_admin %}
        <!-- ================= ABA CONFIGURAÇÕES ================= -->
        <section id="aba-configuracoes" class="tab-content space-y-6">
            <div class="grid grid-cols-1 md:grid-cols-3 gap-6">
                <div class="bg-white/95 p-5 rounded-3xl card-glow border border-indigo-100">
                    <h3 class="text-sm font-black uppercase text-blue-950 mb-3 pb-2 border-b">Ministérios / Deptos</h3>
                    <form action="/config/departamento/novo" method="POST" class="flex gap-2 mb-3">
                        <input type="text" name="nome" placeholder="Novo departamento" required class="h-11 px-3 text-sm border-2 border-slate-200 rounded-xl w-full">
                        <button type="submit" class="bg-blue-900 text-white px-4 font-black rounded-xl">+</button>
                    </form>
                    <ul class="divide-y divide-slate-100 text-sm max-h-56 overflow-y-auto">
                        {% for d in lista_deptos %}<li class="py-2.5 flex justify-between font-semibold"><span>{{ d['nome'] }}</span><a href="/config/departamento/apagar/{{ d['id'] }}" class="text-rose-600 font-bold">✕</a></li>{% endfor %}
                    </ul>
                </div>
                <div class="bg-white/95 p-5 rounded-3xl card-glow border border-indigo-100">
                    <h3 class="text-sm font-black uppercase text-emerald-950 mb-3 pb-2 border-b">Categorias Financeiras</h3>
                    <form action="/config/categoria/novo" method="POST" class="space-y-2 mb-3 text-sm">
                        <div class="flex gap-2">
                            <select name="tipo" class="h-11 px-2 border-2 border-slate-200 rounded-xl bg-white w-1/3 text-xs font-black"><option value="Entrada">Entrada</option><option value="Saída">Saída</option></select>
                            <input type="text" name="nome" placeholder="Nome" required class="h-11 px-3 border-2 border-slate-200 rounded-xl w-2/3">
                        </div>
                        <button type="submit" class="w-full h-10 bg-emerald-600 text-white font-black rounded-xl">+ Guardar</button>
                    </form>
                    <ul class="divide-y divide-slate-100 text-sm max-h-56 overflow-y-auto">
                        {% for c in lista_categorias %}<li class="py-2.5 flex justify-between font-semibold"><span><strong class="text-emerald-700">{{ c['tipo'] }}</strong>: {{ c['nome'] }}</span><a href="/config/categoria/apagar/{{ c['id'] }}" class="text-rose-600 font-bold">✕</a></li>{% endfor %}
                    </ul>
                </div>
                <div class="bg-white/95 p-5 rounded-3xl card-glow border border-indigo-100">
                    <h3 class="text-sm font-black uppercase text-amber-950 mb-3 pb-2 border-b">Zonas & Bairros</h3>
                    <form action="/config/zona/novo" method="POST" class="flex gap-2 mb-3">
                        <input type="text" name="nome" placeholder="Nova zona..." required class="h-11 px-3 text-sm border-2 border-slate-200 rounded-xl w-full">
                        <button type="submit" class="bg-amber-600 text-white px-4 font-black rounded-xl">+</button>
                    </form>
                    <ul class="divide-y divide-slate-100 text-sm max-h-56 overflow-y-auto">
                        {% for z in lista_zonas %}<li class="py-2.5 flex justify-between font-semibold"><span>{{ z['nome'] }}</span><a href="/config/zona/apagar/{{ z['id'] }}" class="text-rose-600 font-bold">✕</a></li>{% endfor %}
                    </ul>
                </div>
            </div>
        </section>

        <!-- ================= ABA UTILIZADORES ================= -->
        <section id="aba-usuarios" class="tab-content space-y-6">
            <div class="grid grid-cols-1 md:grid-cols-12 gap-6">
                <div class="md:col-span-5 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5 sm:p-6">
                    <h3 class="text-base sm:text-lg font-black text-slate-900 mb-4 pb-3 border-b border-indigo-50">👤 Criar Acesso</h3>
                    <form action="/usuarios/novo" method="POST" class="space-y-4">
                        <input type="text" name="usuario" required placeholder="Utilizador *" class="w-full h-12 px-4 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none">
                        <input type="password" name="senha" required placeholder="Palavra-passe *" class="w-full h-12 px-4 text-sm sm:text-base border-2 border-slate-200 rounded-2xl outline-none">
                        <select name="cargo" required class="w-full h-12 px-3 text-sm sm:text-base border-2 border-slate-200 rounded-2xl bg-white font-black text-blue-950">
                            <option value="Secretário">Secretário (Apenas Cadastro de Membros e Eventos)</option>
                            <option value="Líder">Líder de Ministério (Apenas Cadastro)</option>
                            <option value="Tesoureiro">Tesoureiro (Apenas Tesouraria e Dashboard)</option>
                            <option value="Pastor">Pastor (Administrador Geral)</option>
                        </select>
                        <button type="submit" class="w-full h-12 bg-gradient-to-r from-emerald-600 to-teal-700 text-white font-black rounded-2xl shadow-lg transition">Criar Utilizador</button>
                    </form>
                </div>
                <div class="md:col-span-7 bg-white/95 rounded-3xl card-glow border border-indigo-100 p-5 sm:p-6">
                    <h3 class="text-base sm:text-lg font-black text-slate-900 mb-4 pb-3 border-b border-indigo-50">Utilizadores Registados</h3>
                    <div class="overflow-x-auto rounded-2xl border border-indigo-50">
                        <table class="w-full text-left text-sm">
                            <thead class="bg-gradient-to-r from-blue-900 to-indigo-950 text-white font-black text-xs uppercase">
                                <tr><th class="p-3">Login</th><th class="p-3">Cargo / Permissão</th><th class="p-3 text-center">Ação</th></tr>
                            </thead>
                            <tbody class="divide-y divide-slate-100 bg-white">
                                {% for u in lista_usuarios %}
                                <tr>
                                    <td class="p-3 font-extrabold text-slate-900">{{ u['usuario'] }}</td>
                                    <td class="p-3"><span class="px-3 py-1 rounded-xl text-xs font-black inline-block bg-blue-100 text-blue-900">{{ u['cargo'] }}</span></td>
                                    <td class="p-3 text-center">{% if u['usuario'] != 'admin' %}<a href="/usuarios/apagar/{{ u['id'] }}" onclick="return confirm('Eliminar?');" class="text-rose-600 font-bold">✕</a>{% else %}<span class="text-slate-400 text-xs font-black">Principal</span>{% endif %}</td>
                                </tr>
                                {% endfor %}
                            </tbody>
                        </table>
                    </div>
                </div>
            </div>
        </section>
        {% endif %}

    </main>

    <div id="modalMembro" class="fixed inset-0 bg-slate-950/70 backdrop-blur-sm hidden items-center justify-center p-4 z-50">
        <div class="bg-white rounded-3xl max-w-lg w-full p-6 shadow-2xl relative border border-slate-200">
            <button onclick="fecharModalMembro()" class="absolute top-5 right-5 w-8 h-8 rounded-full bg-slate-100 text-slate-500 font-bold text-base flex items-center justify-center">&times;</button>
            <div id="conteudoModalMembro" class="space-y-4"></div>
        </div>
    </div>

    <script>
        function mostrarPreview(event) {
            const input = event.target;
            if (input.files && input.files[0]) {
                const reader = new FileReader();
                reader.onload = function(e) {
                    const preview = document.getElementById('previewFoto');
                    const placeholder = document.getElementById('placeholderFoto');
                    preview.src = e.target.result;
                    preview.classList.remove('hidden');
                    placeholder.classList.add('hidden');
                };
                reader.readAsDataURL(input.files[0]);
            }
        }

        const membrosDados = {{ membros_json | safe }};

        function verCartaoMembro(id) {
            const m = membrosDados.find(x => x.id === id);
            if (!m) return;

            const fotoHtml = m.foto_path 
                ? `<img src="${m.foto_path}" class="w-28 h-28 object-cover rounded-3xl mx-auto border-4 border-white shadow-xl ring-4 ring-blue-600/30">`
                : `<div class="w-28 h-28 rounded-3xl bg-gradient-to-tr from-blue-600 to-indigo-700 text-white flex items-center justify-center text-3xl font-black mx-auto border-4 border-white shadow-xl ring-4 ring-blue-600/30">${m.nome.substring(0,2).toUpperCase()}</div>`;

            document.getElementById('conteudoModalMembro').innerHTML = `
                <div class="text-center pb-4 border-b border-indigo-50">
                    ${fotoHtml}
                    <h3 class="text-xl font-black text-slate-900 mt-3">${m.nome}</h3>
                    <div class="flex justify-center gap-2 mt-2 flex-wrap">
                        <span class="bg-gradient-to-r from-blue-600 to-indigo-700 text-white text-xs font-black px-3.5 py-1 rounded-full">${m.posicao_atual || 'Membro'}</span>
                        <span class="bg-indigo-50 text-indigo-900 border border-indigo-200 text-xs font-black px-3.5 py-1 rounded-full">${m.departamento || 'Geral'}</span>
                    </div>
                </div>
                <div class="grid grid-cols-2 gap-3 text-sm text-slate-700 bg-slate-50 p-4 rounded-2xl border font-medium">
                    <div><span class="text-xs text-slate-400 font-bold block uppercase">Segmento</span><strong>${m.faixa_etaria || '-'}</strong></div>
                    <div><span class="text-xs text-slate-400 font-bold block uppercase">Contacto</span><strong>${m.telefone || '-'}</strong></div>
                    <div><span class="text-xs text-slate-400 font-bold block uppercase">Documento</span><strong>${m.tipo_documento || 'BI'}: ${m.numero_documento || '-'}</strong></div>
                    <div><span class="text-xs text-slate-400 font-bold block uppercase">Naturalidade</span><strong>${m.naturalidade || '-'}</strong></div>
                    <div><span class="text-xs text-slate-400 font-bold block uppercase">Bairro / Zona</span><strong>${m.bairro || '-'}</strong></div>
                    <div><span class="text-xs text-slate-400 font-bold block uppercase">Ano Conversão</span><strong>${m.ano_conversao || '-'}</strong></div>
                    <div><span class="text-xs text-slate-400 font-bold block uppercase">Data Batismo</span><strong>${m.data_batismo || '-'}</strong></div>
                    <div><span class="text-xs text-slate-400 font-bold block uppercase">Filiação</span><strong>${m.filiacao || '-'}</strong></div>
                </div>
                <div class="p-4 bg-blue-50/70 rounded-2xl border border-blue-200">
                    <span class="text-xs font-black text-blue-950 uppercase tracking-wider block mb-1">Histórico de Progressões Eclesiásticas</span>
                    <p class="text-slate-800 text-xs font-semibold whitespace-pre-line leading-relaxed">${m.progressoes || 'Nenhuma alteração de cargo registada.'}</p>
                </div>
                <button onclick="window.print()" class="w-full h-12 bg-slate-900 text-white font-black text-sm rounded-2xl shadow-md transition">🖨️ Imprimir Cartão de Membro</button>
            `;
            document.getElementById('modalMembro').classList.remove('hidden');
            document.getElementById('modalMembro').classList.add('flex');
        }

        function fecharModalMembro() {
            document.getElementById('modalMembro').classList.add('hidden');
            document.getElementById('modalMembro').classList.remove('flex');
        }

        const todasCategorias = {{ categorias_json | safe }};

        function atualizarCategoriasPorTipo() {
            const el = document.getElementById('selectTipoTransacao');
            if (!el) return;
            const tipo = el.value;
            const selectCat = document.getElementById('selectCategoriaTransacao');
            selectCat.innerHTML = '';
            const filtradas = todasCategorias.filter(c => c.tipo === tipo);
            filtradas.forEach(c => {
                const opt = document.createElement('option');
                opt.value = c.nome;
                opt.textContent = c.nome;
                selectCat.appendChild(opt);
            });
        }
        document.addEventListener('DOMContentLoaded', atualizarCategoriasPorTipo);

        function trocarAba(abaId) {
            document.querySelectorAll('.tab-content').forEach(el => el.classList.remove('active'));
            document.querySelectorAll('.tab-btn').forEach(btn => {
                btn.className = 'tab-btn px-4 py-2.5 rounded-2xl transition flex items-center gap-2 text-indigo-100 hover:bg-white/10 hover:text-white';
            });
            const abaAlvo = document.getElementById('aba-' + abaId);
            if (abaAlvo) abaAlvo.classList.add('active');
            const activeBtn = document.getElementById('btn-' + abaId);
            if (activeBtn) activeBtn.className = 'tab-btn px-4 py-2.5 rounded-2xl transition flex items-center gap-2 bg-gradient-to-r from-blue-600 to-indigo-600 text-white font-black shadow-md shadow-blue-900/40';
        }

        function filtrarTabela(inputId, tabelaId) {
            const filtro = document.getElementById(inputId).value.toLowerCase();
            const tr = document.getElementById(tabelaId).getElementsByTagName('tr');
            for (let i = 1; i < tr.length; i++) {
                tr[i].style.display = tr[i].innerText.toLowerCase().includes(filtro) ? '' : 'none';
            }
        }

        {% if pode_tesouraria %}
        // Paleta de Cores Eclesiástica Rica
        const paletaCores = ['#059669', '#2563eb', '#d97706', '#7c3aed', '#db2777', '#0891b2', '#ea580c', '#4f46e5'];
        const paletaSaidas = ['#e11d48', '#f97316', '#dc2626', '#9333ea', '#c026d3', '#b91c1c', '#64748b'];

        // 1. Gráfico Evolução Mensal Geral
        new Chart(document.getElementById('graficoEvolucaoMensal').getContext('2d'), {
            type: 'line',
            data: {
                labels: ['Jan', 'Fev', 'Mar', 'Abr', 'Mai', 'Jun', 'Jul', 'Ago', 'Set', 'Out', 'Nov', 'Dez'],
                datasets: [
                    { label: 'Entradas', data: {{ evolucao_entradas | safe }}, borderColor: '#059669', backgroundColor: 'rgba(5, 150, 105, 0.1)', borderWidth: 3, tension: 0.35, fill: true },
                    { label: 'Saídas', data: {{ evolucao_saidas | safe }}, borderColor: '#e11d48', backgroundColor: 'rgba(225, 29, 72, 0.05)', borderWidth: 3, tension: 0.35, fill: true }
                ]
            },
            options: { responsive: true, maintainAspectRatio: false }
        });

        // 2. Gráfico Situação dos Departamentos (Entradas vs Saídas)
        new Chart(document.getElementById('graficoFinancasDepto').getContext('2d'), {
            type: 'bar',
            data: {
                labels: {{ depto_fin_labels | safe }},
                datasets: [
                    { label: 'Entradas', data: {{ depto_fin_entradas | safe }}, backgroundColor: '#059669', borderRadius: 6 },
                    { label: 'Saídas', data: {{ depto_fin_saidas | safe }}, backgroundColor: '#e11d48', borderRadius: 6 }
                ]
            },
            options: { responsive: true, maintainAspectRatio: false }
        });

        // 3. Gráfico Fontes de Fundos (Entradas por Categoria)
        new Chart(document.getElementById('graficoFontesFundos').getContext('2d'), {
            type: 'doughnut',
            data: {
                labels: {{ fontes_labels | safe }},
                datasets: [{
                    data: {{ fontes_valores | safe }},
                    backgroundColor: paletaCores,
                    borderWidth: 2,
                    borderColor: '#ffffff'
                }]
            },
            options: {
                responsive: true,
                maintainAspectRatio: false,
                plugins: { legend: { position: 'bottom', labels: { boxWidth: 12, font: { size: 11, weight: 'bold' } } } }
            }
        });

        // 4. Gráfico Destino de Saídas (Despesas por Categoria)
        new Chart(document.getElementById('graficoSaidasCategorias').getContext('2d'), {
            type: 'doughnut',
            data: {
                labels: {{ saidas_labels | safe }},
                datasets: [{
                    data: {{ saidas_valores | safe }},
                    backgroundColor: paletaSaidas,
                    borderWidth: 2,
                    borderColor: '#ffffff'
                }]
            },
            options: {
                responsive: true,
                maintainAspectRatio: false,
                plugins: { legend: { position: 'bottom', labels: { boxWidth: 12, font: { size: 11, weight: 'bold' } } } }
            }
        });
        {% endif %}
    </script>
</body>
</html>
"""

with open(os.path.join("templates", "dashboard.html"), "w", encoding="utf-8") as f:
    f.write(html_code)

print("✓ 2/2: templates/dashboard.html atualizado com os 4 gráficos analíticos de tesouraria!")