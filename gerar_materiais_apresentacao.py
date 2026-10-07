import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt
from pptx.dml.color import RGBColor as PptxRGBColor

# ==============================================================================
# 1. GERAÇÃO DA APRESENTAÇÃO DE SLIDES (SIGAD_Apresentacao_Executiva.pptx)
# ==============================================================================
prs = Presentation()
prs.slide_width = PptxInches(13.333)  # Widescreen 16:9
prs.slide_height = PptxInches(7.5)

cor_azul = PptxRGBColor(30, 58, 138)
cor_dourada = PptxRGBColor(217, 119, 6)
cor_texto = PptxRGBColor(15, 23, 42)
cor_subtexto = PptxRGBColor(100, 116, 139)

slides_dados = [
    {
        "titulo": "SIGAD",
        "subtitulo": "Sistema Integrado de Gestão Eclesiástica\nPlataforma Administrativa de Padrão Nacional",
        "pontos": [
            "Governança Institucional, Membresia e Transparência Financeira",
            "Gestão Unificada de Liderança, Discipulado e Missão",
            "Solução Tecnológica Segura, Moderna e Adaptada ao Contexto Nacional"
        ]
    },
    {
        "titulo": "1. Visão Geral & Níveis de Acesso",
        "subtitulo": "Controlo rigoroso com separação clara de responsabilidades",
        "pontos": [
            "SuperAdmin Nacional / Provincial: Visão estatística consolidada de todas as congregações.",
            "Pastor Presidente / Liderança Local: Painel executivo, aprovação pastoral e direção estratégica.",
            "Secretaria Geral: Gestão de fichas de membros, departamentos, células e credenciais.",
            "Tesouraria: Lançamento de dízimos/ofertas, conciliação e controlo de métodos digitais."
        ]
    },
    {
        "titulo": "2. Passo 1: Secretaria & Membresia",
        "subtitulo": "Censo fidedigno e gestão demográfica do rebanho",
        "pontos": [
            "Registo Unificado: Fichas detalhadas com contactos, fotos, documentos e filiação.",
            "Segmentação Inteligente: Filtros por departamentos (Homens, Senhoras, Jovens, Crianças).",
            "Estrutura Geográfica: Acompanhamento de membros organizados por distritos, zonas e células.",
            "Emissão Instantânea: Cartões de membro e credenciais em PDF padronizados para impressão."
        ]
    },
    {
        "titulo": "3. Passo 2: Finanças Digitais Sem Custos",
        "subtitulo": "Transparência total sem dependência de APIs ou taxas intermediárias",
        "pontos": [
            "Múltiplos Canais de Entrada: Dinheiro físico (caixa), M-Pesa, e-Mola e Contas Bancárias.",
            "Rastreio por Código de Transação: Registo da referência de pagamento via SMS para conciliação.",
            "Classificação Contabilística: Distribuição de receitas por dízimos, ofertas, votos e departamentos.",
            "Comprovativos & Relatórios: Emissão imediata de recibos em PDF e exportação para Excel."
        ]
    },
    {
        "titulo": "4. Passo 3: Funil de Discipulado & Novos Convertidos",
        "subtitulo": "Garantindo que nenhuma alma alcançada se perca no processo",
        "pontos": [
            "Etapa 1 - Novo Decidido: Registo imediato no ato da entrega ou apelo congregacional.",
            "Etapa 2 - Em Discipulado / Aulas: Alocação direta de obreiro responsável e formação doutrinária.",
            "Etapa 3 - Apto para Batismo: Validação de aproveitamento bíblico pelo corpo pastoral.",
            "Etapa 4 - Batizado: Transição automática com 1 clique para a membresia plena da congregação."
        ]
    },
    {
        "titulo": "5. Passo 4: Painel Executivo SuperAdmin",
        "subtitulo": "Visão estratégica em tempo real para Convenções e Sedes Nacionais",
        "pontos": [
            "Quadro de Comando Consolidado: Métricas de crescimento congregacional de todo o país.",
            "Filtro Provincial e Distrital: Análise comparativa de congregações por região.",
            "Dossiês Executivos Oficiais: Emissão instantânea de relatórios consolidados em formato PDF.",
            "Tomada de Decisão Baseada em Dados: Alocação estratégica de obreiros e investimentos de missão."
        ]
    },
    {
        "titulo": "6. Vantagens Estratégicas do SIGAD",
        "subtitulo": "Por que o SIGAD é a escolha ideal para o crescimento ministerial?",
        "pontos": [
            "Independência Tecnológica: Sistema leve, funcional mesmo em ligações de baixa velocidade.",
            "Zero Taxas Bancárias Extras: Suporte digital manual de alta eficácia com M-Pesa e e-Mola.",
            "Padronização Administrativa: Uniformidade de relatórios e registos entre a sede e as congregações.",
            "Segurança de Dados: Backups seguros e controlo individualizado de permissões."
        ]
    }
]

for item in slides_dados:
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Layout em branco
    
    # Caixa de Título
    tb_titulo = slide.shapes.add_textbox(PptxInches(1.0), PptxInches(0.8), PptxInches(11.3), PptxInches(1.2))
    p_t = tb_titulo.text_frame.paragraphs[0]
    p_t.text = item["titulo"]
    p_t.font.name = "Arial"
    p_t.font.size = PptxPt(36)
    p_t.font.bold = True
    p_t.font.color.rgb = cor_azul
    
    # Subtítulo
    p_sub = tb_titulo.text_frame.add_paragraph()
    p_sub.text = item["subtitulo"]
    p_sub.font.name = "Arial"
    p_sub.font.size = PptxPt(16)
    p_sub.font.color.rgb = cor_dourada
    
    # Caixa de Conteúdo
    tb_conteudo = slide.shapes.add_textbox(PptxInches(1.0), PptxInches(2.5), PptxInches(11.3), PptxInches(4.2))
    tf_conteudo = tb_conteudo.text_frame
    tf_conteudo.word_wrap = True
    
    for i, pt_txt in enumerate(item["pontos"]):
        p_c = tf_conteudo.paragraphs[0] if i == 0 else tf_conteudo.add_paragraph()
        p_c.text = f"•  {pt_txt}"
        p_c.font.name = "Arial"
        p_c.font.size = PptxPt(18)
        p_c.font.color.rgb = cor_texto
        p_c.space_before = PptxPt(14)

prs.save("SIGAD_Apresentacao_Executiva.pptx")
print("✓ 1. SIGAD_Apresentacao_Executiva.pptx gerado com sucesso!")


# ==============================================================================
# 2. GERAÇÃO DO DOSSIÊ EXECUTIVO (.docx)
# ==============================================================================
doc = Document()

# Configuração de Margens
sections = doc.sections
for s in sections:
    s.top_margin = Inches(1.0)
    s.bottom_margin = Inches(1.0)
    s.left_margin = Inches(1.0)
    s.right_margin = Inches(1.0)

# Cabeçalho Principal
titulo_doc = doc.add_heading("SIGAD — Sistema Integrado de Gestão Eclesiástica", level=0)
titulo_doc.alignment = WD_ALIGN_PARAGRAPH.CENTER

p_sub = doc.add_paragraph("DOSSIÊ TÉCNICO & MANUAL EXECUTIVO DE APRESENTAÇÃO PASSO A PASSO\nPlataforma de Gestão Eclesiástica de Padrão Nacional")
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sub.runs[0].font.color.rgb = RGBColor(30, 58, 138)
p_sub.runs[0].font.bold = True

doc.add_paragraph("—" * 45).alignment = WD_ALIGN_PARAGRAPH.CENTER

# Introdução
doc.add_heading("1. Introdução e Propósito Estrutural", level=1)
doc.add_paragraph(
    "O SIGAD foi concebido para suprir as exigências administrativas contemporâneas de igrejas e convenções "
    "ministeriais. Aliando solidez de processos, rigor doutrinário e modernização tecnológica, a plataforma "
    "centraliza a gestão de membresia, a supervisão de lideranças e o rigor financeiro em uma interface ágil e acessível."
)

# Níveis de Acesso
doc.add_heading("2. Governança e Perfis de Acesso", level=1)
doc.add_paragraph(
    "A integridade das informações é sustentada por barreiras de segurança por papel funcional:\n"
    "• Nível Executivo Nacional (SuperAdmin): Acesso a consolidações estatísticas, métricas de crescimento e relatórios provinciais unificados.\n"
    "• Liderança Congregacional (Pastor Presidente): Visão integral da congregação local, acompanhamento pastoral e aprovação de atos administrativos.\n"
    "• Secretaria: Inscrição de fiéis, organização de células/zonas, controlo de presenças e impressão de credenciais oficiais.\n"
    "• Tesouraria: Registo de receitas e despesas, controlo de carteiras móveis e emissão de recibos e balancetes."
)

# Passo a Passo Operacional
doc.add_heading("3. Passo a Passo Operacional por Módulos", level=1)

doc.add_heading("Etapa 1: Gestão de Membresia e Secretaria", level=2)
doc.add_paragraph(
    "1. Registo de Membros: Formulário completo abrangendo contactos, documentação civil, filiação, datas de conversão e batismo.\n"
    "2. Distribuição Departamental: Segmentação automática por homens, senhoras, juventude, adolescentes e crianças.\n"
    "3. Credenciação: Geração direta de cartões de membro em PDF com foto e QR Code para identificação eclesiástica imediata."
)

doc.add_heading("Etapa 2: Módulo Financeiro e Carteiras Móveis", level=2)
doc.add_paragraph(
    "1. Lançamentos sem Intermediários: Registro de dízimos e ofertas em dinheiro físico, M-Pesa, e-Mola e transferências bancárias.\n"
    "2. Conciliação por Referência: Registo do código alfanumérico da transação móvel para verificação ágil com o extrato da operadora.\n"
    "3. Prestação de Contas: Emissão instantânea de recibos individuais em PDF e exportação contínua de extratos para planilhas eletrónicas."
)

doc.add_heading("Etapa 3: Funil de Novos Convertidos e Discipulado", level=2)
doc.add_paragraph(
    "1. Registo da Decisão: Recepção imediata de dados do novo convertido após o apelo no culto.\n"
    "2. Atribuição de Discipulador: Indicação formal do obreiro responsável pelas visitas e acompanhamento pessoal.\n"
    "3. Progressão de Fases: Acompanhamento da passagem pelas classes de doutrina até à autorização para o batismo nas águas e admissão no rol de membros."
)

doc.add_heading("Etapa 4: Painel Estatístico Consolidado (SuperAdmin)", level=2)
doc.add_paragraph(
    "1. Centralização Nacional: A sede provincial ou geral acompanha indicadores de batismos, crescimento de membros e atividade das congregações.\n"
    "2. Relatório Executivo em PDF: Emissão de dossiês gerenciais completos com gráficos e tabelas padronizados para convenções e reuniões ministeriais."
)

doc.save("SIGAD_Dossie_Executivo.docx")
print("✓ 2. SIGAD_Dossie_Executivo.docx gerado com sucesso!")


# ==============================================================================
# 3. GERAÇÃO DO GUIA RÁPIDO DO UTILIZADOR (.html pronto para impressão PDF)
# ==============================================================================
html_guia = '''<!DOCTYPE html>
<html lang="pt">
<head>
    <meta charset="UTF-8">
    <title>SIGAD - Guia Rápido do Utilizador</title>
    <style>
        body { font-family: 'Segoe UI', system-ui, sans-serif; line-height: 1.6; color: #1e293b; background: #f8fafc; padding: 30px; margin: 0; }
        .page { max-width: 800px; margin: 0 auto; background: white; padding: 40px; border-radius: 16px; box-shadow: 0 4px 16px rgba(0,0,0,0.06); }
        .header { display: flex; align-items: center; justify-content: space-between; border-bottom: 2px solid #1e3a8a; padding-bottom: 15px; margin-bottom: 25px; }
        .title { color: #1e3a8a; margin: 0; font-size: 24px; font-weight: 800; }
        .subtitle { color: #d97706; margin: 0; font-size: 13px; font-weight: 700; text-transform: uppercase; }
        .card { background: #f1f5f9; border-left: 4px solid #1e3a8a; padding: 15px; border-radius: 8px; margin-bottom: 20px; }
        .card-gold { border-left-color: #d97706; }
        .card-green { border-left-color: #059669; }
        h3 { color: #0f172a; margin-top: 0; font-size: 16px; }
        ul { margin: 5px 0 0; padding-left: 20px; font-size: 14px; }
        li { margin-bottom: 6px; }
        .badge { display: inline-block; background: #1e3a8a; color: white; padding: 2px 8px; border-radius: 12px; font-size: 11px; font-weight: bold; }
        @media print { body { background: white; padding: 0; } .page { box-shadow: none; padding: 0; } }
    </style>
</head>
<body>
    <div class="page">
        <div class="header">
            <div>
                <h1 class="title">SIGAD</h1>
                <p class="subtitle">Guia de Apoio Rápido às Igrejas Locais</p>
            </div>
            <div>
                <span class="badge">Versão Oficial</span>
            </div>
        </div>

        <div class="card">
            <h3>Secretaria: Cadastrar e Atualizar Membros</h3>
            <ul>
                <li>Aceda à aba <b>Membros</b> e clique em <b>Novo Membro</b>.</li>
                <li>Preencha nome completo, contacto de telefone, endereço/bairro e departamento.</li>
                <li>Gere a credencial oficial através do botão <b>📄 Cartão de Membro (PDF)</b>.</li>
            </ul>
        </div>

        <div class="card card-gold">
            <h3>Tesouraria: Lançamento de Entradas (M-Pesa, e-Mola ou Dinheiro)</h3>
            <ul>
                <li>No painel financeiro, selecione o tipo (<b>Entrada</b>) e a categoria (<b>Dízimo</b> ou <b>Oferta</b>).</li>
                <li>Escolha o método correspondente: <b>💵 Dinheiro</b>, <b>📱 M-Pesa</b>, <b>📱 e-Mola</b> ou <b>🏦 Banco</b>.</li>
                <li>Se for pagamento digital, insira o <b>Código da Transação</b> fornecido pelo SMS do membro.</li>
                <li>Clique em <b>Gravar Movimento</b> e emita o recibo em PDF de imediato.</li>
            </ul>
        </div>

        <div class="card card-green">
            <h3>Discipulado: Gestão de Novos Convertidos</h3>
            <ul>
                <li>Ao registar uma nova conversão, marque a fase como <b>Novo Decidido</b>.</li>
                <li>Defina o nome do <b>Discipulador Responsável</b> para garantir acompanhamento direto.</li>
                <li>Conforme a frequência nas aulas bíblicas, avance a fase para <b>Em Aulas</b> e depois <b>Pronto para Batismo</b>.</li>
                <li>Após a cerimónia nas águas, clique em <b>Batizado</b> para integrá-lo formalmente como Membro Efetivo.</li>
            </ul>
        </div>

        <p style="text-align: center; font-size: 12px; color: #94a3b8; margin-top: 30px;">
            SIGAD • Sistema Integrado de Gestão Eclesiástica • Padrão Administrativo Nacional
        </p>
    </div>
</body>
</html>
'''

with open("SIGAD_Guia_Rapido_Utilizador.html", "w", encoding="utf-8") as f:
    f.write(html_guia)
print("✓ 3. SIGAD_Guia_Rapido_Utilizador.html gerado com sucesso!")